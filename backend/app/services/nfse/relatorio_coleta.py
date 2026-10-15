"""Planilha de uma coleta automática de NFS-e (fase 2, 25/09/2026).

Uma planilha por rodada, salva na pasta das NFS-e da rede em
`_relatorios/AAAA-MM/`, e anexada ao informativo. As abas repetem o que o
detalhe da tela mostra, com as notas inteiras:

- Visão Geral: os números da rodada;
- Notas novas: cada nota que entrou, com as colunas do relatório ABRASF;
- Por empresa: o resultado de cada empresa consultada;
- Não consultadas: falha e adiada, com o motivo;
- Certificados: os que vencem em até 30 dias (ou já venceram).

"Nota nova da rodada" é a nota de origem ADN gravada ENTRE o início e o fim da
execução, de uma empresa que a rodada consultou. Sai do banco, não do JSON da
execução: o JSON guarda contagens, e a planilha precisa da nota.
"""
import io
import json
import logging
import os
from datetime import datetime

from openpyxl import Workbook
from sqlalchemy import func

from ...core.fuso import hora_local
from ...core.templating import moeda_br
from ...domain.models.nfse import NfseDocument
from .. import planilha_estilo as estilo
from . import storage

logger = logging.getLogger(__name__)

PASTA = "_relatorios"

TIPOS_ERRO = {
    "adiada": "Adiada pelo ADN",
    "certificado": "Certificado / rede",
    "resposta": "Resposta do ADN",
    "cadastro": "Cadastro",
    "inesperado": "Erro inesperado",
}

# As colunas do relatório ABRASF da tela (routes/nfse.py), com o CNPJ dito no
# título: é o título que faz a planilha formatar a coluna como TEXTO, e CNPJ
# como número perde o zero da frente.
TITULOS_TROCADOS = {"prestador": "CNPJ Prestador", "tomador": "CNPJ Tomador"}


def _detalhe(execucao) -> dict:
    try:
        return json.loads(execucao.detalhe or "{}")
    except ValueError:
        return {}


def _tenants_da_rodada(execucao) -> list[str]:
    return [e.get("tenant_id") for e in _detalhe(execucao).get("empresas", []) if e.get("tenant_id")]


def _filtro_rodada(execucao):
    fim = execucao.fim or datetime.utcnow()
    return (NfseDocument.origem == "adn",
            NfseDocument.created_at >= execucao.inicio,
            NfseDocument.created_at <= fim,
            NfseDocument.tenant_id.in_(_tenants_da_rodada(execucao) or ["-"]))


def notas_da_execucao(db, execucao) -> list:
    return (db.query(NfseDocument).filter(*_filtro_rodada(execucao))
            .order_by(NfseDocument.tenant_id, NfseDocument.dt_emissao).all())


def valores_por_empresa(db, execucao) -> dict:
    """tenant_id -> {"qtd", "valor", "retencoes"} das notas novas da rodada.

    Retenções = IR + PIS + COFINS + CSLL + INSS retidos na nota. ISS fica de
    fora: a retenção dele depende do município e vem como indicador, não como
    valor retido.
    """
    if not execucao.notas_novas:
        return {}
    c = func.coalesce
    linhas = (db.query(NfseDocument.tenant_id, func.count(),
                       func.sum(c(NfseDocument.valor_servico, 0)),
                       func.sum(c(NfseDocument.ir, 0) + c(NfseDocument.pis, 0)
                                + c(NfseDocument.cofins, 0) + c(NfseDocument.csll, 0)
                                + c(NfseDocument.inss, 0)))
              .filter(*_filtro_rodada(execucao))
              .group_by(NfseDocument.tenant_id).all())
    return {str(t): {"qtd": q, "valor": float(v or 0), "retencoes": float(r or 0)}
            for t, q, v, r in linhas}


def nome_arquivo(execucao) -> str:
    local = hora_local(execucao.inicio)
    return f"coleta_nfse_{local:%Y-%m-%d_%H%M}_{execucao.id[:8]}.xlsx"


def _data(valor):
    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y")
    return valor


def gerar_planilha(db, execucao, nome_automacao: str, rotulo_status: str) -> bytes:
    from ...api.routes.nfse import COLUNAS   # mesmas colunas da tela de NFS-e

    detalhe = _detalhe(execucao)
    empresas = detalhe.get("empresas", [])
    certificados = detalhe.get("certificados_vencendo", [])
    valores = valores_por_empresa(db, execucao)
    nomes = {e.get("tenant_id"): e.get("empresa", "") for e in empresas}

    wb = Workbook()
    wb.remove(wb.active)

    # Notas novas
    ws = wb.create_sheet("Notas novas")
    ws.append(["Empresa"] + [TITULOS_TROCADOS.get(campo, rotulo) for campo, rotulo in COLUNAS])
    for n in notas_da_execucao(db, execucao):
        ws.append([nomes.get(str(n.tenant_id), "")]
                  + [_data(getattr(n, campo, None)) for campo, _r in COLUNAS])

    # Por empresa
    ws = wb.create_sheet("Por empresa")
    ws.append(["Empresa", "CNPJ", "Situação", "Notas novas", "Notas atualizadas",
               "Valor das notas novas", "Retenções", "Motivo"])
    for e in empresas:
        v = valores.get(e.get("tenant_id"), {})
        situacao = (TIPOS_ERRO.get(e.get("tipo_erro"), "Falha") if e.get("erro")
                    else "Consultada")
        ws.append([e.get("empresa"), e.get("cnpj"), situacao, e.get("novos", 0),
                   e.get("atualizados", 0), v.get("valor", 0.0), v.get("retencoes", 0.0),
                   e.get("erro") or ""])

    # Não consultadas
    ws = wb.create_sheet("Não consultadas")
    ws.append(["Empresa", "CNPJ", "Tipo", "Motivo"])
    for e in empresas:
        if e.get("erro"):
            ws.append([e.get("empresa"), e.get("cnpj"),
                       TIPOS_ERRO.get(e.get("tipo_erro"), "Falha"), e.get("erro")])

    # Certificados
    ws = wb.create_sheet("Certificados")
    ws.append(["Empresa", "CNPJ", "Validade", "Dias", "Situação"])
    for c in certificados:
        validade = c.get("validade") or ""
        ws.append([c.get("empresa"), c.get("cnpj"),
                   f"{validade[8:10]}/{validade[5:7]}/{validade[:4]}" if validade else "",
                   c.get("dias"),
                   "Vencido" if c.get("vencido") else "Vence em breve"])

    estilo.estilizar_planilha(wb, cores={
        "Notas novas": estilo.VERDE, "Não consultadas": estilo.VERMELHO,
        "Certificados": "D97706"})

    total_valor = sum(v["valor"] for v in valores.values())
    total_ret = sum(v["retencoes"] for v in valores.values())
    local = hora_local(execucao.inicio)
    estilo.montar_capa(wb, "Coleta de NFS-e", [
        ("Automação", nome_automacao),
        ("Início", f"{local:%d/%m/%Y %H:%M}"),
        ("Origem", "Disparada à mão" if execucao.gatilho == "manual" else "Agenda"),
        ("Situação", rotulo_status),
        ("Empresas consultadas", f"{execucao.empresas_ok} de {execucao.empresas_total}"),
        ("Notas novas", int(execucao.notas_novas or 0)),
        ("Notas atualizadas", int(execucao.notas_atualizadas or 0)),
        ("Valor das notas novas", f"R$ {moeda_br(total_valor)}"),
        ("Retenções (IR, PIS, COFINS, CSLL, INSS)", f"R$ {moeda_br(total_ret)}"),
        ("Empresas com falha", int(execucao.empresas_falha or 0)),
        ("Consultas adiadas pelo ADN", int(execucao.empresas_adiadas or 0)),
        ("Certificados vencendo em 30 dias", int(execucao.certificados_vencendo or 0)),
    ], subtitulo=execucao.mensagem or "")

    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


def salvar(execucao, conteudo: bytes) -> str | None:
    """Grava na pasta das NFS-e. Devolve o caminho relativo, ou None se a
    pasta da rede estiver fora — a coleta não para por isso."""
    local = hora_local(execucao.inicio)
    rel = os.path.join(PASTA, f"{local:%Y-%m}", nome_arquivo(execucao))
    destino = os.path.join(storage.raiz(), rel)
    try:
        os.makedirs(os.path.dirname(destino), exist_ok=True)
        with open(destino, "wb") as f:
            f.write(conteudo)
        return rel
    except OSError:
        logger.warning("Não foi possível salvar o relatório da coleta em %s", destino,
                       exc_info=True)
        return None
