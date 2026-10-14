"""Informativo: resumo das auditorias do período (25/09/2026).

Para quem acompanha o escritório: quantas auditorias rodaram, quanto de
inconsistência apareceu e QUANTO DINHEIRO está em jogo — a estornar, a
aproveitar, a recolher —, empresa por empresa.

Três regras, as mesmas do Radar de Atividade e pelos mesmos motivos:

1. **Nunca abre `result_data` para somar.** O valor em jogo vem das colunas
   `vl_*` de `audit_tasks`, que o motor grava ao concluir. Auditoria anterior a
   essas colunas tem o relatório lido UMA vez, aqui, e as colunas preenchidas —
   da segunda vez em diante já é coluna. NULO nunca vira zero.
2. **Rodar a mesma competência várias vezes é rotina.** Com
   `ultima_por_competencia` (o padrão), entra só a auditoria mais recente de
   cada empresa e competência; somar todas faria a empresa que foi corrigida
   parecer a pior justamente por causa do retrabalho.
3. **O que não concluiu aparece.** Auditoria que falhou, esgotou o tempo ou
   está parada esperando a confirmação da IA vai numa lista própria — o resumo
   que só mostra o que deu certo esconde exatamente o que precisa de gente.
"""
import io
import json
import logging

from openpyxl import Workbook
from sqlalchemy import select, update

from ...core.competencia import de_data
from ...core.fuso import hora_local
from ...core.templating import build_env, moeda_br
from ...domain.models.audit import AuditTask
from ...domain.models.auth_models import Tenant, UserProfile
from .. import planilha_estilo as estilo
from ..audit_engine.resumo_financeiro import valores_para_colunas
from . import janela as jan

logger = logging.getLogger(__name__)

SITUACOES = {
    "erro": "Falhou",
    "falha": "Falhou",
    "timeout": "Tempo esgotado",
    "aguardando_ia": "Parada esperando a confirmação da IA",
    "pendente": "Em andamento",
    "processando": "Em andamento",
}
COLUNAS_VALOR = ("vl_estornar", "vl_aproveitar", "vl_recolher", "vl_conferir")
LIMITE_EMAIL = 25

_env = None


def _ambiente():
    global _env
    if _env is None:
        _env = build_env()
    return _env


def _competencia_rotulo(comp: str | None) -> str:
    return f"{comp[5:7]}/{comp[:4]}" if comp and len(comp) == 7 else "—"


def preencher_valores(db, linhas) -> int:
    """Preenche as colunas de auditorias antigas, lendo o relatório de UMA por
    vez (nunca todos na memória). Devolve quantas foram preenchidas."""
    preenchidas = 0
    for linha in linhas:
        bruto = db.execute(select(AuditTask.result_data)
                           .where(AuditTask.id == linha.id)).scalar()
        if not bruto:
            continue
        try:
            resultado = json.loads(bruto)
        except ValueError:
            continue
        valores = valores_para_colunas(resultado)
        resumo = resultado.get("resumo") or {}
        if linha.notas_processadas is None:
            valores["notas_processadas"] = int(resumo.get("notas_processadas") or 0)
        if linha.inconsistencias is None:
            valores["inconsistencias"] = int(resumo.get("inconsistencias") or 0)
        if linha.competencia is None:
            valores["competencia"] = de_data(resultado.get("periodo_inicio"))
        db.execute(update(AuditTask).where(AuditTask.id == linha.id).values(**valores))
        db.commit()
        del resultado, bruto
        preenchidas += 1
    return preenchidas


def _consultar(db, inicio, fim):
    """As auditorias da janela — coluna a coluna, SEM result_data."""
    return db.execute(select(
        AuditTask.id, AuditTask.tenant_id, AuditTask.user_id, AuditTask.status,
        AuditTask.created_at, AuditTask.competencia, AuditTask.file_name,
        AuditTask.notas_processadas, AuditTask.inconsistencias,
        AuditTask.vl_estornar, AuditTask.vl_aproveitar,
        AuditTask.vl_recolher, AuditTask.vl_conferir,
    ).where(AuditTask.created_at >= inicio, AuditTask.created_at < fim)
     .order_by(AuditTask.created_at.desc())).all()


def coletar(db, automacao, config: dict, inicio, fim) -> dict:
    escopo = None
    if getattr(automacao, "escopo", "todas") == "selecionadas":
        escopo = {x.strip() for x in (automacao.empresas or "").split(",") if x.strip()}

    def no_escopo(linha):
        return escopo is None or str(linha.tenant_id) in escopo

    linhas = [l for l in _consultar(db, inicio, fim) if no_escopo(l)]
    concluidas = [l for l in linhas if l.status == "concluido"]

    faltando = [l for l in concluidas if any(getattr(l, c) is None for c in COLUNAS_VALOR)]
    if faltando:
        preencher_valores(db, faltando)
        linhas = [l for l in _consultar(db, inicio, fim) if no_escopo(l)]
        concluidas = [l for l in linhas if l.status == "concluido"]

    repetidas = 0
    if config.get("ultima_por_competencia", True):
        vistas, unicas = set(), []
        for l in concluidas:                     # já vêm da mais recente para a mais antiga
            chave = (str(l.tenant_id), l.competencia or f"sem-competencia:{l.id}")
            if chave in vistas:
                repetidas += 1
                continue
            vistas.add(chave)
            unicas.append(l)
        concluidas = unicas

    empresas = {str(t.id): t for t in db.query(Tenant).all()}
    pessoas = {str(p.id): (p.name or "") for p in db.query(UserProfile.id, UserProfile.name).all()}

    def base(l):
        t = empresas.get(str(l.tenant_id))
        return {"task_id": l.id, "tenant_id": str(l.tenant_id),
                "empresa": t.name if t else "(empresa removida)",
                "cnpj": (t.cnpj or "") if t else "",
                "competencia": l.competencia or "",
                "competencia_rotulo": _competencia_rotulo(l.competencia),
                "quando": f"{hora_local(l.created_at):%d/%m/%Y %H:%M}",
                "por": pessoas.get(str(l.user_id or ""), ""),
                "arquivo": l.file_name or ""}

    auditorias = []
    for l in concluidas:
        # Coluna que continuou nula depois do preenchimento é relatório que não
        # deu para ler: a auditoria entra na contagem, mas com "sem valor
        # apurado" — somá-la como R$ 0,00 afirmaria que não havia nada em jogo.
        sem_valor = any(getattr(l, c) is None for c in COLUNAS_VALOR)
        v = {c: float(getattr(l, c) or 0.0) for c in COLUNAS_VALOR}
        auditorias.append({
            **base(l),
            "notas": int(l.notas_processadas or 0),
            "inconsistencias": int(l.inconsistencias or 0),
            "estornar": v["vl_estornar"], "aproveitar": v["vl_aproveitar"],
            "recolher": v["vl_recolher"], "conferir": v["vl_conferir"],
            "em_jogo": round(v["vl_estornar"] + v["vl_aproveitar"] + v["vl_recolher"], 2),
            "sem_valor": sem_valor,
        })
    auditorias.sort(key=lambda a: (-a["em_jogo"], -a["inconsistencias"], a["empresa"]))

    nao_concluidas = [{**base(l), "situacao": SITUACOES.get(l.status, l.status)}
                      for l in linhas if l.status != "concluido"]

    totais = {
        "auditorias": len(auditorias),
        "empresas": len({a["tenant_id"] for a in auditorias}),
        "notas": sum(a["notas"] for a in auditorias),
        "inconsistencias": sum(a["inconsistencias"] for a in auditorias),
        "estornar": round(sum(a["estornar"] for a in auditorias), 2),
        "aproveitar": round(sum(a["aproveitar"] for a in auditorias), 2),
        "recolher": round(sum(a["recolher"] for a in auditorias), 2),
        "conferir": round(sum(a["conferir"] for a in auditorias), 2),
        "nao_concluidas": len(nao_concluidas),
        "repetidas_descartadas": repetidas,
        "sem_valor": sum(1 for a in auditorias if a["sem_valor"]),
    }
    totais["em_jogo"] = round(totais["estornar"] + totais["aproveitar"] + totais["recolher"], 2)
    return {"tipo": "resumo_auditorias", "periodo": jan.rotulo(inicio, fim),
            "periodo_curto": jan.rotulo_curto(inicio, fim), "config": config,
            "totais": totais, "auditorias": auditorias, "nao_concluidas": nao_concluidas}


def desfecho(dados: dict) -> tuple[str, str, bool]:
    """(status, mensagem, tem_novidade)."""
    t = dados["totais"]
    if not t["auditorias"] and not t["nao_concluidas"]:
        return "sem_novidade", "Nenhuma auditoria no período.", False
    partes = [f"{t['auditorias']} auditoria{'s' if t['auditorias'] != 1 else ''} de "
              f"{t['empresas']} empresa{'s' if t['empresas'] != 1 else ''}",
              f"{t['inconsistencias']} inconsistência{'s' if t['inconsistencias'] != 1 else ''}",
              f"R$ {moeda_br(t['em_jogo'])} em jogo"]
    if t["nao_concluidas"]:
        partes.append(f"{t['nao_concluidas']} não concluída{'s' if t['nao_concluidas'] != 1 else ''}")
    return "ok", " · ".join(partes), True


# ------------------------------------------------------------------ saídas

def nome_arquivo(execucao) -> str:
    return f"resumo_auditorias_{hora_local(execucao.inicio):%Y-%m-%d_%H%M}_{execucao.id[:8]}.xlsx"


def gerar_planilha(dados: dict, nome_automacao: str) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Auditorias")
    ws.append(["Empresa", "CNPJ", "Competência", "Data da auditoria", "Por",
               "Notas analisadas", "Inconsistências", "Valor a estornar",
               "Valor a aproveitar", "Valor a recolher", "Valor em jogo",
               "Valor a conferir", "Arquivo", "Observação"])
    for a in dados["auditorias"]:
        # Sem valor apurado, as células de valor ficam VAZIAS, não zeradas.
        valores = ([None] * 5 if a.get("sem_valor") else
                   [a["estornar"], a["aproveitar"], a["recolher"], a["em_jogo"], a["conferir"]])
        ws.append([a["empresa"], a["cnpj"], a["competencia_rotulo"], a["quando"], a["por"],
                   a["notas"], a["inconsistencias"], *valores, a["arquivo"],
                   "Relatório ilegível: valor não apurado" if a.get("sem_valor") else ""])

    ws = wb.create_sheet("Não concluídas")
    ws.append(["Empresa", "CNPJ", "Competência", "Data da auditoria", "Por", "Situação", "Arquivo"])
    for a in dados["nao_concluidas"]:
        ws.append([a["empresa"], a["cnpj"], a["competencia_rotulo"], a["quando"], a["por"],
                   a["situacao"], a["arquivo"]])

    estilo.estilizar_planilha(wb, cores={"Não concluídas": estilo.VERMELHO})
    t = dados["totais"]
    estilo.montar_capa(wb, "Resumo das auditorias", [
        ("Automação", nome_automacao),
        ("Período", dados["periodo"]),
        ("Auditorias", int(t["auditorias"])),
        ("Empresas", int(t["empresas"])),
        ("Notas analisadas", int(t["notas"])),
        ("Inconsistências", int(t["inconsistencias"])),
        ("A estornar", f"R$ {moeda_br(t['estornar'])}"),
        ("A aproveitar", f"R$ {moeda_br(t['aproveitar'])}"),
        ("A recolher", f"R$ {moeda_br(t['recolher'])}"),
        ("Em jogo (estornar + aproveitar + recolher)", f"R$ {moeda_br(t['em_jogo'])}"),
        ("A conferir (exposição, não é imposto)", f"R$ {moeda_br(t['conferir'])}"),
        ("Não concluídas", int(t["nao_concluidas"])),
        *([("Sem valor apurado (relatório ilegível)", int(t["sem_valor"]))]
          if t.get("sem_valor") else []),
    ], subtitulo=("Só a auditoria mais recente de cada empresa e competência"
                  if dados["config"].get("ultima_por_competencia", True)
                  else "Todas as auditorias do período, inclusive as repetidas"))
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


def assunto(dados: dict) -> str:
    t = dados["totais"]
    if not t["auditorias"] and not t["nao_concluidas"]:
        return f"[Synax] Resumo das auditorias {dados['periodo_curto']} — nenhuma auditoria no período"
    partes = [f"{t['auditorias']} auditoria{'s' if t['auditorias'] != 1 else ''}",
              f"R$ {moeda_br(t['em_jogo'])} em jogo"]
    if t["nao_concluidas"]:
        partes.append(f"{t['nao_concluidas']} não concluída{'s' if t['nao_concluidas'] != 1 else ''}")
    return f"[Synax] Resumo das auditorias {dados['periodo_curto']} — " + " · ".join(partes)


def montar_email(dados: dict, nome_automacao: str, com_anexo: bool) -> tuple[str, str, str]:
    contexto = {**dados, "nome": nome_automacao, "com_anexo": com_anexo,
                "limite": LIMITE_EMAIL, "mostradas": dados["auditorias"][:LIMITE_EMAIL]}
    html = _ambiente().get_template("emails/informativo_resumo.html").render(contexto)
    return assunto(dados), html, _texto(contexto)


def _texto(d: dict) -> str:
    t = d["totais"]
    linhas = [f"Resumo das auditorias — {d['nome']}", f"Período: {d['periodo']}", "",
              f"Auditorias: {t['auditorias']} ({t['empresas']} empresa(s)) · "
              f"inconsistências: {t['inconsistencias']}",
              f"A estornar: R$ {moeda_br(t['estornar'])} · a aproveitar: "
              f"R$ {moeda_br(t['aproveitar'])} · a recolher: R$ {moeda_br(t['recolher'])}",
              f"Em jogo: R$ {moeda_br(t['em_jogo'])}", ""]
    if d["mostradas"]:
        linhas.append("POR EMPRESA (maior valor em jogo primeiro)")
        for a in d["mostradas"]:
            valor = ("valor não apurado" if a.get("sem_valor")
                     else f"R$ {moeda_br(a['em_jogo'])} em jogo")
            linhas.append(f"   - {a['empresa']} · competência {a['competencia_rotulo']} · "
                          f"{a['inconsistencias']} inconsistência(s) · {valor}")
        if len(d["auditorias"]) > len(d["mostradas"]):
            linhas.append(f"   ... e mais {len(d['auditorias']) - len(d['mostradas'])} na planilha.")
    else:
        linhas.append("Nenhuma auditoria concluída no período.")
    if d["nao_concluidas"]:
        linhas += ["", "NÃO CONCLUÍDAS"]
        for a in d["nao_concluidas"]:
            linhas.append(f"   - {a['empresa']} · {a['quando']} · {a['situacao']}")
    linhas += ["", ("A planilha completa segue em anexo." if d["com_anexo"] else ""),
               "Enviado automaticamente pelo Synax (Administração → Automações)."]
    return "\n".join(l for l in linhas if l is not None)
