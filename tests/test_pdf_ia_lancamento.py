"""CFOP 1933 x 2933 das notas de serviço lidas do PDF — cálculo autônomo.

As duas UFs saem da PRÓPRIA NOTA (a IA lê prestador e tomador do documento).
Nenhuma consulta a banco, nenhuma empresa selecionada na tela.

Duas versões anteriores erraram aqui, e é isso que esta suíte impede de voltar:

1. `UF_EMPRESA = "ES"` fixo no módulo (relato de 26/08/2026: "temos uma empresa
   do RJ, nas notas RJ x RJ deveria ser 1933 e ele estava lendo 2933, porque a
   base dele foi feita como se todas as empresas fossem no ES"). No PDF real de
   abril, 19 das 59 notas mudavam de CFOP por causa disso.
2. A UF passou a vir do `Tenant.uf` — cálculo certo, mas exigia a empresa
   cadastrada, e as 170 estavam sem UF: a tela não deixava fazer nada.

Terceira trava: quando falta alguma UF, o CFOP é presumido 1933 (nada bloqueia,
tudo entra no TXT) MAS a nota sai marcada. Presunção silenciosa foi como o
problema 1 durou meses.
"""
import os, sys, io, ast

sys.path.insert(0, os.path.abspath('.'))

import dotenv
dotenv.load_dotenv = lambda *a, **k: False
import dotenv.main
dotenv.main.load_dotenv = lambda *a, **k: False

from backend.app.services.nfse import pdf_ia_lancamento as mod

results = []
def check(name, cond, detail=""):
    results.append((name, bool(cond), detail))
    print(("  OK  " if cond else " FALHA") + f" | {name}" + (f"  [{detail}]" if detail and not cond else ""))


def nota(**kw):
    base = {
        "cnpj_prestador": "99888777000100", "fornecedor": "FORNECEDOR GENERICO LTDA",
        "uf_prestador": "ES", "cnpj_tomador": "11222333000181", "uf_tomador": "ES",
        "data_emissao": "05/07/2026", "data_entrada": "08/07/2026",
        "num_nota": "1", "valor": 1000.0, "retencao_ir": 15.0, "retencao_pis": 6.5,
        "retencao_cofins": 30.0, "retencao_csll": 10.0,
    }
    base.update(kw)
    return base


# ============================================================================
print("\n=== 1. A regra, isolada: as duas UFs vêm da nota ===")
check("mesma UF -> 1933, lido", mod.cfop_da_nota("ES", "ES") == ("1933", False))
check("UF diferente -> 2933, lido", mod.cfop_da_nota("SP", "ES") == ("2933", False))
check("sem UF do prestador -> presumido", mod.cfop_da_nota("", "ES") == ("1933", True))
check("sem UF do tomador -> presumido", mod.cfop_da_nota("SP", "") == ("1933", True))
check("sem nenhuma das duas -> presumido", mod.cfop_da_nota("", "") == ("1933", True))
check("espaço e minúscula normalizam", mod.cfop_da_nota(" rj ", "RJ") == ("1933", False))
check("UF de 3 letras é inválida -> presumido", mod.cfop_da_nota("SPO", "ES")[1] is True)
check("UF numérica é inválida -> presumido", mod.cfop_da_nota("32", "ES")[1] is True)


# ============================================================================
print("\n=== 2. O BUG DO RELATO: nota RJ x RJ é interna, venha de onde vier ===")
rj = nota(num_nota="900", uf_prestador="RJ", uf_tomador="RJ",
          fornecedor="PRESTADOR CARIOCA LTDA")
txt = mod.gerar_txt_importacao([rj])
check("prestador RJ + tomador RJ = 1933 (era 2933)", ";1933;" in txt, txt)
check("e NÃO sai 2933", ";2933;" not in txt, txt)
check("sem presunção: as duas UFs foram lidas",
      mod.preparar_notas([rj])[0]["_presumido"] is False)

# O ponto da reformulação: o resultado não depende de nada fora da nota.
sp_para_rj = nota(num_nota="901", uf_prestador="SP", uf_tomador="RJ")
check("prestador SP + tomador RJ = 2933",
      ";2933;" in mod.gerar_txt_importacao([sp_para_rj]))

# Varredura: mesma UF é sempre interna; par diferente é sempre interestadual.
UFS = ["AC","AL","AP","AM","BA","CE","DF","ES","GO","MA","MT","MS","MG","PA","PB",
       "PR","PE","PI","RJ","RN","RS","RO","RR","SC","SP","SE","TO"]
check(f"mesma UF é 1933 nas {len(UFS)} UFs",
      all(mod.cfop_da_nota(uf, uf) == ("1933", False) for uf in UFS))
check("qualquer par de UFs diferentes é 2933",
      all(mod.cfop_da_nota(a, b) == ("2933", False) for a in UFS for b in UFS if a != b))

# Nenhuma UF literal decidindo nada, e nenhum resquício do desenho antigo.
_fonte = io.open(mod.__file__, encoding="utf-8").read()
_doc = ast.get_docstring(ast.parse(_fonte)) or ""
_codigo = "\n".join(
    l for l in _fonte.replace(_doc, "", 1).split("BASE_CNPJ = {")[0].splitlines()
    if not l.strip().startswith("#")
)
check("nenhuma UF fixa no código executável",
      all(f'"{uf}"' not in _codigo for uf in ["ES", "RJ", "SP", "MG", "BA"]),
      [uf for uf in ["ES", "RJ", "SP", "MG", "BA"] if f'"{uf}"' in _codigo])
check("não existe constante UF_EMPRESA", "UF_EMPRESA" not in _codigo)
# O módulo inteiro sem docstring nem comentários: a docstring cita `Tenant.uf`
# ao contar o desenho antigo, e isso é documentação, não dependência.
_codigo_todo = "\n".join(
    l for l in _fonte.replace(_doc, "", 1).splitlines()
    if not l.strip().startswith("#")
)
check("o módulo não toca em banco de dados",
      all(t not in _codigo_todo.lower() for t in ["tenant", "session", "sqlalchemy", "query("]),
      [t for t in ["tenant", "session", "sqlalchemy", "query("] if t in _codigo_todo.lower()])
check("nenhuma função exige uf_empresa", "uf_empresa" not in _codigo_todo)


# ============================================================================
print("\n=== 3. Nada bloqueia: toda nota entra no arquivo ===")
mistura = [
    nota(num_nota="10", uf_prestador="ES", uf_tomador="ES"),
    nota(num_nota="11", uf_prestador="", uf_tomador="ES"),
    nota(num_nota="12", uf_prestador="SP", uf_tomador="ES"),
    nota(num_nota="13", uf_prestador="SP", uf_tomador=""),
]
linhas = [l for l in mod.gerar_txt_importacao(mistura).split("\n") if l]
check("as 4 notas entram no TXT (nenhuma some)", len(linhas) == 4, len(linhas))
check("nota 10 (ES/ES) -> 1933", ";1933;" in linhas[0], linhas[0])
check("nota 11 (sem UF prest.) -> 1933 presumido", ";1933;" in linhas[1], linhas[1])
check("nota 12 (SP/ES) -> 2933", ";2933;" in linhas[2], linhas[2])
check("nota 13 (sem UF tomador) -> 1933 presumido", ";1933;" in linhas[3], linhas[3])

prep = mod.preparar_notas(mistura)
check("só as notas 11 e 13 ficam marcadas como presumidas",
      [n["num_nota"] for n in prep if n["_presumido"]] == ["11", "13"],
      [n["num_nota"] for n in prep if n["_presumido"]])
check("a marca diz QUAL UF faltou (prestador)",
      "UF do prestador" in mod.motivo_presuncao(prep[1]), mod.motivo_presuncao(prep[1]))
check("a marca diz QUAL UF faltou (tomador)",
      "UF do tomador" in mod.motivo_presuncao(prep[3]), mod.motivo_presuncao(prep[3]))
check("nota lida não tem marca nenhuma", mod.motivo_presuncao(prep[0]) == "")

r = mod.resumo(mistura)
check("resumo conta 4 notas", r["total"] == 4, r)
check("resumo separa presumidas das internas de verdade",
      r["internas"] == 1 and r["presumidas"] == 2 and r["interestaduais"] == 1, r)
check("resumo soma o valor total", r["valor_total"] == 4000.0, r["valor_total"])
check("resumo lista as UFs de tomador encontradas", r["tomadores"] == ["ES"], r["tomadores"])


# ============================================================================
print("\n=== 4. Cadastro de fornecedores é opcional, nunca pré-requisito ===")
desconhecido = nota(num_nota="23", cnpj_prestador="55555555555555",
                    fornecedor="EMPRESA NUNCA VISTA LTDA", uf_prestador="MG",
                    uf_tomador="ES")
txt = mod.gerar_txt_importacao([desconhecido])
check("fornecedor fora do cadastro é lançado normalmente", txt.startswith("55555555555555;"))
check("MG x ES -> 2933", ";2933;" in txt)
check("com os marcadores de dado ausente", "CIDADE NAO INFORMADA" in txt)
check("e o código de serviço padrão 176", ";176;2933;" in txt, txt)

# Quando o fornecedor ESTÁ no cadastro, ele completa e confirma.
solvis = nota(num_nota="20", cnpj_prestador="00105063000102", uf_prestador="",
              fornecedor="SOLVIS", uf_tomador="ES")
txt = mod.gerar_txt_importacao([solvis])
check("cadastro resolve a UF que a IA não leu", ";2933;" in txt, txt)
check("e a nota deixa de ser presumida", mod.preparar_notas([solvis])[0]["_presumido"] is False)
check("usa o código de serviço do cadastro (179)", ";179;2933;" in txt, txt)
check("usa cidade e endereço do cadastro", "CURITIBA" in txt and "WESTPHALEN" in txt)

por_nome = nota(num_nota="21", cnpj_prestador="00000000000000",
                fornecedor="INDEXMED TECNOLOGIA LTDA", uf_prestador="", uf_tomador="ES")
txt = mod.gerar_txt_importacao([por_nome])
check("acha o fornecedor pela razão social", txt.startswith("34177449000141;"), txt[:40])
check("INDEXMED é de SP -> 2933", ";2933;" in txt)

uf_errada = nota(num_nota="22", cnpj_prestador="00105063000102",
                 uf_prestador="ES", uf_tomador="ES")
check("cadastro (PR) vence a UF lida do PDF (ES) -> 2933",
      ";2933;" in mod.gerar_txt_importacao([uf_errada]))

# Compatibilidade: análise salva antes da mudança usava `cnpj` em vez de
# `cnpj_prestador`. Não pode virar linha sem CNPJ no arquivo do Domínio.
antiga = {"cnpj": "12345678000199", "fornecedor": "FORMATO ANTIGO LTDA",
          "uf_prestador": "ES", "num_nota": "99", "valor": 100.0,
          "data_emissao": "01/07/2026"}
check("resultado no formato antigo ainda lança com o CNPJ certo",
      mod.gerar_txt_importacao([antiga]).startswith("12345678000199;"),
      mod.gerar_txt_importacao([antiga])[:30])


# ============================================================================
print("\n=== 5. Layout da linha do Domínio ===")
uma = nota(num_nota="30", cnpj_prestador="54090954000161", valor=1500.5,
           uf_prestador="ES", uf_tomador="ES",
           retencao_ir=22.51, retencao_pis=9.75, retencao_cofins=45.02, retencao_csll=15.01)
campos = mod.gerar_txt_importacao([uma]).split(";")
check("CNPJ no campo 1", campos[0] == "54090954000161", campos[0])
check("razão social do cadastro no campo 2", campos[1] == "CAFFEU E THOMAZ DISTRIBUIDORA", campos[1])
check("UF do prestador no campo 3", campos[2] == "ES", campos[2])
check("número da nota no campo 6", campos[5] == "30", campos[5])
check("emissão no campo 8", campos[7] == "05/07/2026", campos[7])
check("entrada no campo 9", campos[8] == "08/07/2026", campos[8])
check("código de serviço no campo 11", campos[10] == "176", campos[10])
check("CFOP no campo 12", campos[11] == "1933", campos[11])
check("valor com vírgula decimal no campo 13", campos[12] == "1500,50", campos[12])
check("retenções somam as QUATRO (22,51+9,75+45,02+15,01 = 92,29)",
      campos[24] == "92,29", campos[24])

sem_entrada = nota(num_nota="31", data_entrada="")
check("sem data de entrada, usa a de emissão",
      mod.gerar_txt_importacao([sem_entrada]).split(";")[8] == "05/07/2026")
check("valor inválido não quebra o arquivo",
      mod.formatar_moeda("abc") == "0,00" and mod.formatar_moeda(None) == "0,00")
check("total de retenções ignora campo sujo",
      mod.total_retencoes({"retencao_ir": "x", "retencao_pis": 5.0}) == 5.0)
check("lista vazia gera TXT vazio, sem estourar", mod.gerar_txt_importacao([]) == "")
check("None não estoura", mod.gerar_txt_importacao(None) == "")


# ============================================================================
print("\n=== 6. Planilha de conferência ===")
import openpyxl
validadas = [
    nota(num_nota="40", uf_prestador="RJ", uf_tomador="RJ",
         cnpj_prestador="66666666000191", fornecedor="CARIOCA LTDA"),
    nota(num_nota="41", uf_prestador="", uf_tomador="",
         cnpj_prestador="77777777777777", fornecedor="SEM UF LTDA"),
]
divs = [{"cnpj_prestador": "88888888888888", "num_nota": "42",
         "data_emissao": "01/07/2026",
         "motivo_divergencia": "Resumo diz 100, nota diz 120."}]

wb = openpyxl.load_workbook(mod.gerar_excel_auditoria(validadas, divs))
check("duas abas", wb.sheetnames == ["Notas Validadas", "Divergências"], wb.sheetnames)

ws = wb["Notas Validadas"]
cab = [c.value for c in ws[1]]
check("cabeçalho mostra os DOIS lados",
      "UF Prest." in cab and "UF Tom." in cab and "CNPJ Tomador" in cab, cab)
check("coluna de CFOP", cab[9] == "CFOP Calculado", cab[9])
check("coluna de observação", cab[10] == "Observação", cab[10])
check("nota RJ x RJ aparece como 1933", ws.cell(2, 10).value == "1933", ws.cell(2, 10).value)
check("nota lida não tem observação", not ws.cell(2, 11).value, ws.cell(2, 11).value)
check("nota sem UF sai 1933 MAS com o aviso escrito",
      ws.cell(3, 10).value == "1933" and "presumido" in (ws.cell(3, 11).value or ""),
      (ws.cell(3, 10).value, ws.cell(3, 11).value))
check("e a linha presumida fica destacada em amarelo",
      ws.cell(3, 1).fill.start_color.rgb.endswith("FEF3C7"),
      ws.cell(3, 1).fill.start_color.rgb)
check("retenções da planilha somam as quatro (61,50)", ws.cell(2, 9).value == 61.5,
      ws.cell(2, 9).value)

ws_div = wb["Divergências"]
check("o motivo da divergência chega na planilha (lia a chave errada antes)",
      ws_div.cell(2, 4).value == "Resumo diz 100, nota diz 120.", ws_div.cell(2, 4).value)

wb_vazio = openpyxl.load_workbook(mod.gerar_excel_auditoria([], []))
check("resultado vazio ainda cria as duas abas",
      wb_vazio.sheetnames == ["Notas Validadas", "Divergências"], wb_vazio.sheetnames)


# ============================================================================
fails = [n for n, ok, _ in results if not ok]
print("\n" + "=" * 60)
print(f"RESULTADO: {len(results)-len(fails)}/{len(results)} passaram")
if fails:
    print("FALHAS:")
    for f in fails:
        print(f"  - {f}")
    sys.exit(1)
print("OK")
