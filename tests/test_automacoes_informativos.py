# -*- coding: utf-8 -*-
"""Informativos por e-mail: resumo das auditorias e erros do sistema.

Pedido do Sandro em 25/09/2026: na mesma aba das Automações, além da coleta de
NFS-e, dois informativos configuráveis — o resumo das auditorias do período
(gestão, semanal ou mensal) e os erros do sistema (TI, diário e só quando
houver erro). O que esta suíte trava:

1. **A janela é honesta.** "Mês anterior" usa o mês de Brasília, não o de UTC;
   "desde a última rodada" começa onde a última BEM-SUCEDIDA terminou, então
   nada cai no vão entre dois envios; e o rótulo não diz "até 01/09" quando o
   dia 1º não entrou.

2. **Mensal existe e não pula mês.** Dia de 1 a 28 — automação marcada para o
   dia 31 não rodaria em abril, e ninguém seria avisado.

3. **O resumo não abre `result_data` de todo mundo, e nulo não vira zero.** Ele
   soma colunas; auditoria antiga tem o relatório lido UMA vez e as colunas
   preenchidas. Relatório ilegível sai como "sem valor apurado", nunca como
   R$ 0,00. A mesma competência rodada duas vezes conta uma vez, e o que não
   concluiu aparece em lista própria.

4. **O informe de erros usa as duas fontes da aba.** Relatórios de diagnóstico
   E o log cru, agrupado — sem o log, "nenhum erro" viria justamente no dia em
   que o diagnóstico estava desligado. Texto de exceção sai escapado.

5. **O agendador trava por TIPO.** Uma coleta de uma hora não segura o resumo;
   dois resumos não rodam juntos. Sem novidade, o e-mail é "pulado" e isso fica
   registrado; o "Rodar agora" envia sempre.

6. **A tela** oferece os três tipos, valida cada um e não deixa o tipo de uma
   automação existente ser trocado por um campo escondido.
"""
import io
import json
import os
import shutil
import sys
import tempfile
import threading
import types
import uuid
from datetime import datetime, timedelta

sys.path.insert(0, os.path.abspath('.'))

import dotenv
dotenv.load_dotenv = lambda *a, **k: False
import dotenv.main
dotenv.main.load_dotenv = lambda *a, **k: False

os.environ["DATABASE_URL"] = "sqlite:///./_informativos_test.db"
os.environ["SUPABASE_URL"] = "http://fake"
os.environ["SUPABASE_KEY"] = "fake"
os.environ["SUPABASE_SERVICE_KEY"] = "fake"
os.environ["FERNET_SECRET_KEY"] = "hn5Ck1J8pVXHRV9V3Q6mYbXW4y2S1TqLpRZ0oN7dG8A="
_TMP = tempfile.mkdtemp(prefix="informativos_test_")
os.environ["NFSE_STORAGE_DIR"] = _TMP
if os.path.exists("_informativos_test.db"):
    os.remove("_informativos_test.db")

from sqlalchemy import event, types as _satypes  # noqa: E402
import sqlalchemy.dialects.postgresql as pg  # noqa: E402


class _UUIDShim(_satypes.TypeDecorator):
    impl = _satypes.Uuid
    cache_ok = True

    def __init__(self, *a, **kw):
        kw.pop('as_uuid', None)
        super().__init__()

    def process_bind_param(self, value, dialect):
        return uuid.UUID(value) if isinstance(value, str) else value


pg.UUID = _UUIDShim

import backend.app.infrastructure.auth.supabase_client as sb  # noqa: E402
TOKENS = {}


class _Auth:
    def get_user(self, token):
        u = TOKENS.get(token.replace("Bearer ", "").strip())
        if not u:
            raise Exception("invalid")
        return types.SimpleNamespace(user=u)


class _Client:
    auth = _Auth()


sb.get_supabase_client = lambda: _Client()
sb.get_supabase_admin_client = lambda: _Client()
import backend.app.api.dependencies.auth as dep_auth  # noqa: E402
dep_auth.get_supabase_client = lambda: _Client()

from backend.app.infrastructure.database.session import engine as _eng  # noqa: E402
if _eng.url.get_backend_name() != "sqlite":
    raise SystemExit("ABORTADO: engine nao e SQLite")

import backend.app.main as m  # noqa: E402
from backend.app.infrastructure.database.session import SessionLocal  # noqa: E402
from backend.app.domain.models.base import Base  # noqa: E402
from backend.app.domain.models.audit import AuditTask  # noqa: E402
from backend.app.domain.models.auth_models import Tenant, UserProfile, UserTenant  # noqa: E402
from backend.app.domain.models.automacao import NfseAutomacao, NfseAutomacaoExecucao  # noqa: E402
from backend.app.core.fuso import FUSO_LOCAL_HORAS  # noqa: E402
from backend.app.services import configuracoes as cfg  # noqa: E402
from backend.app.services import email_envio, erros_sistema  # noqa: E402
from backend.app.services.audit_engine.resumo_financeiro import valores_para_colunas  # noqa: E402
from backend.app.services.automacoes import informe_erros as ie  # noqa: E402
from backend.app.services.automacoes import janela as jan  # noqa: E402
from backend.app.services.automacoes import resumo_auditorias as ra  # noqa: E402
from backend.app.services.automacoes import tipos  # noqa: E402
from backend.app.services.nfse import automacao as svc  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
import openpyxl  # noqa: E402

OK = FALHAS = 0
_falhou = []


def check(nome, condicao, detalhe=""):
    global OK, FALHAS
    if condicao:
        OK += 1
        print(f"  OK   | {nome}")
    else:
        FALHAS += 1
        _falhou.append(nome)
        print(f" FALHA | {nome}  [{detalhe}]")


# ============================================================================
print("\n=== 1. A janela: mes de Brasilia, e nada cai no vao entre dois envios ===")

agora = datetime(2026, 9, 25, 15, 0)          # 12:00 em Brasilia
check("7 dias", jan.janela("7d", agora) == (agora - timedelta(days=7), agora))
check("24 horas", jan.janela("24h", agora) == (agora - timedelta(days=1), agora))
ini, fim = jan.janela("mes_anterior", agora)
check("mes anterior = agosto inteiro, nas fronteiras de Brasilia (03:00 UTC)",
      (ini, fim) == (datetime(2026, 8, 1, 3, 0), datetime(2026, 9, 1, 3, 0)), (ini, fim))
check("o rotulo do mes fechado nao parece incluir o dia 1o",
      jan.rotulo(ini, fim) == "01/08/2026 00:00 a 31/08/2026 23:59", jan.rotulo(ini, fim))
check("nem o curto, que vai no assunto do e-mail",
      jan.rotulo_curto(ini, fim) == "01/08–31/08", jan.rotulo_curto(ini, fim))
ini2, fim2 = jan.janela("mes_anterior", datetime(2026, 9, 1, 2, 0))
check("02h UTC do dia 1o ainda e dia 31 em Brasilia: o mes anterior e julho",
      (ini2, fim2) == (datetime(2026, 7, 1, 3, 0), datetime(2026, 8, 1, 3, 0)), (ini2, fim2))
check("mes atual ate agora",
      jan.janela("mes_atual", agora) == (datetime(2026, 9, 1, 3, 0), agora))
ult = agora - timedelta(hours=30)
check("desde a ultima: comeca onde a ultima terminou",
      jan.janela("desde_ultima", agora, ult) == (ult, agora))
check("primeira vez: um dia na diaria, sete na semanal, trinta na mensal",
      [jan.janela("desde_ultima", agora, None, f)[0] for f in ("diaria", "semanal", "mensal")]
      == [agora - timedelta(days=d) for d in (1, 7, 30)])
check("valor desconhecido cai no 'desde a ultima' (seguro)",
      jan.janela("inventado", agora, ult) == (ult, agora))
check("fim anterior no futuro (relogio mexido) nao inverte a janela",
      jan.janela("desde_ultima", agora, agora + timedelta(hours=1))[0] == agora - timedelta(days=1))


# ============================================================================
print("\n=== 2. Mensal: dia de 1 a 28, e a virada do ano ===")

p = svc.proxima_execucao
check("dia 1 as 08:00, a partir de 25/09: 1o/10 as 11:00 UTC",
      p("mensal", "08:00", [], agora, 1) == datetime(2026, 10, 1, 11, 0), p("mensal", "08:00", [], agora, 1))
check("no proprio dia, antes da hora: hoje",
      p("mensal", "08:00", [], datetime(2026, 10, 1, 10, 59), 1) == datetime(2026, 10, 1, 11, 0))
check("exatamente na hora: o mes seguinte (estritamente depois)",
      p("mensal", "08:00", [], datetime(2026, 10, 1, 11, 0), 1) == datetime(2026, 11, 1, 11, 0))
check("dezembro vira janeiro do ano seguinte",
      p("mensal", "08:00", [], datetime(2026, 12, 15, 12, 0), 1) == datetime(2027, 1, 1, 11, 0))
check("dia 28 existe em fevereiro",
      p("mensal", "08:00", [], datetime(2027, 2, 1, 12, 0), 28) == datetime(2027, 2, 28, 11, 0))
check("sem dia, 0, 29, 31 ou lixo: nao inventa data",
      all(p("mensal", "08:00", [], agora, d) is None for d in (None, 0, 29, 31, "x")))
check("o dia do mes nao mexe na diaria",
      p("diaria", "07:00", [], datetime(2026, 9, 24, 9, 0), 15) == datetime(2026, 9, 24, 10, 0))
check("descricao legivel da mensal",
      svc.descrever_agenda("mensal", "08:00", None, 1) == "Todo dia 1º do mês às 08:00"
      and svc.descrever_agenda("mensal", "18:30", None, 15) == "Todo dia 15 do mês às 18:30",
      svc.descrever_agenda("mensal", "08:00", None, 1))
check("proxima_da le o dia gravado na automacao",
      svc.proxima_da(NfseAutomacao(frequencia="mensal", hora="08:00", dia_mes=5), agora)
      == datetime(2026, 10, 5, 11, 0))


# ============================================================================
print("\n=== 3. Os tipos e o que cada um aceita ===")

# O quarto tipo, "resumo_emails", nasceu em 28/09/2026 e é travado na suíte
# própria (tests/test_email_recebimento.py). Aqui a conta segue completa: o
# que esta lista protege é ninguém sumir com um tipo sem perceber. O quinto, "validar_sped" (23/10/2026), é
# travado em tests/test_validar_sped_pva.py.
check("cinco tipos", set(tipos.TIPOS) == {"coleta_nfse", "resumo_auditorias",
                                         "erros_sistema", "resumo_emails", "validar_sped"},
      sorted(tipos.TIPOS))
check("coleta nao tem mensal; os informativos tem",
      "mensal" not in tipos.frequencias("coleta_nfse")
      and "mensal" in tipos.frequencias("resumo_auditorias")
      and "mensal" in tipos.frequencias("erros_sistema"))
velha = types.SimpleNamespace(tipo="resumo_auditorias", config='{"periodo": "30d", "invasor": 1}')
conf = tipos.ler_config(velha)
check("config gravada e completada pelo padrao, sem chave estranha",
      conf == {"periodo": "30d", "ultima_por_competencia": True, "so_com_novidade": False}, conf)
check("JSON quebrado vira o padrao do tipo",
      tipos.ler_config(types.SimpleNamespace(tipo="erros_sistema", config="{")) == tipos.padrao("erros_sistema"))
_c, _e = tipos.config_do_formulario("resumo_auditorias", {"periodo": "semana-passada"})
check("periodo fora da lista e recusado", "período" in _e, _e)
_c, _e = tipos.config_do_formulario("erros_sistema", {"periodo": "24h", "gravidade": "alta",
                                                     "incluir_log": "", "so_com_novidade": "1"})
check("formulario vira config", not _e and _c == {"periodo": "24h", "gravidade": "alta",
                                                   "incluir_log": False, "so_com_novidade": True}, _c)
_d = tipos.descrever("erros_sistema", _c)
check("e o cartao a descreve", "Últimas 24 horas" in _d and "sem o log cru" in _d
      and "só quando houver erro" in _d, _d)


# ============================================================================
print("\n=== 4. Resumo das auditorias: colunas, nunca o relatorio de todo mundo ===")

Base.metadata.create_all(bind=_eng)
db = SessionLocal()
AGORA = datetime.utcnow()
T1, T2, T3 = (str(uuid.uuid4()) for _ in range(3))
ANA, ADM, OPE = (str(uuid.uuid4()) for _ in range(3))
db.add_all([
    Tenant(id=T1, name="ALFA <b>COMERCIO</b>", cnpj="11111111000111"),
    Tenant(id=T2, name="BETA DISTRIBUIDORA", cnpj="22222222000122"),
    Tenant(id=T3, name="GAMA TRANSPORTES", cnpj="33333333000133"),
    UserProfile(id=ANA, tenant_id=T1, role="operador", name="Ana Fiscal"),
    UserProfile(id=ADM, tenant_id=T1, role="admin", name="Admin"),
    UserProfile(id=OPE, tenant_id=T1, role="operador", name="Operador"),
])
db.add(UserTenant(user_id=OPE, tenant_id=T1))
db.commit()

# Relatorio de auditoria ANTIGA (colunas nulas): 50 a estornar, 30 a
# aproveitar, 20 a recolher e 500 a conferir.
RELATORIO_ANTIGO = json.dumps({
    "resumo": {"notas_processadas": 10, "inconsistencias": 3},
    "periodo_inicio": "2026-08-01",
    "icms_issues": [{"sped_icms": 150.0, "xml_icms": 100.0}],
    "ipi_issues": [{"sped_val": 0.0, "xml_val": 30.0}],
    "icms_st_notes": [{"vl_icms_st": 20.0}],
    "missing_xmls": [{"vl_doc": 500.0}],
})
check("(a conta das colunas e a da Visao Geral)",
      valores_para_colunas(json.loads(RELATORIO_ANTIGO))
      == {"vl_estornar": 50.0, "vl_aproveitar": 30.0, "vl_recolher": 20.0, "vl_conferir": 500.0})

TAREFAS = {}


def tarefa(chave, tenant, status, horas_atras, competencia=None, notas=None, incons=None,
           valores=None, relatorio=None):
    t = AuditTask(id=str(uuid.uuid4()), tenant_id=tenant, user_id=ANA, status=status,
                  file_name=f"{chave}.txt", competencia=competencia, notas_processadas=notas,
                  inconsistencias=incons, result_data=relatorio)
    for coluna, valor in (valores or {}).items():
        setattr(t, coluna, valor)
    t.created_at = AGORA - timedelta(hours=horas_atras)
    t.updated_at = t.created_at
    db.add(t)
    TAREFAS[chave] = t.id


def vl(e, a, r, c=0.0):
    return {"vl_estornar": e, "vl_aproveitar": a, "vl_recolher": r, "vl_conferir": c}


tarefa("antiga", T1, "concluido", 72, relatorio=RELATORIO_ANTIGO)
# A mesma empresa e competencia, rodada de novo depois: vale esta.
tarefa("refeita", T1, "concluido", 24, "2026-08", 10, 1, vl(10.0, 0.0, 0.0), '{"lixo": 1}')
tarefa("beta", T2, "concluido", 48, "2026-07", 40, 7, vl(0.0, 200.0, 5.5))
tarefa("ilegivel", T3, "concluido", 30, "2026-05", 5, 2, relatorio="{isso nao e json")
tarefa("falhou", T3, "erro", 20)
tarefa("parada", T3, "aguardando_ia", 10, "2026-08")
tarefa("velha", T2, "concluido", 24 * 20, "2026-06", 1, 1, vl(999.0, 0.0, 0.0))
db.commit()

SQL = []


@event.listens_for(_eng, "before_cursor_execute")
def _capturar(conn, cursor, statement, params, contexto, executemany):
    SQL.append(statement)


config_padrao = tipos.padrao("resumo_auditorias")
automacao_todas = types.SimpleNamespace(escopo="todas", empresas=None)
ini7, fim7 = jan.janela("7d", datetime.utcnow())
SQL.clear()
dados = ra.coletar(db, automacao_todas, config_padrao, ini7, fim7)
lendo_relatorio = [s for s in SQL if "result_data" in s]
check("le o result_data SO das duas auditorias sem coluna (uma a uma, por id)",
      len(lendo_relatorio) == 2 and all("audit_tasks.id =" in s for s in lendo_relatorio),
      lendo_relatorio)

db.expire_all()
antiga = db.get(AuditTask, TAREFAS["antiga"])
check("a auditoria antiga ficou com as colunas preenchidas (da proxima vez, e coluna)",
      (antiga.vl_estornar, antiga.vl_aproveitar, antiga.vl_recolher, antiga.vl_conferir)
      == (50.0, 30.0, 20.0, 500.0),
      (antiga.vl_estornar, antiga.vl_aproveitar, antiga.vl_recolher, antiga.vl_conferir))
check("e com os contadores e a competencia que faltavam",
      (antiga.notas_processadas, antiga.inconsistencias, antiga.competencia) == (10, 3, "2026-08"))
ilegivel = db.get(AuditTask, TAREFAS["ilegivel"])
check("relatorio ilegivel: a coluna CONTINUA nula (nao vira zero)", ilegivel.vl_estornar is None)

t = dados["totais"]
check("a mesma competencia rodada duas vezes conta UMA (a mais recente)",
      t["auditorias"] == 3 and t["repetidas_descartadas"] == 1
      and "refeita.txt" in [a["arquivo"] for a in dados["auditorias"]]
      and "antiga.txt" not in [a["arquivo"] for a in dados["auditorias"]], t)
check("fora da janela nao entra", "velha.txt" not in [a["arquivo"] for a in dados["auditorias"]])
check("totais: 215,50 em jogo (10 + 200 + 5,50)",
      (t["estornar"], t["aproveitar"], t["recolher"], t["em_jogo"]) == (10.0, 200.0, 5.5, 215.5), t)
check("notas e inconsistencias somadas", (t["notas"], t["inconsistencias"], t["empresas"]) == (55, 10, 3), t)
check("maior valor em jogo primeiro",
      [a["empresa"] for a in dados["auditorias"]][0] == "BETA DISTRIBUIDORA")
sem = [a for a in dados["auditorias"] if a["arquivo"] == "ilegivel.txt"][0]
check("a ilegivel entra marcada 'sem valor apurado'", sem["sem_valor"] and t["sem_valor"] == 1, sem)
situacoes = sorted(a["situacao"] for a in dados["nao_concluidas"])
check("o que nao concluiu aparece, com a situacao em portugues",
      situacoes == ["Falhou", "Parada esperando a confirmação da IA"], situacoes)
check("desfecho 'ok' com o dinheiro na mensagem",
      ra.desfecho(dados)[0] == "ok" and "R$ 215,50 em jogo" in ra.desfecho(dados)[1],
      ra.desfecho(dados))

SQL.clear()
ra.coletar(db, automacao_todas, config_padrao, ini7, fim7)
check("na segunda vez so a ilegivel e relida", len([s for s in SQL if "result_data" in s]) == 1)
event.remove(_eng, "before_cursor_execute", _capturar)

todas = ra.coletar(db, automacao_todas, {**config_padrao, "ultima_por_competencia": False}, ini7, fim7)
check("sem 'ultima por competencia': as repetidas somam",
      todas["totais"]["auditorias"] == 4 and todas["totais"]["em_jogo"] == 315.5, todas["totais"])
so_beta = ra.coletar(db, types.SimpleNamespace(escopo="selecionadas", empresas=T2),
                     config_padrao, ini7, fim7)
check("escopo: so as empresas escolhidas",
      [a["empresa"] for a in so_beta["auditorias"]] == ["BETA DISTRIBUIDORA"]
      and so_beta["nao_concluidas"] == [], so_beta["totais"])
vazio = ra.coletar(db, automacao_todas, config_padrao, AGORA - timedelta(days=100),
                   AGORA - timedelta(days=99))
check("periodo sem nada: 'sem novidade', nunca 'ok' com zero",
      ra.desfecho(vazio)[:1] == ("sem_novidade",) and ra.desfecho(vazio)[2] is False)

wb = openpyxl.load_workbook(io.BytesIO(ra.gerar_planilha(dados, "Resumo semanal")))
check("planilha: capa, auditorias e nao concluidas",
      wb.sheetnames == ["Visão Geral", "Auditorias", "Não concluídas"], wb.sheetnames)
linhas = list(wb["Auditorias"].iter_rows(min_row=2, values_only=True))
gama = [r for r in linhas if r[0] == "GAMA TRANSPORTES"][0]
check("sem valor apurado: celulas de valor VAZIAS, com a observacao",
      gama[7:12] == (None,) * 5 and "ilegível" in (gama[13] or ""), gama)
check("e as outras com o valor", [r for r in linhas if r[0] == "BETA DISTRIBUIDORA"][0][10] == 205.5)

assunto, html_mail, texto_mail = ra.montar_email(dados, "Resumo semanal", com_anexo=True)
check("assunto com o dinheiro e o que nao concluiu",
      assunto.startswith("[Synax] Resumo das auditorias") and "R$ 215,50 em jogo" in assunto
      and "2 não concluídas" in assunto, assunto)
check("nome de empresa com HTML sai escapado",
      "&lt;b&gt;COMERCIO&lt;/b&gt;" in html_mail and "<b>COMERCIO</b>" not in html_mail)
check("o e-mail traz o valor, as nao concluidas e a linha sem valor",
      all(s in html_mail for s in ("215,50", "Parada esperando a confirmação da IA",
                                   "sem valor apurado", "A planilha completa segue em anexo")))
check("e o texto puro tambem", "Em jogo: R$ 215,50" in texto_mail and "valor não apurado" in texto_mail)
muitas = dict(dados, auditorias=[dict(dados["auditorias"][0], empresa=f"EMPRESA {i}") for i in range(30)])
check("lista longa: o e-mail para no limite e diz quantas ficaram na planilha",
      "e mais 5 auditorias na planilha" in ra.montar_email(muitas, "x", com_anexo=True)[1])

CC = io.open(os.path.join("backend", "app", "services", "audit_engine", "cross_checker.py"),
             encoding="utf-8").read()
check("o motor grava o valor em jogo ao concluir, sem derrubar a auditoria se falhar",
      "valores_para_colunas(results)" in CC
      and "Não foi possível calcular o valor em jogo" in CC)


# ============================================================================
print("\n=== 5. Erros do sistema: relatorios E log, agrupado e escapado ===")

PASTA = os.path.join(_TMP, "erros_ia")
os.makedirs(PASTA)
LOG = os.path.join(_TMP, "erros_ia.log")
_dir_real, _log_real = erros_sistema._diretorio, erros_sistema._arquivo_log
erros_sistema._diretorio = lambda: PASTA
erros_sistema._arquivo_log = lambda: LOG
LOCAL = datetime.utcnow() + timedelta(hours=FUSO_LOCAL_HORAS)   # o disco fala Brasilia


def relatorio(horas_atras, origem, tipo_, gravidade, resumo):
    quando = LOCAL - timedelta(hours=horas_atras)
    corpo = f"# Erro\n\n## Em uma frase\n{resumo}\n"
    if gravidade:
        corpo += f"\n## Gravidade\n**{gravidade}** — explicação\n"
    io.open(os.path.join(PASTA, f"{quando:%Y-%m-%d_%H-%M-%S}__{origem}__{tipo_}.md"), "w",
            encoding="utf-8").write(corpo)


relatorio(2, "envio-de-arquivos", "HTTPException", "Alta", "Upload falhou <script>alert(1)</script>")
relatorio(3, "auditoria", "KeyError", "Média", "Chave ausente no relatório")
relatorio(4, "curadoria", "ValueError", "Baixa", "Valor inválido")
relatorio(5, "radar", "TypeError", "", "Sem gravidade escrita")
relatorio(72, "antigo", "OSError", "Alta", "Fora da janela")

_linhas_log = []


def no_log(horas_atras, origem, mensagem, nivel="ERROR"):
    quando = LOCAL - timedelta(hours=horas_atras)
    _linhas_log.append(f"{quando:%Y-%m-%d %H:%M:%S},123 {nivel} {origem} — {mensagem}")


for _h in (1, 2, 3):
    no_log(_h, "backend.app.services.sieg", "Timeout ao consultar o cofre")
no_log(2.5, "backend.app.api.routes.uploads", "Falha ao ler <b>XML</b> do fornecedor")
no_log(1.5, "backend.app.x", "so um aviso", nivel="WARNING")
no_log(48, "backend.app.velho", "fora da janela")
io.open(LOG, "w", encoding="utf-8").write("\n".join(_linhas_log) + "\n")

check("gravidade lida do texto do diagnostico",
      (ie.nivel("Alta"), ie.nivel("Média"), ie.nivel("media — algo"), ie.nivel("Crítica"),
       ie.nivel("Baixa"), ie.nivel("")) == (3, 2, 2, 3, 1, 0))

ini24, fim24 = jan.janela("24h", datetime.utcnow())
base_erros = tipos.padrao("erros_sistema")
er = ie.coletar(base_erros, ini24, fim24)
check("os quatro relatorios do dia, nao o de tres dias atras",
      er["totais"]["relatorios"] == 4 and er["totais"]["altas"] == 1, er["totais"])
check("o log so com ERROR/CRITICAL da janela: 4 ocorrencias em 2 tipos",
      (er["totais"]["ocorrencias_log"], er["totais"]["tipos_log"]) == (4, 2), er["totais"])
check("a mesma falha 3 vezes e UMA linha com a contagem, e vem primeiro",
      er["log"][0]["vezes"] == 3 and er["log"][0]["origem"] == "backend.app.services.sieg", er["log"][:1])
check("gravidade 'media e alta' corta baixa e sem gravidade",
      ie.coletar({**base_erros, "gravidade": "media"}, ini24, fim24)["totais"]["relatorios"] == 2)
check("gravidade 'so alta'",
      ie.coletar({**base_erros, "gravidade": "alta"}, ini24, fim24)["totais"]["relatorios"] == 1)
sem_log = ie.coletar({**base_erros, "incluir_log": False}, ini24, fim24)
check("sem o log cru, so os relatorios", sem_log["log"] == [] and sem_log["totais"]["ocorrencias_log"] == 0)
check("desfecho com as duas fontes",
      ie.desfecho(er) == ("ok", "4 erros diagnosticados (1 de gravidade alta) · 4 ocorrências no log (2 tipos)", True),
      ie.desfecho(er))
nada = ie.coletar(base_erros, AGORA - timedelta(days=100), AGORA - timedelta(days=99))
check("periodo sem erro: 'sem novidade'", ie.desfecho(nada) == ("sem_novidade", "Nenhum erro no período.", False))
assunto_e, html_e, texto_e = ie.montar_email(er, "Erros TI")
check("assunto de alerta quando ha erro", assunto_e.startswith("⚠ [Synax] Erros do sistema")
      and "4 diagnosticados · 4 no log" in assunto_e, assunto_e)
check("texto de excecao sai escapado no e-mail",
      "&lt;script&gt;" in html_e and "<script>alert" not in html_e and "&lt;b&gt;XML&lt;/b&gt;" in html_e)
check("o texto puro traz a contagem", "3x backend.app.services.sieg" in texto_e, texto_e)
check("sem erro, o assunto nao assusta", not ie.assunto(nada).startswith("⚠")
      and "nenhum erro no período" in ie.assunto(nada))


# ============================================================================
print("\n=== 6. O agendador: trava por tipo, 'pulado' registrado, manual envia ===")

ENVIADOS = []


class FakeSMTP:
    def __init__(self, host, porta, timeout=None, **kw):
        pass

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def ehlo(self):
        pass

    def starttls(self, context=None):
        pass

    def login(self, usuario, senha):
        pass

    def send_message(self, msg):
        ENVIADOS.append(msg)
        return {}


_smtp_real = email_envio.smtplib.SMTP
email_envio.smtplib.SMTP = FakeSMTP
for _chave, _valor in ((cfg.SMTP_HOST, "MCO2"), (cfg.SMTP_PORTA, "25"),
                       (cfg.SMTP_SEGURANCA, "nenhuma"),
                       (cfg.SMTP_REMETENTE, "synax@vmcontadores.com.br")):
    cfg.gravar(db, _chave, _valor, "teste")

ID_RES, ID_ERR = str(uuid.uuid4()), str(uuid.uuid4())
db.add_all([
    NfseAutomacao(id=ID_RES, tipo="resumo_auditorias", nome="Resumo semanal", frequencia="semanal",
                  dias_semana="0", hora="08:00", escopo="todas", ativo=True,
                  destinatarios="gestao@vm.com.br", config=json.dumps(tipos.padrao("resumo_auditorias"))),
    NfseAutomacao(id=ID_ERR, tipo="erros_sistema", nome="Erros TI", frequencia="diaria",
                  hora="08:00", escopo="todas", ativo=True, destinatarios="ti@vm.com.br",
                  config=json.dumps(tipos.padrao("erros_sistema"))),
])
db.commit()


def rodar(aid, gatilho="agenda"):
    eid_ = svc.executar(SessionLocal, aid, gatilho)
    db.expire_all()
    return db.get(NfseAutomacaoExecucao, eid_)


e_res = rodar(ID_RES)
check("rodada do resumo: tipo e status", e_res.tipo == "resumo_auditorias" and e_res.status == "ok",
      (e_res.tipo, e_res.status, e_res.mensagem))
check("a janela fica gravada (7 dias)", e_res.periodo_fim - e_res.periodo_inicio == timedelta(days=7))
check("o detalhe guarda os numeros do e-mail",
      json.loads(e_res.detalhe)["totais"]["em_jogo"] == 215.5 and e_res.empresas_total == 3)
check("e-mail enviado, com a planilha", e_res.email_status == "enviado" and len(ENVIADOS) == 1,
      (e_res.email_status, e_res.email_detalhe))
_anexos = list(ENVIADOS[-1].iter_attachments())
check("anexo e xlsx de verdade", len(_anexos) == 1 and _anexos[0].get_filename().startswith("resumo_auditorias_")
      and _anexos[0].get_content()[:2] == b"PK", [a.get_filename() for a in _anexos])

ENVIADOS.clear()
e1 = rodar(ID_ERR)
check("informe de erros: sai com o que achou", e1.status == "ok" and e1.email_status == "enviado"
      and len(ENVIADOS) == 1, (e1.status, e1.email_status))
e2 = rodar(ID_ERR)
check("'desde a ultima': a segunda comeca onde a primeira terminou",
      e2.periodo_inicio == e1.periodo_fim, (e2.periodo_inicio, e1.periodo_fim))
check("sem erro novo: 'Nada a informar' e o e-mail PULADO, registrado",
      e2.status == "sem_novidade" and e2.email_status == "pulado" and len(ENVIADOS) == 1
      and "só quando há novidade" in (e2.email_detalhe or ""), (e2.status, e2.email_status))
e3 = rodar(ID_ERR, "manual")
check("o 'Rodar agora' envia mesmo sem novidade",
      e3.status == "sem_novidade" and e3.email_status == "enviado" and len(ENVIADOS) == 2
      and "nenhum erro no período" in ENVIADOS[-1]["Subject"], (e3.email_status, len(ENVIADOS)))

_coletar_real = ie.coletar


def _explode(*a, **k):
    raise RuntimeError("disco sumiu")


ie.coletar = _explode
e4 = rodar(ID_ERR)
ie.coletar = _coletar_real
check("falha inesperada: 'falhou', com mensagem de informativo (nao de coleta)",
      e4.status == "falhou" and "informativo" in (e4.mensagem or "") and "baixado" not in e4.mensagem,
      (e4.status, e4.mensagem))
e5 = rodar(ID_ERR)
check("a rodada que falhou nao conta: a proxima cobre desde a ultima BEM-SUCEDIDA",
      e5.periodo_inicio == e3.periodo_fim, (e5.periodo_inicio, e3.periodo_fim))

_trava_res = svc._trava("resumo_auditorias")
_trava_res.acquire()
try:
    check("resumo em andamento: outro resumo, sem esperar, nao roda",
          svc.executar(SessionLocal, ID_RES, "manual", esperar=False) is None)
    _ok, _txt = svc.rodar_agora(SessionLocal, ID_RES)
    check("e o botao diz que ha informativo em andamento", not _ok and "informativo deste tipo" in _txt, _txt)
    check("mas o informe de erros roda: a trava e por tipo",
          svc.executar(SessionLocal, ID_ERR, "manual", esperar=False) is not None)
finally:
    _trava_res.release()
svc._executando.acquire()
try:
    check("uma coleta de NFS-e andando nao segura o resumo",
          svc.executar(SessionLocal, ID_RES, "manual", esperar=False) is not None)
finally:
    svc._executando.release()

for _aid in (ID_RES, ID_ERR):
    db.get(NfseAutomacao, _aid).proxima_execucao = datetime.utcnow() - timedelta(minutes=5)
db.commit()
_antes = db.query(NfseAutomacaoExecucao).count()
_n = svc.rodar_devidas(SessionLocal, datetime.utcnow(), em_paralelo=True)
for _t in list(threading.enumerate()):
    if _t.name.startswith("automacao-"):
        _t.join(timeout=60)
db.expire_all()
check("o agendador dispara as duas, cada uma na sua thread",
      _n == 2 and db.query(NfseAutomacaoExecucao).count() - _antes == 2,
      (_n, db.query(NfseAutomacaoExecucao).count() - _antes))
check("e a proxima de cada uma foi para o futuro",
      all(db.get(NfseAutomacao, _aid).proxima_execucao > datetime.utcnow() for _aid in (ID_RES, ID_ERR)))

db.add(NfseAutomacaoExecucao(id=str(uuid.uuid4()), automacao_id=ID_RES, tipo="resumo_auditorias",
                             gatilho="agenda", status="rodando", inicio=datetime.utcnow()))
db.commit()
svc.marcar_interrompidas(db)
_presa = (db.query(NfseAutomacaoExecucao).filter_by(automacao_id=ID_RES, status="interrompida").first())
check("informativo interrompido pelo reinicio: a mensagem fala de informativo",
      _presa is not None and "informativo" in _presa.mensagem and "baixado" not in _presa.mensagem,
      _presa.mensagem if _presa else None)


# ============================================================================
print("\n=== 7. A tela: tres tipos, validacao de cada um ===")

TOKENS["adm"] = types.SimpleNamespace(id=ADM, email="admin@synax.com")
TOKENS["ope"] = types.SimpleNamespace(id=OPE, email="oper@synax.com")

from backend.app.api.dependencies.db import get_rls_db as _real_rls  # noqa: E402


def _fake_rls():
    ss = SessionLocal()
    try:
        yield ss
    finally:
        ss.close()


m.app.dependency_overrides[_real_rls] = _fake_rls
c = TestClient(m.app, raise_server_exceptions=False)


def as_(token):
    c.cookies.clear()
    if token:
        c.cookies.set("access_token", f"Bearer {token}")
    return c


r = as_("adm").get("/admin/automacoes")
check("o menu 'Nova automacao' oferece todos os tipos",
      r.status_code == 200 and all(f"/admin/automacoes/nova?tipo={k}" in r.text for k in tipos.TIPOS))
check("o cartao do resumo diz o tipo, a agenda e a configuracao",
      all(s in r.text for s in ("Resumo das auditorias", "Seg às 08:00", "Últimos 7 dias", "Próximo envio")))
check("o historico mostra o que o informativo levantou", "R$ 215,50 em jogo" in r.text)
check("rodada sem novidade: 'Nada a informar', com o e-mail pulado marcado",
      "Nada a informar" in r.text and "fa-bell-slash" in r.text)

r = as_("adm").get("/admin/automacoes/nova?tipo=resumo_auditorias")
check("modal do resumo: periodo, 'Todo mes' e o tipo escondido",
      all(s in r.text for s in ("Período coberto", 'value="mensal"',
                                'name="tipo" value="resumo_auditorias"',
                                "Resumo semanal das auditorias", 'name="ultima_por_competencia"')))
r = as_("adm").get("/admin/automacoes/nova?tipo=erros_sistema")
check("modal de erros: gravidade e log, sem escolha de empresa",
      "Gravidade mínima" in r.text and 'name="incluir_log"' in r.text
      and 'type="radio" name="escopo"' not in r.text and 'name="empresas"' not in r.text)
r = as_("adm").get("/admin/automacoes/nova?tipo=inventado")
check("tipo desconhecido vira coleta, sem 'Todo mes'",
      'name="tipo" value="coleta_nfse"' in r.text and 'value="mensal"' not in r.text)
check("sem tipo: a coleta de sempre",
      "Nova automação de coleta" in as_("adm").get("/admin/automacoes/nova").text)

BASE_RES = {"tipo": "resumo_auditorias", "nome": "Resumo mensal", "frequencia": "mensal",
            "hora": "08:00", "escopo": "todas", "ativo": "1", "periodo": "mes_anterior",
            "ultima_por_competencia": "1"}


def salvar(dados_):
    return as_("adm").post("/admin/automacoes/salvar", data=dados_)


r = salvar({**BASE_RES, "dia_mes": "1"})
check("informativo sem destinatario e recusado, no proprio modal",
      "Informe ao menos um destinatário" in r.text and r.headers.get("HX-Retarget") == "#modal-container")
r = salvar({**BASE_RES, "destinatarios": "gestao@vm.com.br"})
check("mensal sem dia: pede o dia", "dia do mês" in r.text)
r = salvar({**BASE_RES, "destinatarios": "gestao@vm.com.br", "dia_mes": "31"})
check("dia 31 e recusado (nao existe em todo mes)", "de 1 a 28" in r.text)
r = salvar({**BASE_RES, "destinatarios": "gestao@vm.com.br", "dia_mes": "1", "periodo": "semana-passada"})
check("periodo forjado e recusado", "Escolha o período" in r.text)
r = salvar({**BASE_RES, "destinatarios": "gestao@vm.com.br", "dia_mes": "1"})
nova = db.query(NfseAutomacao).filter_by(nome="Resumo mensal").first()
check("grava: tipo, dia do mes e configuracao",
      nova is not None and nova.tipo == "resumo_auditorias" and nova.dia_mes == 1
      and json.loads(nova.config) == {"periodo": "mes_anterior", "ultima_por_competencia": True,
                                      "so_com_novidade": False},
      (nova.tipo, nova.dia_mes, nova.config) if nova else None)
_local = nova.proxima_execucao + timedelta(hours=FUSO_LOCAL_HORAS)
check("a proxima e num dia 1o, as 08:00 de Brasilia",
      _local.day == 1 and (_local.hour, _local.minute) == (8, 0)
      and nova.proxima_execucao > datetime.utcnow(), nova.proxima_execucao)
check("o cartao diz 'Todo dia 1o do mes'", "Todo dia 1º do mês às 08:00" in r.text)

r = as_("adm").get(f"/admin/automacoes/{nova.id}/editar")
check("editar reabre o resumo com o que foi gravado",
      "Editar automação" in r.text and 'name="tipo" value="resumo_auditorias"' in r.text
      and 'value="mes_anterior" selected' in r.text)
salvar({**BASE_RES, "automacao_id": nova.id, "tipo": "coleta_nfse",
        "destinatarios": "gestao@vm.com.br", "dia_mes": "5"})
db.expire_all()
_nova = db.get(NfseAutomacao, nova.id)
check("o tipo de automacao existente e o do banco (campo escondido nao converte)",
      _nova.tipo == "resumo_auditorias" and _nova.dia_mes == 5, (_nova.tipo, _nova.dia_mes))

salvar({"tipo": "erros_sistema", "nome": "Erros diarios", "frequencia": "diaria", "hora": "07:30",
        "escopo": "selecionadas", "empresas": [T1], "ativo": "1", "periodo": "24h",
        "gravidade": "media", "incluir_log": "1", "so_com_novidade": "1",
        "destinatarios": "ti@vm.com.br"})
_err = db.query(NfseAutomacao).filter_by(nome="Erros diarios").first()
check("erro do sistema nao e de empresa: escopo forcado para 'todas'",
      _err is not None and _err.escopo == "todas" and _err.empresas is None
      and json.loads(_err.config)["gravidade"] == "media")
r = salvar({"nome": "Coleta mensal", "frequencia": "mensal", "hora": "07:00", "escopo": "todas",
            "dia_mes": "1"})
check("a coleta nao aceita mensal", "Escolha se a coleta é diária ou semanal" in r.text)

r = as_("adm").get(f"/admin/automacoes/execucao/{e_res.id}")
check("detalhe do resumo: por empresa, nao concluidas e o valor",
      all(s in r.text for s in ("Por empresa", "Não concluídas", "BETA DISTRIBUIDORA", "215,50",
                                "sem valor apurado")))
check("com o nome da empresa escapado", "&lt;b&gt;COMERCIO" in r.text and "<b>COMERCIO" not in r.text)
check("e o botao da planilha fora do boost", 'hx-boost="false"' in r.text and "/planilha" in r.text)
r = as_("adm").get(f"/admin/automacoes/execucao/{e_res.id}/planilha")
check("planilha do resumo: xlsx baixado, nunca renderizado",
      r.status_code == 200 and r.content[:2] == b"PK"
      and "attachment" in r.headers.get("content-disposition", "")
      and r.headers.get("x-content-type-options") == "nosniff", r.status_code)
_wb = openpyxl.load_workbook(io.BytesIO(r.content))
check("montada do detalhe da rodada: os numeros que foram no e-mail",
      _wb["Auditorias"].max_row == 4, _wb["Auditorias"].max_row)

r = as_("adm").get(f"/admin/automacoes/execucao/{e1.id}")
check("detalhe dos erros: diagnosticados com link para o relatorio e o log agrupado",
      all(s in r.text for s in ("Erros diagnosticados", "/admin/erros/", "3×"))
      and "&lt;script&gt;" in r.text and "<script>alert" not in r.text)
check("informe de erros nao tem planilha",
      as_("adm").get(f"/admin/automacoes/execucao/{e1.id}/planilha").status_code == 404)

r = as_("adm").post(f"/admin/automacoes/{ID_ERR}/rodar")
check("'Rodar agora' do informativo responde na hora",
      r.status_code == 200 and "Informativo em preparação" in r.text and 'hx-trigger="every 4s"' in r.text)
svc._ultima_manual.join(timeout=30)

check("operador nao abre o modal", as_("ope").get("/admin/automacoes/nova?tipo=resumo_auditorias").status_code == 403)
check("nem grava informativo", as_("ope").post("/admin/automacoes/salvar",
                                               data={**BASE_RES, "dia_mes": "1"}).status_code == 403)


# ============================================================================
print("\n=== 8. Colunas novas chegam a banco que ja tinha as tabelas ===")

from sqlalchemy import create_engine, inspect as _inspect, text  # noqa: E402
from backend.app.infrastructure.database import schema_upgrades  # noqa: E402

CAMINHO_VELHO = "./_informativos_schema_velho.db"
if os.path.exists(CAMINHO_VELHO):
    os.remove(CAMINHO_VELHO)
eng_velho = create_engine("sqlite:///" + CAMINHO_VELHO)
TABELAS = ((NfseAutomacao, "nfse_automacoes", None),
           (NfseAutomacaoExecucao, "nfse_automacao_execucoes", None),
           # Em audit_tasks, so as de 25/09: as anteriores tem suite propria.
           (AuditTask, "audit_tasks", ["vl_estornar", "vl_aproveitar", "vl_recolher", "vl_conferir"]))
for modelo, tabela, colunas in TABELAS:
    modelo.__table__.create(bind=eng_velho)
    with eng_velho.begin() as _conn:
        for _nome in (colunas or [n for n, _ddl in schema_upgrades.COLUNAS_NOVAS.get(tabela, [])]):
            _conn.execute(text(f"ALTER TABLE {tabela} DROP COLUMN {_nome}"))
with eng_velho.begin() as _conn:
    _conn.execute(text("INSERT INTO nfse_automacoes (id, nome, ativo, frequencia, hora, escopo, "
                       "created_at, updated_at) VALUES ('velha', 'Coleta antiga', 1, 'diaria', "
                       "'07:00', 'todas', '2026-09-24 10:00:00', '2026-09-24 10:00:00')"))
schema_upgrades.aplicar(eng_velho)
for modelo, tabela, _c in TABELAS:
    faltando = ({col.name for col in modelo.__table__.columns}
                - {col["name"] for col in _inspect(eng_velho).get_columns(tabela)})
    check(f"{tabela}: nenhuma coluna do model falta depois do startup", not faltando, faltando)
with eng_velho.connect() as _conn:
    _tipo_velha = _conn.execute(text("SELECT tipo FROM nfse_automacoes WHERE id = 'velha'")).scalar()
check("automacao que ja existia vira 'coleta_nfse' (o que ela sempre foi)",
      _tipo_velha == "coleta_nfse", _tipo_velha)
check("o valor em jogo nasce SEM default: nulo continua significando 'nunca medido'",
      all("DEFAULT" not in ddl.upper() for n, ddl in schema_upgrades.COLUNAS_NOVAS["audit_tasks"]
          if n.startswith("vl_")))
eng_velho.dispose()
try:
    os.remove(CAMINHO_VELHO)
except OSError:
    pass


# ============================================================================
email_envio.smtplib.SMTP = _smtp_real
erros_sistema._diretorio, erros_sistema._arquivo_log = _dir_real, _log_real
db.close()
_eng.dispose()
shutil.rmtree(_TMP, ignore_errors=True)
if os.path.exists("_informativos_test.db"):
    try:
        os.remove("_informativos_test.db")
    except OSError:
        pass

print("\n" + "=" * 62)
print(f"RESULTADO: {OK}/{OK + FALHAS} passaram")
for f in _falhou:
    print("  -", f)
sys.exit(1 if FALHAS else 0)
