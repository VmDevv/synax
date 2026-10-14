# -*- coding: utf-8 -*-
"""Lote de NFS-e em segundo plano e a pasta das NFS-e por EMPRESA.

Relato do Sandro em 01/10/2026: "no VmHub o sistema não para quando eu coloco
para baixar diversas empresas; no Synax ele para do nada". O lote rodava dentro
do clique, com 240 s para todas as empresas — cabiam 5 ou 6 — e o limite de 6
cliques a cada 15 minutos barrava, calado, quem clicava de novo. E as notas iam
para outro servidor, numa pasta por CNPJ do TOMADOR: no primeiro lote grande,
72 das 100 pastas criadas tinham o CPF de um cliente da empresa no nome.

O que esta suíte trava:

1. A pasta: a da EMPRESA (a que já existe na rede termina com o CNPJ e é
   REAPROVEITADA), uma subpasta por dia, o arquivo com a CHAVE da nota, e o
   evento com nome próprio — senão ele sobrescreveria a nota.
2. O ADN: os bytes gravados são os que vieram; 429 é "adiada", não "vazio";
   "acabou" e "faltou tempo" saem diferentes.
3. O lote: o clique só começa e volta; cada empresa sai com o status dela (o
   VmHub dizia "finalizado com sucesso" com empresa falhando); um lote por vez,
   nunca junto da coleta automática; "Parar" para de verdade sem perder nota.
"""
import base64
import zipfile
import gzip
import os
import shutil
import sys
import tempfile
import threading
import time
import types
import uuid

sys.path.insert(0, os.path.abspath('.'))

import dotenv  # noqa: E402
dotenv.load_dotenv = lambda *a, **k: False
import dotenv.main  # noqa: E402
dotenv.main.load_dotenv = lambda *a, **k: False

_TMP_STORAGE = tempfile.mkdtemp(prefix="nfse_lote_")
_CERT_DIR = tempfile.mkdtemp(prefix="nfse_lote_cert_")
os.environ["NFSE_STORAGE_DIR"] = _TMP_STORAGE
os.environ["CERTIFICADOS_DIR"] = _CERT_DIR
os.environ["DATABASE_URL"] = "sqlite:///./_nfse_lote_test.db"
os.environ["SUPABASE_URL"] = "http://fake"
os.environ["SUPABASE_KEY"] = "fake"
os.environ["SUPABASE_SERVICE_KEY"] = "fake"
os.environ["FERNET_SECRET_KEY"] = "hn5Ck1J8pVXHRV9V3Q6mYbXW4y2S1TqLpRZ0oN7dG8A="
if os.path.exists("_nfse_lote_test.db"):
    os.remove("_nfse_lote_test.db")

from sqlalchemy import types as _satypes  # noqa: E402
import sqlalchemy.dialects.postgresql as pg  # noqa: E402


class _UUIDShim(_satypes.TypeDecorator):
    impl = _satypes.Uuid
    cache_ok = True

    def __init__(self, *a, **kw):
        kw.pop('as_uuid', None)
        super().__init__()

    def process_bind_param(self, value, dialect):
        return uuid.UUID(value) if isinstance(value, str) else value


pg.UUID = _UUIDShim

import backend.app.infrastructure.auth.supabase_client as sb  # noqa: E402
TOKENS = {}


class _Auth:
    def get_user(self, token):
        u = TOKENS.get(token.replace("Bearer ", "").strip())
        if not u:
            raise Exception("invalid")
        return types.SimpleNamespace(user=u)


class _Client:
    auth = _Auth()


sb.get_supabase_client = lambda: _Client()
sb.get_supabase_admin_client = lambda: _Client()
import backend.app.api.dependencies.auth as dep_auth  # noqa: E402
dep_auth.get_supabase_client = lambda: _Client()

from backend.app.infrastructure.database.session import engine as _eng  # noqa: E402
if _eng.url.get_backend_name() != "sqlite":
    raise SystemExit("ABORTADO: engine nao e SQLite")

import backend.app.main as m  # noqa: E402
from backend.app.infrastructure.database.session import SessionLocal, engine  # noqa: E402
from backend.app.domain.models.base import Base  # noqa: E402
from backend.app.domain.models.auth_models import Tenant, UserProfile, UserTenant  # noqa: E402
from backend.app.domain.models.nfse import NfseDocument, NfseCursor  # noqa: E402
from backend.app.domain.models.certificates import TenantCertificate  # noqa: E402
from backend.app.core.security import encrypt_data  # noqa: E402
from backend.app.services.nfse import storage, adn_client  # noqa: E402
from backend.app.services.nfse import lote as lote_nfse  # noqa: E402
from backend.app.services.nfse import automacao  # noqa: E402
import backend.app.api.routes.nfse as rota_nfse  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

Base.metadata.create_all(bind=engine)

results = []


def check(name, cond, detail=""):
    results.append((name, cond, detail))
    print(("  OK  " if cond else " FALHA") + f" | {name}" + (f"  [{detail}]" if detail and not cond else ""))


CHAVE = "32049062244709461000103000000000000123028384347067"


def nota_nacional(numero: str, chave: str = CHAVE, toma: str = "06377159692") -> str:
    doc_toma = f"<CNPJ>{toma}</CNPJ>" if len(toma) == 14 else f"<CPF>{toma}</CPF>"
    return f"""<?xml version="1.0" encoding="utf-8"?>
<NFSe versao="1.00" xmlns="http://www.sped.fazenda.gov.br/nfse">
 <infNFSe Id="NFS{chave}">
  <xLocEmi>São Mateus</xLocEmi><nNFSe>{numero}</nNFSe><cLocIncid>3204906</cLocIncid>
  <cStat>100</cStat><dhProc>2026-09-10T09:00:00-03:00</dhProc>
  <emit><CNPJ>11222333000181</CNPJ><xNome>PRESTADOR EXEMPLO LTDA</xNome>
    <enderNac><cMun>3204906</cMun><UF>ES</UF></enderNac></emit>
  <valores><vBC>100.00</vBC><pAliqAplic>2.00</pAliqAplic><vISSQN>2.00</vISSQN><vLiq>100.00</vLiq></valores>
  <DPS versao="1.00"><infDPS Id="DPS1">
    <dhEmi>2026-09-08T13:30:39-03:00</dhEmi><dCompet>2026-09-01</dCompet>
    <prest><CNPJ>11222333000181</CNPJ></prest>
    <toma>{doc_toma}<xNome>CLIENTE DA NOTA</xNome></toma>
    <serv><cServ><cTribNac>141201</cTribNac><xDescServ>Servico</xDescServ></cServ></serv>
    <valores><vServPrest><vServ>100.00</vServ></vServPrest></valores>
  </infDPS></DPS>
 </infNFSe>
</NFSe>"""


EVENTO = f"""<?xml version="1.0" encoding="utf-8"?>
<evento versao="1.00" xmlns="http://www.sped.fazenda.gov.br/nfse">
 <infEvento Id="EVT{CHAVE}101101001"><chNFSe>{CHAVE}</chNFSe>
  <dhProc>2026-09-11T10:00:00-03:00</dhProc></infEvento>
</evento>"""


print("\n=== 1. Nome da empresa a partir do arquivo do certificado ===")
casos = {
    "BIANCARDI AUTO PECAS LTDA 20149617000180.pfx": "BIANCARDI AUTO PECAS LTDA",
    "3R PARTICIPACOES E EMPREENDIMENTOS LTDA_47307450000195.pfx": "3R PARTICIPACOES E EMPREENDIMENTOS LTDA",
    "46658392 BRUNELA AZEREDO COUTINHO SUEIRO 46658392000181 SENHA 13579046.pfx":
        "46658392 BRUNELA AZEREDO COUTINHO SUEIRO",
    "ADEMILSE GUIDINI 73214019791 SENHA 13579046.pfx": "ADEMILSE GUIDINI",
    "AI TRANSPORTES MUNICIPAIS LTDA.pfx": "AI TRANSPORTES MUNICIPAIS LTDA",
    "1.pfx": "",
}
for arquivo, esperado in casos.items():
    obtido = storage.nome_do_certificado(arquivo)
    check(f"'{arquivo[:40]}' -> '{esperado}'", obtido == esperado, f"obtido={obtido!r}")

print("\n=== 2. A pasta da empresa que JA existe na rede e reaproveitada ===")
existentes = ["01-10 notas NFS-E", "06250952000140",
              "ADEL - ASSOCIACAO PARA O DESENVOLVIMENTO 06250952000140",
              "BIANCARDI AUTO PECAS LTDA 20149617000180",
              "BIANCARDI AUTO PECAS LTDA 20149617000261"]
check("acha a pasta pelo CNPJ, mesmo com nome diferente do certificado",
      storage.pasta_da_empresa("06.250.952/0001-40", "ADEL ASSOCIACAO X", existentes)
      == "ADEL - ASSOCIACAO PARA O DESENVOLVIMENTO 06250952000140")
check("pasta so com o CNPJ (leiaute antigo, por tomador) NAO e a da empresa",
      storage.pasta_da_empresa("06250952000140", "X", ["06250952000140"]) == "X 06250952000140")
check("filial tem a propria pasta",
      storage.pasta_da_empresa("20149617000261", "BIANCARDI", existentes)
      == "BIANCARDI AUTO PECAS LTDA 20149617000261")
novo = storage.pasta_da_empresa("11222333000181", 'EMPRESA: NOVA/"TESTE" COM UM NOME QUE PASSA DE QUARENTA', [])
check("pasta nova: nome limpo, cortado em 40 e o CNPJ no fim",
      novo == "EMPRESA NOVATESTE COM UM NOME QUE PASSA 11222333000181", novo)
check("o dia e AAAA-MM-DD", len(storage.dia_de_hoje()) == 10 and storage.dia_de_hoje()[4] == "-")

print("\n=== 3. Nome do arquivo: a CHAVE da nota; evento nao sobrescreve a nota ===")
check("nota -> NFSe_<chave>.xml",
      storage.nome_do_arquivo(nota_nacional("77").encode(), "NSU_1_10.xml") == f"NFSe_{CHAVE}.xml")
nome_ev = storage.nome_do_arquivo(EVENTO.encode(), "NSU_1_11.xml")
check("evento -> Evento_<id>.xml, diferente da nota",
      nome_ev.startswith("Evento_") and nome_ev != f"NFSe_{CHAVE}.xml", nome_ev)
check("nota ABRASF (sem chave nacional) -> NFSe_<prestador>_<numero>.xml",
      storage.nome_do_arquivo(b"<ConsultarNfseResposta/>", "nota.xml", "03813799000123", "298231")
      == "NFSe_03813799000123_298231.xml")
check("XML sem identificacao usa a reserva",
      storage.nome_do_arquivo(b"<x/>", "NSU_1_12.xml") == "NSU_1_12.xml")
check("XML ilegivel usa a reserva", storage.nome_do_arquivo(b"nao e xml", "nota.xml") == "nota.xml")
rel = storage.gravar_em("../../fora", "../a.xml", b"<x/>")
check("gravar nao escapa da raiz", rel is not None and os.path.realpath(
    os.path.join(_TMP_STORAGE, rel)).startswith(os.path.realpath(_TMP_STORAGE)), rel)

print("\n=== 4. Leitura acha a nota gravada antes da mudanca ===")
_env = os.environ.pop("NFSE_STORAGE_DIR")
try:
    check("sem raiz configurada, a leitura procura tambem na raiz antiga",
          storage.RAIZES_ANTERIORES[0] in storage._raizes_de_leitura())
    check("a raiz nova e a pasta NFS-E- do fiscal", storage.raiz().endswith(r"ARQUIVOS APURAÇÃO\NFS-E-"))
finally:
    os.environ["NFSE_STORAGE_DIR"] = _env
check("com raiz configurada, so ela", storage._raizes_de_leitura() == [_TMP_STORAGE])

print("\n=== 5. ADN: bytes como vieram, 429 e motivo da parada ===")
latin = '<?xml version="1.0" encoding="ISO-8859-1"?><a>S\xe3o Mateus</a>'.encode("latin-1")
texto = adn_client._descompactar(base64.b64encode(gzip.compress(latin)).decode())
check("byte que nao e UTF-8 volta IDENTICO", adn_client.bytes_do_xml(texto) == latin)


class _Resp:
    def __init__(self, status, corpo=None):
        self.status_code = status
        self._c = corpo
        self.text = "" if corpo is None else str(corpo)

    def json(self):
        return self._c


def _rodar_adn(respostas, **kw):
    fila = list(respostas)

    class _Cli:
        def __enter__(self): return self
        def __exit__(self, *a): return False
        def get(self, url, **k): return fila.pop(0)

    import contextlib
    real_http, real_mtls = adn_client.httpx.Client, adn_client._material_mtls
    adn_client.httpx.Client = lambda **k: _Cli()
    adn_client._material_mtls = contextlib.contextmanager(lambda p, s: (yield ("c", "k")))
    estado = {}
    try:
        docs = list(adn_client.baixar_lote(b"x", "s", "11222333000181", 0, estado=estado, **kw))
        return docs, estado, None
    except Exception as err:
        return [], estado, err
    finally:
        adn_client.httpx.Client, adn_client._material_mtls = real_http, real_mtls


def _pagina(n0, qtd):
    return {"LoteDFe": [{"NSU": n0 + i, "ArquivoXml": base64.b64encode(
        gzip.compress(b"<x/>")).decode()} for i in range(qtd)]}


_, est, err = _rodar_adn([_Resp(429, {"erro": "Too Many Requests"})])
check("HTTP 429 vira ADIADA (ConsumoIndevido), nao 'sem notas'",
      isinstance(err, adn_client.ConsumoIndevido), repr(err))
docs, est, _ = _rodar_adn([_Resp(200, _pagina(1, 3))])
check("pagina incompleta: parada = fim", est.get("parada") == "fim" and len(docs) == 3, est)
docs, est, _ = _rodar_adn([_Resp(200, _pagina(1, 50)), _Resp(200, _pagina(51, 50))], max_paginas=2)
check("teto de paginas: parada = paginas (ainda ha mais)", est.get("parada") == "paginas", est)
docs, est, _ = _rodar_adn([_Resp(404)])
check("404: parada = fim", est.get("parada") == "fim", est)

# ------------------------------------------------------------------ cenário
db = SessionLocal()
tA, tB, tC, tD, tE, tF = (uuid.uuid4() for _ in range(6))
EMPRESAS = {
    tA: ("ALFA SERVICOS LTDA", "48135251000100"),   # 2 notas + 1 evento
    tB: ("BETA LTDA", "22222222000191"),            # nada novo
    tC: ("GAMA LTDA", "33333333000191"),            # 1 nota e o ADN adia
    tD: ("DELTA LTDA", "44444444000191"),           # certificado sem senha
    tE: ("EPSILON LTDA", "55555555000191"),         # certificado inativo
    tF: ("ZETA LTDA", "66666666000191"),            # faltou tempo: continua
}
for tid, (nome, cnpj) in EMPRESAS.items():
    db.add(Tenant(id=tid, name=nome, cnpj=cnpj))
# Só para o teste de parar: vem DEPOIS da ZETA na ordem do lote (por nome).
tZ = uuid.uuid4()
db.add(Tenant(id=tZ, name="ZZ ULTIMA LTDA", cnpj="77777777000191"))
u_adm, u_ope, u_ope2 = uuid.uuid4(), uuid.uuid4(), uuid.uuid4()
db.add_all([UserProfile(id=u_adm, tenant_id=None, role="admin", name="Sandro"),
            UserProfile(id=u_ope, tenant_id=None, role="operador", name="Livia"),
            UserProfile(id=u_ope2, tenant_id=None, role="operador", name="Outro")])
db.commit()
db.add_all([UserTenant(user_id=u_ope, tenant_id=tA), UserTenant(user_id=u_ope2, tenant_id=tB)])
for tid, (nome, cnpj) in EMPRESAS.items():
    db.add(TenantCertificate(tenant_id=str(tid), cnpj=cnpj, arquivo=f"{nome} {cnpj}.pfx",
                             senha_encrypted=encrypt_data("" if tid == tD else "senha"),
                             ativo=tid != tE))
    with open(os.path.join(_CERT_DIR, f"{nome} {cnpj}.pfx"), "wb") as f:
        f.write(b"pfx")
db.commit()
db.close()
TOKENS["adm"] = types.SimpleNamespace(id=str(u_adm), email="adm@x.com")
TOKENS["ope"] = types.SimpleNamespace(id=str(u_ope), email="ope@x.com")
TOKENS["ope2"] = types.SimpleNamespace(id=str(u_ope2), email="ope2@x.com")

# A pasta da ALFA já existe na rede com outro nome (criada pelo robô do SIEG).
os.makedirs(os.path.join(_TMP_STORAGE, "ALFA SERVICOS LTD 48135251000100"))

from backend.app.api.dependencies.db import get_rls_db as _real_rls  # noqa: E402


def _fake_rls():
    s = SessionLocal()
    try:
        yield s
    finally:
        s.close()


m.app.dependency_overrides[_real_rls] = _fake_rls
c = TestClient(m.app, raise_server_exceptions=False)


def as_(tok):
    c.cookies.clear()
    if tok:
        c.cookies.set("access_token", f"Bearer {tok}")
    return c


SEGURA = threading.Event()   # segura o ADN falso para ver o lote "rodando"
SEGURA.set()
LENTO = {"on": False}


def _fake_baixar(pfx, senha, cnpj, ultimo_nsu=0, max_paginas=40, limite_seg=None, estado=None):
    SEGURA.wait(10)
    estado = estado if estado is not None else {}
    estado["parada"] = "fim"
    if cnpj == "48135251000100":
        yield (10, nota_nacional("1", toma="48135251000100"))
        yield (11, nota_nacional("2", CHAVE[:-1] + "8", toma="48135251000100"))
        yield (12, EVENTO)
    elif cnpj == "33333333000191":
        yield (20, nota_nacional("3", CHAVE[:-1] + "9"))
        raise adn_client.ConsumoIndevido("O ADN limitou as consultas (HTTP 429).")
    elif cnpj == "66666666000191":
        for i in range(30 if LENTO["on"] else 2):
            if LENTO["on"]:
                time.sleep(0.05)
            yield (100 + i, f"<x>{i}</x>")
        estado["parada"] = "tempo"


adn_client.baixar_lote = _fake_baixar
rota_nfse.adn_client.baixar_lote = _fake_baixar

print("\n=== 6. O clique so COMECA o lote e volta na hora ===")
SEGURA.clear()
todas = [str(t) for t in EMPRESAS]
t0 = time.time()
r = as_("adm").post("/nfse/consultar-lote", data={"tenant_ids": todas})
gasto = time.time() - t0
check("responde 200 com o painel", r.status_code == 200 and 'id="nfse-lote"' in r.text, r.status_code)
check(f"volta sem esperar o ADN ({gasto:.2f}s)", gasto < 5)
check("o painel ja mostra o lote em andamento", "Em andamento" in r.text and "every 2s" in r.text)
check("a tabela tem uma linha por empresa", all(n in r.text for n, _ in EMPRESAS.values()))

r = as_("adm").get("/nfse/drawer")
check("enquanto roda, a gaveta avisa e desliga o botao do lote",
      "Há um lote em andamento" in r.text and ":disabled=\"aba === 'empresas' && true\"" in r.text)

r = as_("adm").post("/nfse/consultar-lote", data={"tenant_ids": todas})
check("segundo lote ao mesmo tempo e recusado com o motivo",
      r.status_code == 200 and "Já há um lote em andamento" in r.text)

r = as_("ope2").get("/nfse/lote")
check("operador que nao comecou o lote so fica sabendo que ha um rodando",
      "Há um lote de NFS-e em andamento" in r.text and "ALFA" not in r.text)

SEGURA.set()
lote_nfse.esperar()
atual = lote_nfse.atual()
check("o lote termina", atual is not None and not atual.rodando and atual.status == "concluido",
      atual.status if atual else None)

print("\n=== 7. Cada empresa sai com o status dela ===")
por_nome = {ln.empresa: ln for ln in atual.linhas}
alfa = por_nome["ALFA SERVICOS LTDA"]
check("ALFA: SUCESSO com 3 arquivos (2 notas + 1 evento)", alfa.status == "SUCESSO" and alfa.qtd == 3,
      f"{alfa.status} {alfa.qtd}")
check("BETA: SUCESSO (Vazio)", por_nome["BETA LTDA"].status == "SUCESSO (Vazio)", por_nome["BETA LTDA"].status)
gama = por_nome["GAMA LTDA"]
check("GAMA: ADIADA pelo ADN, com o que ja veio gravado",
      gama.status == "ADIADA" and gama.qtd == 1 and "gravado" in gama.destino, f"{gama.status} {gama.destino}")
check("DELTA: FALHOU (Sem Senha)", por_nome["DELTA LTDA"].status == "FALHOU (Sem Senha)")
eps = por_nome["EPSILON LTDA"]
check("EPSILON: FALHOU com o motivo do certificado", eps.status == "FALHOU" and "inativo" in eps.destino,
      eps.destino)
zeta = por_nome["ZETA LTDA"]
check("ZETA: SUCESSO (continua) — faltou tempo nao e 'em dia'",
      zeta.status == "SUCESSO (continua)" and "rode o lote de novo" in zeta.detalhe, zeta.status)
fim = atual.log[-1][1]
check("o fim do log NAO diz sucesso com empresa falhando",
      fim.startswith("Lote NFS-e finalizado") and "2 com falha" in fim and "1 adiada" in fim, fim)
check("o log tem uma linha por passo", len(atual.log) >= 2 + 2 * len(EMPRESAS), len(atual.log))

print("\n=== 8. Os arquivos estao na pasta da empresa, por dia ===")
dia = storage.dia_de_hoje()
pasta_alfa = os.path.join(_TMP_STORAGE, "ALFA SERVICOS LTD 48135251000100", dia)
arquivos = sorted(os.listdir(pasta_alfa)) if os.path.isdir(pasta_alfa) else []
check("reaproveitou a pasta que ja existia (achada pelo CNPJ)", len(arquivos) == 3, arquivos)
check("nota gravada com a chave", f"NFSe_{CHAVE}.xml" in arquivos, arquivos)
check("evento gravado a parte", any(a.startswith("Evento_") for a in arquivos), arquivos)
check("o destino na tela e a pasta do dia", alfa.destino == storage.caminho_completo(
    os.path.join("ALFA SERVICOS LTD 48135251000100", dia)), alfa.destino)
check("pasta nova leva o nome do certificado e o CNPJ",
      os.path.isdir(os.path.join(_TMP_STORAGE, "ZETA LTDA 66666666000191", dia)))
check("empresa sem documento nao deixa pasta vazia",
      not any(n.endswith("22222222000191") for n in os.listdir(_TMP_STORAGE)))
check("nenhuma pasta com o CPF do tomador",
      not any(n.startswith("06377159692") for n in os.listdir(_TMP_STORAGE)))
s = SessionLocal()
doc = s.query(NfseDocument).filter_by(tenant_id=str(tA), numero="1").first()
check("a nota entrou no banco apontando para o arquivo novo",
      doc is not None and doc.arquivo_xml.startswith("ALFA SERVICOS LTD 48135251000100"),
      doc.arquivo_xml if doc else None)
cur = s.query(NfseCursor).filter_by(tenant_id=str(tA)).first()
check("o cursor andou ate o evento (ultimo documento gravado)", cur is not None and cur.nsu == 12,
      cur.nsu if cur else None)
s.close()
r = as_("adm").get(f"/nfse/documento/11222333000181/1/{doc.mun_prest or 'sem_municipio'}/xml")
check("o XML da nota abre pelo caminho novo", r.status_code == 200 and b"NFS" in r.content, r.status_code)

print("\n=== 9. Painel depois do fim ===")
r = as_("adm").get("/nfse/lote?acompanhando=1")
check("terminado, o painel para de se atualizar", "every 2s" not in r.text)
check("e recarrega a tabela de notas uma vez", 'hx-get="/nfse/tabela"' in r.text)
r = as_("adm").get("/nfse/lote")
check("aberto depois (sem acompanhar), nao recarrega a tabela", 'hx-get="/nfse/tabela"' not in r.text)
check("o painel mostra status e destino", "SUCESSO (Vazio)" in r.text and "FALHOU (Sem Senha)" in r.text)

print("\n=== 10. Nunca junto da coleta automatica ===")
trava = automacao._trava("coleta_nfse")
trava.acquire()
try:
    r = as_("adm").post("/nfse/consultar-lote", data={"tenant_ids": [str(tB)]})
    check("com a coleta rodando, o lote nao comeca e diz por que",
          "coleta automática" in r.text and lote_nfse.atual() is atual)
finally:
    trava.release()

print("\n=== 11. Parar: o que veio fica, o resto nao e consultado ===")
s = SessionLocal()
s.query(NfseCursor).delete()
s.commit()
s.close()
LENTO["on"] = True
r = as_("ope").post("/nfse/consultar-lote", data={"tenant_ids": [str(tA)]})
check("operador comeca lote da empresa dele", r.status_code == 200 and 'id="nfse-lote"' in r.text
      and "Já há um lote" not in r.text, r.status_code)
lote_nfse.esperar()
check("o lote do operador e dele", lote_nfse.atual().dono_nome == "Livia")
r = as_("ope").post("/nfse/consultar-lote", data={"tenant_ids": [str(tB)]})
check("operador NAO usa empresa alheia", r.status_code == 403, r.status_code)

r = as_("adm").post("/nfse/consultar-lote", data={"tenant_ids": [str(tZ), str(tF)]})
deadline = time.time() + 5
while time.time() < deadline and not any(ln.qtd for ln in lote_nfse.atual().linhas):
    time.sleep(0.02)
r = as_("adm").post("/nfse/lote/parar")
check("parar responde o painel com 'Parando'", r.status_code == 200 and "Parando" in r.text, r.status_code)
lote_nfse.esperar()
parado = lote_nfse.atual()
zeta2, ultima = parado.linhas[0], parado.linhas[1]
check("o lote segue a ordem por nome", zeta2.empresa == "ZETA LTDA", zeta2.empresa)
check("o lote fica 'parado'", parado.status == "parado", parado.status)
check("a empresa em andamento sai INTERROMPIDA com o que ja gravou",
      zeta2.status == "INTERROMPIDA" and 0 < zeta2.qtd < 30, f"{zeta2.status} {zeta2.qtd}")
check("a que nao comecou sai NAO CONSULTADA", ultima.status == "NÃO CONSULTADA", ultima.status)
check("o fim do log diz que o lote foi parado", parado.log[-1][1].startswith("Lote parado"),
      parado.log[-1][1])
s = SessionLocal()
cur = s.query(NfseCursor).filter_by(tenant_id=str(tF)).first()
check("o cursor parou no ultimo documento gravado", cur is not None and cur.nsu == 100 + zeta2.qtd - 1,
      f"{cur.nsu if cur else None} vs {zeta2.qtd}")
s.close()
LENTO["on"] = False

print("\n=== 12. A tela ===")
gaveta = open("frontend/templates/partials/nfse/drawer_import.html", encoding="utf-8").read()
check("o lote responde no painel, nao na tabela", 'hx-target="#nfse-lote"' in gaveta)
check("a janela fecha no clique (antes da resposta)", "hx-on::before-request" in gaveta
      and "aberto = false" in gaveta)
check("sem empresa marcada o pedido nem sai", "lote-sem-empresa" in gaveta)
pagina = open("frontend/templates/pages/nfse.html", encoding="utf-8").read()
check("a pagina carrega o painel do lote", 'id="nfse-lote"' in pagina and 'hx-get="/nfse/lote"' in pagina)
check("erro do lote (429/403) aparece na tela", "htmx:responseError" in pagina)
r = as_(None).get("/nfse/lote")
check("anonimo nao ve o painel", r.status_code in (401, 403), r.status_code)

print("\n=== 13. Certificado: o CNPJ de DENTRO dele; matriz consulta filial ===")
# Um e-CNPJ de teste no padrão ICP-Brasil: CNPJ num otherName do SAN.
import datetime as _dt  # noqa: E402
from cryptography import x509  # noqa: E402
from cryptography.x509.oid import NameOID  # noqa: E402
from cryptography.hazmat.primitives import hashes, serialization  # noqa: E402
from cryptography.hazmat.primitives.asymmetric import rsa  # noqa: E402
from cryptography.hazmat.primitives.serialization import pkcs12  # noqa: E402
from backend.app.services.nfse import certificados as svc_cert  # noqa: E402


def _pfx(cnpj_dentro: str, senha: bytes = b"s") -> bytes:
    chave = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    nome = x509.Name([x509.NameAttribute(NameOID.COMMON_NAME, "EMPRESA DE TESTE")])
    san = x509.SubjectAlternativeName([x509.OtherName(
        x509.ObjectIdentifier("2.16.76.1.3.3"), b"\x04\x0e" + cnpj_dentro.encode())])
    cert = (x509.CertificateBuilder().subject_name(nome).issuer_name(nome)
            .public_key(chave.public_key()).serial_number(x509.random_serial_number())
            .not_valid_before(_dt.datetime(2026, 1, 1)).not_valid_after(_dt.datetime(2027, 1, 1))
            .add_extension(san, critical=False).sign(chave, hashes.SHA256()))
    return pkcs12.serialize_key_and_certificates(
        b"t", chave, cert, None, serialization.BestAvailableEncryption(senha))


PFX_MATRIZ = _pfx("05911108000150")
check("lê o CNPJ do otherName ICP-Brasil", adn_client.cnpj_do_certificado(PFX_MATRIZ, "s") == "05911108000150")
check("senha errada: não lê, não quebra", adn_client.cnpj_do_certificado(PFX_MATRIZ, "x") is None)
check("arquivo que não é certificado: None", adn_client.cnpj_do_certificado(b"pfx", "s") is None)
check("matriz consultando a FILIAL passa (mesma raiz)",
      adn_client.certificado_de_outra_empresa(PFX_MATRIZ, "s", "05911108001806") is None)
msg = adn_client.certificado_de_outra_empresa(PFX_MATRIZ, "s", "48135251000100")
check("certificado de OUTRA empresa é recusado com o motivo", msg and "outra empresa" in msg, msg)
check("a filial vai no cnpjConsulta da URL",
      adn_client._url_do_nsu(50, "05911108001806").endswith("/50?lote=true&cnpjConsulta=05911108001806"))
check("nome do arquivo com CNPJ de outra empresa: aviso",
      svc_cert.cnpj_de_outra_empresa("X 99999999000191.pfx", "48135251000100") == "99999999000191")
check("nome do arquivo da matriz para a filial: sem aviso",
      svc_cert.cnpj_de_outra_empresa("M R 05911108000150.pfx", "05911108001806") is None)
reg = types.SimpleNamespace(arquivo="OUTRA 99999999000191.pfx", ativo=True, validade=None)
diag = svc_cert.diagnosticar(reg, "48135251000100", ["OUTRA 99999999000191.pfx"])
check("diagnóstico marca o certificado de outra empresa",
      not diag["ok"] and "outra empresa" in diag["motivo"], diag)

# No lote: empresa com certificado de outra empresa FALHA sem consultar o ADN.
tW = uuid.uuid4()
s = SessionLocal()
s.add(Tenant(id=tW, name="WAGNER LTDA", cnpj="12345678000195"))
s.add(TenantCertificate(tenant_id=str(tW), cnpj="12345678000195", arquivo="wagner.pfx",
                        senha_encrypted=encrypt_data("s"), ativo=True))
s.commit()
s.close()
with open(os.path.join(_CERT_DIR, "wagner.pfx"), "wb") as f:
    f.write(PFX_MATRIZ)
CHAMADAS_ADN = []
_fake_antes = rota_nfse.adn_client.baixar_lote


def _conta(pfx, senha, cnpj, *a, **k):
    CHAMADAS_ADN.append(cnpj)
    return _fake_antes(pfx, senha, cnpj, *a, **k)


rota_nfse.adn_client.baixar_lote = _conta
as_("adm").post("/nfse/consultar-lote", data={"tenant_ids": [str(tW)]})
lote_nfse.esperar()
ln = lote_nfse.atual().linhas[0]
check("empresa com certificado de outra empresa: FALHOU com o motivo",
      ln.status == "FALHOU" and "outra empresa" in ln.destino, f"{ln.status} {ln.destino}")
check("e o ADN nem foi chamado", "12345678000195" not in CHAMADAS_ADN, CHAMADAS_ADN)
rota_nfse.adn_client.baixar_lote = _fake_antes

print("\n=== 14. Recebida ou emitida: uma regra só ===")
from backend.app.services.nfse import classificacao  # noqa: E402
check("prestador = empresa -> emitida", classificacao.tipo("48135251000100", "48135251000100", "1") == "emitida")
check("tomador = empresa -> recebida", classificacao.tipo("48.135.251/0001-00", "1", "48135251000100") == "recebida")
check("nem um nem outro -> outra", classificacao.tipo("48135251000100", "1", "2") == "outra")
nota = types.SimpleNamespace(prestador="48135251000100", tomador="06377159692",
                             rz_prestador="ALFA", rz_tomador="FULANO")
d = classificacao.descrever(nota, "48135251000100", "ALFA")
check("na emitida, a contraparte é o tomador", d["contraparte_doc"] == "06377159692"
      and d["contraparte_nome"] == "FULANO", d)

print("\n=== 15. A tela diz de quem é a nota e de que lado ===")
# Estas notas foram gravadas fora de lote: a tela precisa de "todas" para
# mostrá-las (por padrão ela abre no último lote — seção 19).
from backend.app.core.templating import moeda_br  # noqa: E402
from datetime import datetime as _datetime  # noqa: E402
s = SessionLocal()
s.add_all([
    NfseDocument(tenant_id=str(tB), prestador="11222333000181", numero="R-1", mun_prest="3204906",
                 tomador="22222222000191", rz_prestador="FORNECEDOR DE SERVICO", valor_servico=100.0,
                 iss=2.0, iss_retido="Sim", ir=1.5, valor_liquido=96.5,
                 dt_emissao=_datetime(2026, 9, 10), origem="adn"),
    NfseDocument(tenant_id=str(tB), prestador="22222222000191", numero="E-1", mun_prest="3204906",
                 tomador="06377159692", rz_tomador="CLIENTE PF", valor_servico=250.0,
                 iss_retido="Não", valor_liquido=250.0, dt_emissao=_datetime(2026, 9, 11),
                 origem="adn"),
])
s.commit()
s.close()
r = as_("adm").get(f"/nfse/tabela?tenant_id={tB}&q_lote=todas")
check("a tabela abre com a empresa escolhida", r.status_code == 200, r.status_code)
check("cada linha diz Recebida ou Emitida", "Recebida</span>" in r.text and "Emitida</span>" in r.text)
check("resumo das recebidas com valor e retenções",
      f"1 nota(s) · R$ {moeda_br(100)}" in r.text and f"ISS retido R$ {moeda_br(2)}" in r.text, r.text[:0])
check("resumo das emitidas", f"1 nota(s) · R$ {moeda_br(250)}" in r.text)
check("na emitida a contraparte é o tomador (CPF)", "CPF 06377159692" in r.text)
check("com a empresa escolhida, a coluna Empresa sai", ">Empresa</th>" not in r.text)
r = as_("adm").get(f"/nfse/tabela?tenant_id={tB}&q_tipo=recebidas&q_lote=todas")
check("filtro Recebidas", "R-1" in r.text and "E-1" not in r.text)
r = as_("adm").get(f"/nfse/tabela?tenant_id={tB}&q_tipo=emitidas&q_lote=todas")
check("filtro Emitidas", "E-1" in r.text and "R-1" not in r.text)
r = as_("adm").get("/nfse/tabela?q_tipo=emitidas&q_lote=todas")
check("filtro Emitidas com todas as empresas", "E-1" in r.text and "R-1" not in r.text)
r = as_("adm").get("/nfse/tabela?q_lote=todas")
check("com todas as empresas, a coluna Empresa aparece com o nome",
      ">Empresa</th>" in r.text and "BETA LTDA" in r.text)
check("e a tela pede uma empresa para mostrar o resumo", "Escolha uma empresa" in r.text)
check("as colunas ABRASF ficam atrás de 'Mais colunas'",
      "nfseMaisColunas" in r.text and 'class="nfse-extra' in r.text and "Base de Cálculo" in r.text)
r = as_("ope2").get(f"/nfse/tabela?tenant_id={tB}&q_tipo=recebidas&q_lote=todas")
check("operador vê o filtro na empresa dele", r.status_code == 200 and "R-1" in r.text, r.status_code)
r = as_("ope").get(f"/nfse/tabela?tenant_id={tB}")
check("operador NÃO vê empresa alheia", r.status_code == 403, r.status_code)

print("\n=== 16. Relatório do lote ===")
import io as _io  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
s = SessionLocal()
s.query(NfseCursor).filter_by(tenant_id=str(tA)).delete()
s.commit()
s.close()
as_("adm").post("/nfse/consultar-lote", data={"tenant_ids": [str(tA), str(tB), str(tD)]})
lote_nfse.esperar()
atual = lote_nfse.atual()
check("o lote guarda o relatório", atual.relatorio is not None and atual.relatorio[:2] == b"PK")
r = as_("adm").get("/nfse/lote?acompanhando=1")
check("o painel oferece o relatório para baixar, sem boost",
      'href="/nfse/lote/relatorio"' in r.text and 'hx-boost="false"' in r.text)
r = as_("adm").get("/nfse/lote/relatorio")
check("o relatório baixa como planilha",
      r.status_code == 200 and r.content[:2] == b"PK"
      and "attachment" in r.headers.get("content-disposition", ""), r.status_code)
wb = load_workbook(_io.BytesIO(r.content))
check("abas: Visão Geral, Notas, Por empresa, Pendências",
      {"Visão Geral", "Notas", "Por empresa", "Pendências"} <= set(wb.sheetnames), wb.sheetnames)
notas = list(wb["Notas"].iter_rows(min_row=2, values_only=True))
check("a aba Notas tem as 2 notas do lote (o evento não é nota)", len(notas) == 2, len(notas))
check("cada nota diz a empresa e se é Recebida",
      all(n[0] == "ALFA SERVICOS LTDA" and n[2] == "Recebida" for n in notas), notas[:1])
por_emp = {row[0]: row for row in wb["Por empresa"].iter_rows(min_row=2, values_only=True)}
alfa_row = por_emp.get("ALFA SERVICOS LTDA")
check("Por empresa: ALFA com 2 recebidas e R$ 200", alfa_row and alfa_row[4] == 2 and alfa_row[5] == 200,
      alfa_row)
pend = [row[0] for row in wb["Pendências"].iter_rows(min_row=2, values_only=True)]
check("Pendências traz a empresa sem senha", "DELTA LTDA" in pend and "ALFA SERVICOS LTDA" not in pend, pend)
copias = []
for base, _p, nomes in os.walk(os.path.join(_TMP_STORAGE, "_relatorios")):
    copias += [n for n in nomes if n.startswith("lote_nfse_")]
check("cópia salva em _relatorios", bool(copias), copias)
r = as_("ope2").get("/nfse/lote/relatorio")
check("quem não começou o lote não baixa o relatório", r.status_code == 404, r.status_code)

as_("adm").post("/nfse/consultar-lote", data={"tenant_ids": [str(tA)], "reiniciar": "1"})
lote_nfse.esperar()
wb = load_workbook(_io.BytesIO(lote_nfse.atual().relatorio))
notas = list(wb["Notas"].iter_rows(min_row=2, values_only=True))
check("'Recomeçar do zero': o relatório lista as notas que já estavam no banco", len(notas) == 2, len(notas))

print("\n=== 17. Script de correção: relê o XML e tira da filial a nota da matriz ===")
import importlib.util as _ilu  # noqa: E402
_spec = _ilu.spec_from_file_location("corrigir_nfse", os.path.join("scripts", "corrigir_nfse.py"))
corrigir = _ilu.module_from_spec(_spec)
_spec.loader.exec_module(corrigir)

tM, tFi = uuid.uuid4(), uuid.uuid4()
s = SessionLocal()
s.add_all([Tenant(id=tM, name="MATRIZ SA", cnpj="05911108000150"),
           Tenant(id=tFi, name="MATRIZ SA FILIAL", cnpj="05911108001806")])
s.commit()
s.add(TenantCertificate(tenant_id=str(tFi), cnpj="05911108001806", arquivo="matriz.pfx",
                        senha_encrypted=encrypt_data("s"), ativo=True))
s.add(NfseCursor(tenant_id=str(tFi), cnpj="05911108001806", ultimo_nsu="50"))
with open(os.path.join(_CERT_DIR, "matriz.pfx"), "wb") as f:
    f.write(PFX_MATRIZ)
# A mesma nota (da matriz) nas duas empresas; e uma nota gravada pelo leitor
# antigo: município vazio e sem tomador, com o XML completo na pasta.
for tid in (tM, tFi):
    s.add(NfseDocument(tenant_id=str(tid), prestador="11222333000181", numero="M-1",
                       mun_prest="3204906", tomador="05911108000150", origem="adn"))
rel_velho = storage.gravar_em("velho", "nota_velha.xml",
                              nota_nacional("V-1", CHAVE[:-1] + "1", toma="05911108000150").encode())
s.add(NfseDocument(tenant_id=str(tM), prestador="11222333000181", numero="V-1", mun_prest="",
                   tomador=None, dt_competencia=None, arquivo_xml=rel_velho, origem="adn"))
s.commit()

docs = s.query(NfseDocument).filter(NfseDocument.tenant_id.in_([str(tM), str(tFi)])).all()
empresas = {str(t.id): (t.name, classificacao.digitos(t.cnpj)) for t in s.query(Tenant).all()}
placar = corrigir.reler(s, docs, todas=False)
velha = [d for d in docs if d.numero == "V-1"][0]
check("relê o XML: tomador e competência preenchidos",
      velha.tomador == "05911108000150" and velha.dt_competencia is not None, placar)
check("o município (parte da chave) é corrigido", velha.mun_prest == "3204906", velha.mun_prest)
linhas, contaminadas = corrigir.zerar_cursores_de_filial(s, empresas, s.query(NfseCursor).all())
check("a filial consultada com o certificado da matriz é achada", str(tFi) in contaminadas, contaminadas)
check("e o cursor dela volta a 0", s.query(NfseCursor).filter_by(tenant_id=str(tFi)).first().ultimo_nsu == "0")
removidas, _fic = corrigir.tirar_notas_alheias(s, docs, empresas, contaminadas)
s.commit()
check("a nota da matriz sai da FILIAL", s.query(NfseDocument).filter_by(tenant_id=str(tFi), numero="M-1").count() == 0)
check("e continua na MATRIZ", s.query(NfseDocument).filter_by(tenant_id=str(tM), numero="M-1").count() == 1)
s.close()

print("\n=== 18. A busca avulsa também recusa certificado de outra empresa ===")
r = as_("adm").post("/nfse/consultar-adn",
                    data={"tenant_id": str(tA), "cnpj": "48135251000100", "senha": "s"},
                    files=[("certificado", ("c.pfx", PFX_MATRIZ, "application/x-pkcs12"))])
check("400 com o motivo", r.status_code == 400 and "outra empresa" in r.text, r.status_code)


print("\n=== 19. A tela abre nas notas do ÚLTIMO lote ===")
# Pedido do Sandro em 02/10/2026: com 5.539 notas guardadas, o que o lote
# acabou de trazer se perdia no histórico.
ultimo = lote_nfse.atual()
s = SessionLocal()
do_lote = s.query(NfseDocument).filter_by(lote_id=ultimo.id).count()
s.close()
check("as notas gravadas pelo lote levam a marca dele", do_lote == 2, do_lote)
r = as_("adm").get("/nfse/tabela")
check("por padrão a tabela mostra só o último lote",
      "Mostrando só as notas do último lote" in r.text and "— 2 nota(s)." in r.text
      and "R-1" not in r.text, r.status_code)
check("o seletor 'Mostrar' vem em 'Último lote'",
      '<option value="ultimo" selected>' in r.text)
r = as_("adm").get("/nfse")
check("a página também abre no último lote", "Mostrando só as notas do último lote" in r.text)
r = as_("adm").get("/nfse/tabela?q_lote=todas")
check("'Todas as notas guardadas' mostra o histórico",
      "R-1" in r.text and "Mostrando só as notas do último lote" not in r.text)
r = as_("adm").post("/nfse/zip", data={"lote_id": ultimo.id})
nomes = zipfile.ZipFile(_io.BytesIO(r.content)).namelist() if r.status_code == 200 else []
check("na visão do lote, o ZIP sem nada marcado traz só o lote", len(nomes) == 2, nomes)

lote_nfse._atual = None          # o servidor reiniciou: o lote saiu da memória
r = as_("adm").get("/nfse/tabela")
check("depois de reiniciar, a tela acha o último lote pelo banco",
      "Mostrando só as notas do último lote" in r.text and "— 2 nota(s)." in r.text)
lote_nfse._atual = ultimo

r = as_("ope2").get("/nfse/tabela")
check("quem não tem lote em empresa sua vê todas, sem o seletor",
      "R-1" in r.text and 'name="q_lote"' not in r.text, r.status_code)

r = as_("adm").post("/nfse/importar", data={"tenant_id": str(tB)},
                    files=[("arquivos", ("u.xml", nota_nacional("U-1", CHAVE[:-1] + "5",
                                                                toma="22222222000191"), "text/xml"))])
check("depois de importar arquivo, a tabela mostra todas (a nota importada aparece)",
      r.status_code == 200 and "U-1" in r.text
      and "Mostrando só as notas do último lote" not in r.text, r.status_code)


engine.dispose()
for caminho in ("_nfse_lote_test.db",):
    if os.path.exists(caminho):
        try:
            os.remove(caminho)
        except OSError:
            pass
shutil.rmtree(_TMP_STORAGE, ignore_errors=True)
shutil.rmtree(_CERT_DIR, ignore_errors=True)

fails = [n for n, ok, _ in results if not ok]
print("\n" + "=" * 60)
print(f"RESULTADO: {len(results) - len(fails)}/{len(results)} passaram")
for f in fails:
    print("  -", f)
sys.exit(1 if fails else 0)
