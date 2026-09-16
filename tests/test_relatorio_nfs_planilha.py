# -*- coding: utf-8 -*-
"""Relatório de NFs: a planilha sai INTEIRA, não só a página aberta.

Pedido do Sandro (27/08/2026): "ela só gera planilha com as notas da página".

Duas causas somadas:

1. A tabela da aba tem `data-paginate="50"` e a paginação esconde as linhas
   das outras páginas com `row.style.display = 'none'`. O `exportGeneralReportCSV`
   pulava toda linha com `display === 'none'` — sem ter como distinguir
   "escondida pela paginação" de "escondida pela busca". Resultado: o CSV saía
   com os 50 registros da tela e o resto do mês ficava de fora, em silêncio.
   Hoje quem marca a linha descartada pela BUSCA é a classe `js-filtrado-fora`,
   e é só ela que o export consulta.

2. O Excel do backend (`/api/export/excel/`) nem tinha a aba do relatório:
   `relatorio_geral` não era exportado. Quem quisesse o mês todo não tinha por
   onde — nem pelo CSV da tela, nem pela planilha completa.

O que este teste tranca: a planilha do backend traz UMA LINHA POR ITEM, todas,
e o export da tela não volta a olhar `style.display`.
"""
import io
import os
import sys

sys.path.insert(0, os.path.abspath('.'))
sys.path.insert(0, os.path.abspath('backend'))

import openpyxl  # noqa: E402

from app.services.audit_engine.excel_exporter import generate_excel_from_results  # noqa: E402

OK = FALHAS = 0
_falhou = []


def check(nome, condicao, detalhe=""):
    global OK, FALHAS
    if condicao:
        OK += 1
        print(f"  OK   | {nome}")
    else:
        FALHAS += 1
        _falhou.append(nome)
        print(f" FALHA | {nome}  [{detalhe}]")


def item(i):
    return {
        "key": str(i).zfill(44),
        "num_nf": str(1000 + i),
        "cnpj_emitente": "11222333000181",
        "uf_emitente": "ES",
        "razao_social": f"FORNECEDOR {i}",
        "descr_item": f"PRODUTO {i}",
        "vl_total": 10.0 * i,
        "bc_icms": 1.0,
        "icms": 2.0,
        "icms_sn": 0.0,
        "ipi": 0.0,
        "ipi_devolvido": 0.0,
        "st": 0.0,
        "fcp_st": 0.0,
        "cfop_xml": "5102",
        "cfop_sped": "1102",
        "icms_mono": 0.0,
    }


# ============================================================================
print("\n=== 1. A planilha do backend traz o relatório inteiro ===")

# 137 itens: de propósito NÃO é múltiplo do tamanho da página (50). Se algum dia
# a exportação voltar a se guiar pela paginação, o número quebra aqui.
ITENS = [item(i) for i in range(137)]
buf = generate_excel_from_results({"resumo": {}, "relatorio_geral": ITENS})
wb = openpyxl.load_workbook(buf)

check("a aba do relatório existe", "Relatório de NFs" in wb.sheetnames, wb.sheetnames)
ws = wb["Relatório de NFs"]
check("uma linha por item, sem parar na primeira página",
      ws.max_row == len(ITENS) + 1, f"max_row={ws.max_row} itens={len(ITENS)}")

cabecalho = [c.value for c in ws[1]]
for coluna in ("N° NF", "CNPJ Emitente", "Razão Social", "Descrição Item",
               "Valor Total", "CFOP XML", "CFOP SPED"):
    check(f"a coluna {coluna!r} está na planilha", coluna in cabecalho, cabecalho)

check("a chave de acesso vai junto, para achar o DANFE",
      "Chave de Acesso" in cabecalho, cabecalho)
check("o nome técnico do campo não vaza para o cabeçalho",
      "num_nf" not in cabecalho and "cfop_xml" not in cabecalho, cabecalho)

primeira = [c.value for c in ws[2]]
ultima = [c.value for c in ws[ws.max_row]]
check("a primeira nota é a primeira do relatório", primeira[0] == "1000", primeira[:2])
check("a última nota também entrou", ultima[0] == "1136", ultima[:2])

print("\n=== 1b. Itens Não Conferidos também vão para a planilha ===")

# É a lista que o contador confere à mão. Sai no Excel COM o motivo: sem ele a
# planilha diz que o item ficou de fora, mas não o que procurar na nota.
NAO_CONFERIDOS = [{
    "key": "9" * 44, "num": str(2000 + i), "num_item": str(i + 1),
    "emit_name": "DIST LAURENCINI", "cnpj_emitente": "36318376000196",
    "cod_item": f"COMPL{i}", "descr": f"PRODUTO COMPLEMENTAR {i}",
    "cfop_sped": "1102", "vl_item": 100.0 + i, "vl_icms": 12.0,
    "motivo": "O item do SPED não casou com nenhum item do XML.",
} for i in range(35)]

wb_nc = openpyxl.load_workbook(generate_excel_from_results(
    {"resumo": {}, "itens_nao_conferidos": NAO_CONFERIDOS}))
check("a aba dos não conferidos existe",
      "Itens Não Conferidos" in wb_nc.sheetnames, wb_nc.sheetnames)
ws_nc = wb_nc["Itens Não Conferidos"]
check("todos os itens saem, não só a primeira página",
      ws_nc.max_row == len(NAO_CONFERIDOS) + 1,
      f"max_row={ws_nc.max_row} itens={len(NAO_CONFERIDOS)}")

cab_nc = [c.value for c in ws_nc[1]]
check("a planilha diz POR QUE o item não foi conferido",
      "Por que não foi conferido" in cab_nc, cab_nc)
for coluna in ("N° NF", "Fornecedor", "Código no SPED", "Descrição no SPED",
               "CFOP SPED", "Valor do Item", "Chave de Acesso"):
    check(f"a coluna {coluna!r} está na planilha", coluna in cab_nc, cab_nc)

check("sem itens não conferidos, sem aba",
      "Itens Não Conferidos" not in openpyxl.load_workbook(
          generate_excel_from_results({"resumo": {}, "itens_nao_conferidos": []})).sheetnames)


print("\n=== 2. Relatório vazio não cria aba fantasma ===")
wb_vazio = openpyxl.load_workbook(
    generate_excel_from_results({"resumo": {}, "relatorio_geral": []}))
check("sem itens, sem aba", "Relatório de NFs" not in wb_vazio.sheetnames,
      wb_vazio.sheetnames)

print("\n=== 3. O export da tela não se guia mais pela paginação ===")
TPL = io.open(os.path.join("frontend", "templates", "partials", "curation",
                           "curation_results.html"), encoding="utf-8").read()

corpo_export = TPL[TPL.find("function exportGeneralReportCSV()"):]
corpo_export = corpo_export[:corpo_export.find("\n}")]

check("o export não olha style.display",
      "style.display === 'none'" not in corpo_export, corpo_export[:400])
check("o export descarta só o que a busca filtrou",
      "js-filtrado-fora" in corpo_export, corpo_export[:400])

corpo_filtro = TPL[TPL.find("function filterGeneralReport("):]
corpo_filtro = corpo_filtro[:corpo_filtro.find("\n}")]
check("a busca marca a linha descartada com a classe",
      "js-filtrado-fora" in corpo_filtro, corpo_filtro[:400])

# A paginação continua escondendo por style.display — é o que faz a tela ficar
# leve. O que não pode é o export voltar a confundir as duas coisas.
check("a paginação segue usando style.display",
      "row.style.display = (i >= start && i < end)" in TPL)

# Botão da aba "Itens Não Conferidos". Nasceu depois da correção do Relatório de
# NFs, então já usa a função genérica que ignora a paginação — o erro de baixar
# só a página aberta não pode se repetir a cada aba nova.
check("existe função de export genérica para qualquer tabela",
      "function exportarTabelaCSV(tabelaId, nomeArquivo)" in TPL)
_generica = TPL[TPL.find("function exportarTabelaCSV"):]
_generica = _generica[:_generica.find("function exportGeneralReportCSV")]
check("a função genérica NÃO olha style.display",
      "style.display" not in _generica, _generica[:400])
check("e respeita o filtro de busca",
      "js-filtrado-fora" in _generica)
check("o CSV sai com BOM, para o Excel não quebrar os acentos",
      "﻿" in _generica, _generica[-400:])

check("a aba de não conferidos tem botão de exportar",
      "exportarTabelaCSV('table-nao-conferidos'" in TPL)
check("e a tabela dela tem id para o botão achar",
      'id="table-nao-conferidos"' in TPL)
check("a tabela mostra o motivo, que é o que se analisa",
      "Por que não foi conferido" in TPL and "item.motivo" in TPL)

check("a tabela do relatório continua paginada na tela",
      'id="table-general-report"' in TPL and 'data-paginate="50"' in TPL)


# ============================================================================
print("\n" + "=" * 62)
print(f"RESULTADO: {OK}/{OK + FALHAS} passaram")
if _falhou:
    print("Falharam: " + ", ".join(_falhou))
print("=" * 62)

sys.exit(1 if FALHAS else 0)
