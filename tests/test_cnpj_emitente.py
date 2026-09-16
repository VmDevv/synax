# -*- coding: utf-8 -*-
"""CNPJ do emitente: o participante escriturado é mesmo quem assinou o XML?

Pedido do Sandro em 11/09/2026. O C100/D100 não traz CNPJ nenhum — traz o
COD_PART, um código do cadastro do cliente que o registro 0150 resolve. Trocar
de participante no Domínio é um clique, e quando isso acontece a auditoria
inteira continua verde: valor, imposto, item e CFOP vêm do XML certo, porque a
nota é encontrada pela CHAVE. O único campo errado é de quem o crédito foi
tomado — e era exatamente o que o motor não olhava.

O que esta suíte trava:

1. **O 0150 passa a guardar CNPJ e CPF.** Antes o parser lia só o nome, e o
   confronto não tinha como existir. `participants` virou dict de dicts, o que
   ameaça os ~15 pontos do `cross_checker` que usavam o valor como texto: a
   seção 4 confere que nenhuma lista do relatório passou a mostrar o dict cru
   no lugar da razão social.

2. **Qual lado conferir depende do IND_EMIT, não do IND_OPER.** Em emissão de
   terceiros o COD_PART é o emitente; em emissão própria ele é o remetente e o
   emitente é a própria empresa. Confundir os dois acusaria divergência em toda
   nota de entrada emitida pelo cliente.

3. **Saída emitida por terceiro não é conferida.** Ali o COD_PART é o
   DESTINATÁRIO, e compará-lo com o emitente inventaria divergência.

4. **Matriz × filial tem mensagem própria.** Mesma raiz de CNPJ é a mesma
   empresa no estabelecimento errado, e o conserto não é trocar o fornecedor.

5. **Modo "só SPED" não confere nada.** Sem XML, o emitente sai do próprio
   0150 — comparar daria igual sempre, e conferência que nunca acusa nada é
   pior do que conferência nenhuma.
"""
import io
import os
import sys
import tempfile
import uuid

sys.path.insert(0, os.path.abspath('.'))

import dotenv
dotenv.load_dotenv = lambda *a, **k: False
import dotenv.main
dotenv.main.load_dotenv = lambda *a, **k: False

os.environ["DATABASE_URL"] = "sqlite:///./_cnpj_emit_test.db"
os.environ["SUPABASE_URL"] = "http://fake"
os.environ["SUPABASE_KEY"] = "fake"
os.environ["SUPABASE_SERVICE_KEY"] = "fake"
os.environ["FERNET_SECRET_KEY"] = "hn5Ck1J8pVXHRV9V3Q6mYbXW4y2S1TqLpRZ0oN7dG8A="

if os.path.exists("_cnpj_emit_test.db"):
    os.remove("_cnpj_emit_test.db")

from sqlalchemy import types as _satypes  # noqa: E402
import sqlalchemy.dialects.postgresql as pg  # noqa: E402


class _UUIDShim(_satypes.TypeDecorator):
    impl = _satypes.Uuid
    cache_ok = True

    def __init__(self, *a, **kw):
        kw.pop('as_uuid', None)
        super().__init__()

    def process_bind_param(self, value, dialect):
        return uuid.UUID(value) if isinstance(value, str) else value


pg.UUID = _UUIDShim

import backend.app.infrastructure.auth.supabase_client as sb  # noqa: E402
sb.get_supabase_client = lambda: type("C", (), {"auth": None})()
sb.get_supabase_admin_client = sb.get_supabase_client

from backend.app.infrastructure.database.session import engine as _eng  # noqa: E402
if _eng.url.get_backend_name() != "sqlite":
    raise SystemExit(f"ABORTADO: engine aponta para {_eng.url.get_backend_name()}.")

import json  # noqa: E402
import shutil  # noqa: E402
from backend.app.infrastructure.database.session import SessionLocal, engine  # noqa: E402
from backend.app.domain.models.base import Base  # noqa: E402
from backend.app.domain.models import rules, audit  # noqa: F401,E402
from backend.app.domain.models.audit import AuditTask  # noqa: E402
from backend.app.services.audit_engine.sped_parser import parse_sped  # noqa: E402
from backend.app.services.audit_engine.cross_checker import (  # noqa: E402
    conferir_cnpj_do_emitente, execute_audit, id_participante, nome_participante)

Base.metadata.create_all(bind=engine)

OK = FALHAS = 0
_falhou = []


def check(nome, condicao, detalhe=""):
    global OK, FALHAS
    if condicao:
        OK += 1
        print(f"  OK   | {nome}")
    else:
        FALHAS += 1
        _falhou.append(nome)
        print(f" FALHA | {nome}  [{detalhe}]")


EMPRESA = "34449753000109"
FORNEC = "86527934000100"
FILIAL_DO_FORNEC = "86527934000271"   # mesma raiz, outro estabelecimento
OUTRO = "99888777000166"
PRODUTOR = "11122233344"              # CPF: participante pessoa física

# Uma chave por nota. Só precisam ser distintas e bater com o nome do arquivo
# de XML — o confronto de emitente NÃO lê o CNPJ embutido na chave.
CHAVE_OK = "35260834449753000109550010000010011000010010"
CHAVE_TROCADA = "35260834449753000109550010000010021000010021"
CHAVE_FILIAL = "35260834449753000109550010000010031000010032"


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

def xml_de(chave, num, cnpj_emit, nome_emit):
    return f"""<?xml version="1.0" encoding="UTF-8"?>
<nfeProc xmlns="http://www.portalfiscal.inf.br/nfe"><NFe><infNFe Id="NFe{chave}">
 <ide><nNF>{num}</nNF><serie>1</serie></ide>
 <emit><CNPJ>{cnpj_emit}</CNPJ><xNome>{nome_emit}</xNome>
  <enderEmit><UF>PR</UF></enderEmit></emit>
 <dest><CNPJ>{EMPRESA}</CNPJ><enderDest><UF>SP</UF></enderDest></dest>
 <det nItem="1"><prod><cProd>P1</cProd><xProd>PARAFUSO</xProd><CFOP>5102</CFOP>
   <NCM>73181500</NCM><vProd>100.00</vProd><qCom>1</qCom><vUnCom>100.00</vUnCom></prod>
  <imposto><ICMS><ICMS00><CST>00</CST><vBC>100.00</vBC><pICMS>18.00</pICMS>
    <vICMS>18.00</vICMS></ICMS00></ICMS></imposto></det>
 <total><ICMSTot><vNF>100.00</vNF><vST>0.00</vST></ICMSTot></total>
</infNFe></NFe></nfeProc>"""


def linha_c100(cod_part, num, chave, ind_emit="1", ind_oper="0"):
    return (f"|C100|{ind_oper}|{ind_emit}|{cod_part}|55|00|1|{num}|{chave}|"
            "18072026|18072026|100,00|0|0,00|0,00|100,00|0|0,00|0,00|0,00|"
            "18,00|0,00|0,00|0,00|0,00|")


LINHA_C170 = ("|C170|1|P1|PARAFUSO|1|UN|100,00|0,00|0|00|1102||100,00|18,00|"
              "18,00|0,00|0,00|0,00|0|00||0,00|0,00|0,00|50|100,00|1,65|1,65|"
              "50|100,00|7,60|7,60|")


def montar_sped():
    """Três notas do MESMO fornecedor no XML, escrituradas em três participantes.

    F1 é o fornecedor certo, F2 é outra empresa, F3 é a filial dele. Como as
    três notas vêm do mesmo emitente no XML, o que separa uma da outra é só o
    COD_PART — que é exatamente o campo sob conferência.
    """
    linhas = [
        f"|0000|017|0|01072026|31072026|EMPRESA TESTE|{EMPRESA}|SP|||",
        f"|0150|F1|FORNECEDOR CERTO|1058|{FORNEC}||||||",
        f"|0150|F2|OUTRA EMPRESA|1058|{OUTRO}||||||",
        f"|0150|F3|FORNECEDOR FILIAL|1058|{FILIAL_DO_FORNEC}||||||",
        f"|0150|PF|JOSE PRODUTOR|1058||{PRODUTOR}|||||",
        "|0200|P1|PARAFUSO SEXTAVADO||UN|||||",
        linha_c100("F1", "1001", CHAVE_OK), LINHA_C170,
        linha_c100("F2", "1002", CHAVE_TROCADA), LINHA_C170,
        linha_c100("F3", "1003", CHAVE_FILIAL), LINHA_C170,
    ]
    fd, caminho = tempfile.mkstemp(suffix=".txt")
    os.close(fd)
    with io.open(caminho, "w", encoding="latin1") as f:
        f.write("\n".join(linhas) + "\n")
    return caminho


# ============================================================================
print("\n=== 1. O registro 0150 passa a trazer o CNPJ do participante ===")

SPED = montar_sped()
dados_sped = parse_sped(SPED)
parts = dados_sped["participants"]

check("o 0150 vira um cadastro, não só o nome", isinstance(parts.get("F1"), dict),
      type(parts.get("F1")).__name__)
check("com o CNPJ do campo 5", parts["F1"].get("cnpj") == FORNEC,
      parts["F1"].get("cnpj"))
check("e a razão social do campo 3",
      parts["F1"].get("nome") == "FORNECEDOR CERTO", parts["F1"].get("nome"))
check("participante pessoa física traz CPF e nenhum CNPJ",
      parts["PF"].get("cpf") == PRODUTOR and parts["PF"].get("cnpj") == "",
      (parts["PF"].get("cpf"), parts["PF"].get("cnpj")))

check("o helper de nome devolve texto, nunca o dict",
      nome_participante(parts, "F1") == "FORNECEDOR CERTO",
      nome_participante(parts, "F1"))
check("COD_PART sem 0150 cai no próprio código, para não sumir da tela",
      nome_participante(parts, "9999") == "9999", nome_participante(parts, "9999"))
check("o helper de identificação prefere o CNPJ",
      id_participante(parts, "F1") == FORNEC, id_participante(parts, "F1"))
check("e usa o CPF quando é o que existe",
      id_participante(parts, "PF") == PRODUTOR, id_participante(parts, "PF"))


# ============================================================================
print("\n=== 2. A conferência, ramo a ramo ===")

XML_CERTO = {"emit_id": FORNEC, "emit_name": "FORNECEDOR CERTO"}


def conferir(cod_part, xml, ind_emit="1", ind_oper="0"):
    entrada = {"num": "1001", "emit": cod_part,
               "ind_emit": ind_emit, "ind_oper": ind_oper}
    return conferir_cnpj_do_emitente("NF-e", CHAVE_OK, entrada, parts, xml, EMPRESA)


check("participante certo não gera apontamento", conferir("F1", XML_CERTO) is None)

trocado = conferir("F2", XML_CERTO)
check("participante de outra empresa é apontado", trocado is not None)
check("e classificado como fornecedor trocado",
      trocado and trocado.get("tipo_divergencia") == "divergente",
      trocado and trocado.get("tipo_divergencia"))
check("com os DOIS CNPJs na linha — é o que se usa para corrigir",
      trocado and trocado.get("sped_cnpj") == OUTRO
      and trocado.get("xml_cnpj") == FORNEC,
      trocado and (trocado.get("sped_cnpj"), trocado.get("xml_cnpj")))
check("e com o COD_PART, que é por onde se acha a nota no Domínio",
      trocado and trocado.get("cod_part") == "F2", trocado and trocado.get("cod_part"))
check("a mensagem nomeia os dois lados",
      trocado and OUTRO in trocado["reason"] and FORNEC in trocado["reason"],
      trocado and trocado.get("reason"))

filial = conferir("F3", XML_CERTO)
check("mesma raiz de CNPJ é estabelecimento trocado, não fornecedor trocado",
      filial and filial.get("tipo_divergencia") == "filial_divergente",
      filial and filial.get("tipo_divergencia"))

sem_cad = conferir("F9", XML_CERTO)
check("COD_PART sem 0150 aparece em vez de sumir calado",
      sem_cad and sem_cad.get("tipo_divergencia") == "sem_cadastro",
      sem_cad and sem_cad.get("tipo_divergencia"))
check("participante sem CNPJ nem CPF também aparece",
      conferir("PF", {"emit_id": FORNEC, "emit_name": "X"}) is not None)

# IND_EMIT = 0: quem emitiu foi a empresa, e o COD_PART ali é o remetente.
propria_ok = conferir("F2", {"emit_id": EMPRESA, "emit_name": "EMPRESA TESTE"},
                      ind_emit="0")
check("emissão própria confere contra o CNPJ da empresa, não contra o COD_PART",
      propria_ok is None, propria_ok and propria_ok.get("reason"))
propria_ma = conferir("F1", XML_CERTO, ind_emit="0")
check("e acusa quando o XML de 'emissão própria' veio de terceiro",
      propria_ma and propria_ma.get("tipo_divergencia") == "emissao_propria_de_terceiro",
      propria_ma and propria_ma.get("tipo_divergencia"))

# IND_OPER = 1: o COD_PART é o destinatário.
check("saída emitida por terceiro não é conferida",
      conferir("F2", XML_CERTO, ind_oper="1") is None)
check("mas saída de emissão própria continua sendo",
      conferir("F2", XML_CERTO, ind_emit="0", ind_oper="1") is not None)

check("XML sem emitente legível não vira erro de cadastro",
      conferir("F2", {"emit_id": "", "emit_name": ""}) is None)
check("pontuação no CNPJ não inventa divergência",
      conferir("F1", {"emit_id": "86.527.934/0001-00", "emit_name": "X"}) is None)


# ============================================================================
print("\n=== 3. Ponta a ponta: a auditoria aponta a nota no fornecedor errado ===")


def rodar(sped_only=False):
    xml_dir = tempfile.mkdtemp(prefix="synax_cnpj_emit_")
    for chave, num in ((CHAVE_OK, "1001"), (CHAVE_TROCADA, "1002"),
                       (CHAVE_FILIAL, "1003")):
        io.open(os.path.join(xml_dir, f"{chave}.xml"), "w", encoding="utf-8").write(
            xml_de(chave, num, FORNEC, "FORNECEDOR CERTO"))

    task_id, tenant = str(uuid.uuid4()), str(uuid.uuid4())
    s = SessionLocal()
    s.add(AuditTask(id=task_id, tenant_id=tenant, status="pendente",
                    file_name="sped.txt"))
    s.commit()
    execute_audit(task_id, tenant, SPED, xml_dir, s, sped_only_mode=sped_only)
    s.expire_all()
    tarefa = s.query(AuditTask).filter(AuditTask.id == task_id).first()
    status, dados = tarefa.status, json.loads(tarefa.result_data or "{}")
    inconsistencias = tarefa.inconsistencias
    s.close()
    shutil.rmtree(xml_dir, ignore_errors=True)
    return status, dados, inconsistencias


STATUS, RESULTADO, INCONSISTENCIAS = rodar()
check("a auditoria conclui", STATUS == "concluido", STATUS)

achados = {a["num"]: a for a in RESULTADO.get("cnpj_emitente_issues", [])}
check("as duas notas erradas são apontadas — e só elas",
      sorted(achados) == ["1002", "1003"], sorted(achados))
check("a nota lançada no participante certo fica de fora", "1001" not in achados)
check("a nota do fornecedor trocado sai classificada",
      achados.get("1002", {}).get("tipo_divergencia") == "divergente",
      achados.get("1002", {}).get("tipo_divergencia"))
check("a da filial também", achados.get("1003", {}).get("tipo_divergencia") == "filial_divergente",
      achados.get("1003", {}).get("tipo_divergencia"))
check("o apontamento leva a chave, para abrir o DANFE no SIEG",
      achados.get("1002", {}).get("key") == CHAVE_TROCADA,
      achados.get("1002", {}).get("key"))
check("e o nome do emitente do XML, que é o fornecedor de verdade",
      achados.get("1002", {}).get("emit_name") == "FORNECEDOR CERTO",
      achados.get("1002", {}).get("emit_name"))
check("o participante do SPED vai pelo nome, não pelo código",
      achados.get("1002", {}).get("sped_name") == "OUTRA EMPRESA",
      achados.get("1002", {}).get("sped_name"))

check("as duas entram na contagem de inconsistências do resumo",
      RESULTADO["resumo"]["inconsistencias"] >= 2,
      RESULTADO["resumo"]["inconsistencias"])
check("e a contagem chega à coluna que o Radar lê",
      INCONSISTENCIAS == RESULTADO["resumo"]["inconsistencias"],
      (INCONSISTENCIAS, RESULTADO["resumo"]["inconsistencias"]))

_, SO_SPED, _ = rodar(sped_only=True)
check("no modo só SPED a conferência não roda — não há com o que cruzar",
      SO_SPED.get("cnpj_emitente_issues") == [],
      SO_SPED.get("cnpj_emitente_issues"))


# ============================================================================
print("\n=== 4. Regressão: o participante virou dict e a tela continua com texto ===")

# `participants` era {COD_PART: nome}. Toda lista do relatório que mostrava o
# fornecedor lia esse valor direto; um dict escapando para qualquer uma delas
# apareceria na tela como "{'nome': ...}" e na planilha como uma célula ilegível.
CAMPOS_DE_NOME = ("emit", "emit_name", "fornecedor", "razao_social", "sped_name",
                  "emit_cnpj", "cnpj_emitente")
vazou = []
for chave, lista in RESULTADO.items():
    if not isinstance(lista, list):
        continue
    for item in lista:
        if not isinstance(item, dict):
            continue
        for campo in CAMPOS_DE_NOME:
            if isinstance(item.get(campo), (dict, list)):
                vazou.append(f"{chave}.{campo}")
check("nenhuma lista do relatório carrega o cadastro cru no lugar do nome",
      not vazou, sorted(set(vazou)))
check("o relatório serializa em JSON (é assim que ele é gravado)",
      isinstance(json.dumps(RESULTADO), str))


# ============================================================================
print("\n=== 5. A categoria existe nos quatro pontos de saída ===")

from backend.app.services.audit_engine.excel_exporter import (  # noqa: E402
    generate_excel_from_results)
from backend.app.services.audit_engine.email_catalog import CATEGORIAS  # noqa: E402
from backend.app.services.audit_engine.resumo_financeiro import SEM_VALOR  # noqa: E402
import openpyxl  # noqa: E402

tela = io.open("frontend/templates/partials/curation/curation_results.html",
               encoding="utf-8").read()
check("a tela tem o botão e o painel da aba",
      tela.count("tab-cnpj-emitente") >= 2, tela.count("tab-cnpj-emitente"))
check("e lê a lista que o motor produz", "result.cnpj_emitente_issues" in tela)

check("o resumo financeiro conhece a categoria",
      any(k == "cnpj_emitente_issues" for k, _, _ in SEM_VALOR))

cat = [c for c in CATEGORIAS if "cnpj_emitente_issues" in c["fontes"]]
check("o gerador de e-mail tem a categoria", len(cat) == 1, len(cat))
if cat:
    linha = cat[0]["formatador"](achados["1002"])
    check("e a linha do e-mail diz qual CNPJ está no lugar de qual",
          OUTRO in linha["detalhe"] and FORNEC in linha["detalhe"], linha["detalhe"])

livro = openpyxl.load_workbook(io.BytesIO(generate_excel_from_results(RESULTADO).read()))
check("a planilha ganha a aba", "CNPJ do Emitente" in livro.sheetnames,
      livro.sheetnames)
if "CNPJ do Emitente" in livro.sheetnames:
    cabecalho = [c.value for c in livro["CNPJ do Emitente"][1]]
    check("com cabeçalho em português, sem nome interno de campo",
          "CNPJ no SPED" in cabecalho and "CNPJ no XML" in cabecalho, cabecalho)
    check("e sem duas colunas chamadas 'Motivo'",
          cabecalho.count("Motivo") <= 1, cabecalho)


# ============================================================================
print("\n" + "=" * 62)
print(f"RESULTADO: {OK}/{OK + FALHAS} passaram")
if _falhou:
    print("Falharam: " + ", ".join(_falhou))
print("=" * 62)

engine.dispose()
for lixo in (SPED, "_cnpj_emit_test.db"):
    if lixo and os.path.exists(lixo):
        try:
            os.remove(lixo)
        except OSError:
            pass

sys.exit(1 if FALHAS else 0)
