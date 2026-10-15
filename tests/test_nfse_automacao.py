# -*- coding: utf-8 -*-
"""Coleta automática de NFS-e: agenda, trava, registro honesto e a tela.

Pedido do Sandro em 24/09/2026 (fase 1 de 3: download agendado + histórico +
tela; relatório e e-mail vêm depois). O que esta suíte trava:

1. **A hora é de Brasília.** O banco guarda UTC; "07:00" na tela é 10:00 UTC.
   E "próxima execução" é SEMPRE depois de agora — nunca o instante presente.

2. **"Consumo indevido" não é "nenhuma nota".** O ADN recusando por excesso de
   chamadas virava lote vazio, e a coleta automática mandaria dizer que a
   empresa não recebeu nota quando a consulta nem aconteceu. Agora é ADIADA.

3. **Nenhuma rodada sai verde com empresa não consultada.** Falha e adiada
   entram no status e no detalhe, com o motivo.

4. **Não roda em dobro.** A próxima execução é trocada no banco com UPDATE
   condicionado; dois processos, uma coleta. Servidor fora por dias roda UMA
   vez ao voltar, não uma por dia perdido.

5. **Só admin, e nada de thread nos testes.** O agendador liga no lifespan do
   servidor, não no import do app.
"""
import contextlib
import io
import json
import os
import shutil
import sys
import tempfile
import types
import uuid
from datetime import datetime, timedelta

sys.path.insert(0, os.path.abspath('.'))

import dotenv
dotenv.load_dotenv = lambda *a, **k: False
import dotenv.main
dotenv.main.load_dotenv = lambda *a, **k: False

os.environ["DATABASE_URL"] = "sqlite:///./_automacao_test.db"
os.environ["SUPABASE_URL"] = "http://fake"
os.environ["SUPABASE_KEY"] = "fake"
os.environ["SUPABASE_SERVICE_KEY"] = "fake"
os.environ["FERNET_SECRET_KEY"] = "hn5Ck1J8pVXHRV9V3Q6mYbXW4y2S1TqLpRZ0oN7dG8A="
_TMP_STORAGE = tempfile.mkdtemp(prefix="nfse_auto_test_")
os.environ["NFSE_STORAGE_DIR"] = _TMP_STORAGE
if os.path.exists("_automacao_test.db"):
    os.remove("_automacao_test.db")

from sqlalchemy import types as _satypes  # noqa: E402
import sqlalchemy.dialects.postgresql as pg  # noqa: E402


class _UUIDShim(_satypes.TypeDecorator):
    impl = _satypes.Uuid
    cache_ok = True

    def __init__(self, *a, **kw):
        kw.pop('as_uuid', None)
        super().__init__()

    def process_bind_param(self, value, dialect):
        return uuid.UUID(value) if isinstance(value, str) else value


pg.UUID = _UUIDShim

import backend.app.infrastructure.auth.supabase_client as sb  # noqa: E402
TOKENS = {}


class _Auth:
    def get_user(self, token):
        u = TOKENS.get(token.replace("Bearer ", "").strip())
        if not u:
            raise Exception("invalid")
        return types.SimpleNamespace(user=u)


class _Client:
    auth = _Auth()


sb.get_supabase_client = lambda: _Client()
sb.get_supabase_admin_client = lambda: _Client()
import backend.app.api.dependencies.auth as dep_auth  # noqa: E402
dep_auth.get_supabase_client = lambda: _Client()

from backend.app.infrastructure.database.session import engine as _eng  # noqa: E402
if _eng.url.get_backend_name() != "sqlite":
    raise SystemExit("ABORTADO: engine nao e SQLite")

import backend.app.main as m  # noqa: E402
from backend.app.infrastructure.database.session import SessionLocal  # noqa: E402
from backend.app.domain.models.base import Base  # noqa: E402
from backend.app.domain.models.auth_models import Tenant, UserProfile, UserTenant  # noqa: E402
from backend.app.domain.models.automacao import NfseAutomacao, NfseAutomacaoExecucao  # noqa: E402
from backend.app.domain.models.certificates import TenantCertificate  # noqa: E402
from backend.app.domain.models.nfse import NfseCursor, NfseDocument  # noqa: E402
from backend.app.services.nfse import adn_client, certificados  # noqa: E402
from backend.app.services.nfse import automacao as svc  # noqa: E402
import backend.app.api.routes.nfse as rota_nfse  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

OK = FALHAS = 0
_falhou = []


def check(nome, condicao, detalhe=""):
    global OK, FALHAS
    if condicao:
        OK += 1
        print(f"  OK   | {nome}")
    else:
        FALHAS += 1
        _falhou.append(nome)
        print(f" FALHA | {nome}  [{detalhe}]")


# ============================================================================
print("\n=== 1. A agenda: hora de Brasilia, sempre depois de agora ===")

p = svc.proxima_execucao
# 24/09/2026 e quinta-feira. 09:00 UTC = 06:00 em Brasilia.
agora = datetime(2026, 9, 24, 9, 0)
check("diaria 07:00, antes do horario: hoje as 07:00 (10:00 UTC)",
      p("diaria", "07:00", [], agora) == datetime(2026, 9, 24, 10, 0), p("diaria", "07:00", [], agora))
check("diaria 07:00, depois do horario: amanha",
      p("diaria", "07:00", [], datetime(2026, 9, 24, 10, 30)) == datetime(2026, 9, 25, 10, 0))
check("exatamente na hora: a PROXIMA e amanha (estritamente depois)",
      p("diaria", "07:00", [], datetime(2026, 9, 24, 10, 0)) == datetime(2026, 9, 25, 10, 0))
check("23:30 em Brasilia cai no dia seguinte em UTC",
      p("diaria", "23:30", [], datetime(2026, 9, 24, 1, 0)) == datetime(2026, 9, 24, 2, 30),
      p("diaria", "23:30", [], datetime(2026, 9, 24, 1, 0)))
# Semanal: segunda (0) e quinta (3). Quinta 08:00 local ja passou das 07:00.
check("semanal seg+qui, quinta depois do horario: vai para segunda",
      p("semanal", "07:00", [0, 3], datetime(2026, 9, 24, 11, 0)) == datetime(2026, 9, 28, 10, 0),
      p("semanal", "07:00", [0, 3], datetime(2026, 9, 24, 11, 0)))
check("semanal seg+qui, quarta: vai para quinta",
      p("semanal", "07:00", [0, 3], datetime(2026, 9, 23, 12, 0)) == datetime(2026, 9, 24, 10, 0))
check("semanal sem dia: nao inventa data", p("semanal", "07:00", [], agora) is None)
check("hora invalida: nao inventa data", p("diaria", "25:00", [], agora) is None)
check("validar_hora normaliza e recusa",
      (svc.validar_hora("7:05"), svc.validar_hora("07:60"), svc.validar_hora("abc"))
      == ("07:05", None, None))
check("descricao legivel da agenda",
      svc.descrever_agenda("semanal", "07:00", "0,2,4") == "Seg, Qua e Sex às 07:00"
      and svc.descrever_agenda("diaria", "18:30", None) == "Todo dia às 18:30",
      svc.descrever_agenda("semanal", "07:00", "0,2,4"))


# ============================================================================
print("\n=== 2. 'Consumo indevido' nao e 'nenhuma nota' ===")

check("a marca saiu da lista de 'sem documentos'",
      not any("CONSUMO" in x for x in adn_client.SEM_DOCUMENTOS), adn_client.SEM_DOCUMENTOS)
corpo = {"StatusProcessamento": "REJEICAO", "Mensagem": "Consumo Indevido"}
check("o corpo e reconhecido como consumo indevido", adn_client._consumo_indevido(corpo))
check("e nao como lote vazio", not adn_client._tem_lote_vazio(corpo))
check("ConsumoIndevido e um tipo de RespostaInesperada (a tela ja trata)",
      issubclass(adn_client.ConsumoIndevido, adn_client.RespostaInesperada))


class _Resp:
    def __init__(self, status, json_):
        self.status_code = status
        self._json = json_
        self.text = json.dumps(json_)
        self.headers = {"content-type": "application/json"}

    def json(self):
        return self._json


class _ClienteFake:
    def __init__(self, resposta):
        self.resposta = resposta

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def get(self, url, **kw):
        return self.resposta


@contextlib.contextmanager
def _material_falso(pfx, senha):
    yield ("cert.pem", "key.pem")


_mtls_real, _http_real = adn_client._material_mtls, adn_client.httpx.Client
adn_client._material_mtls = _material_falso
adn_client.httpx.Client = lambda **kw: _ClienteFake(_Resp(200, corpo))
try:
    list(adn_client.baixar_lote(b"x", "s", "11000000000000", 0))
    check("baixar_lote levanta ConsumoIndevido", False, "nao levantou")
except adn_client.ConsumoIndevido as err:
    check("baixar_lote levanta ConsumoIndevido", True)
    check("com mensagem que diz que nada se perdeu", "Nada se perdeu" in str(err), str(err))
finally:
    adn_client.httpx.Client = _http_real
    adn_client._material_mtls = _mtls_real


# ============================================================================
print("\n=== 3. Massa: empresas, certificados e o ADN de mentira ===")

Base.metadata.create_all(bind=_eng)
db = SessionLocal()
tA, tB, tC, tD, tE, tF = (uuid.uuid4() for _ in range(6))
CNPJ = {tA: "48135251000100", tB: "22222222000191", tC: "33333333000191",
        tD: "44444444000191", tE: "55555555000191", tF: "66666666000191"}
db.add_all([
    Tenant(id=tA, name="ALFA SERVICOS", cnpj=CNPJ[tA]),
    Tenant(id=tB, name="BETA LTDA", cnpj=CNPJ[tB]),
    Tenant(id=tC, name="GAMA COMERCIO", cnpj=CNPJ[tC]),
    Tenant(id=tD, name="DELTA INDUSTRIA", cnpj=CNPJ[tD]),
    Tenant(id=tE, name="EPSILON INATIVA", cnpj=CNPJ[tE], ativo=False),
    Tenant(id=tF, name="ZETA SEM CERTIFICADO", cnpj=CNPJ[tF]),
])
AGORA = datetime.utcnow()
for t, ativo, validade in ((tA, True, AGORA + timedelta(days=300)),
                           (tB, True, AGORA + timedelta(days=200)),
                           (tC, False, AGORA + timedelta(days=100)),
                           (tD, True, AGORA + timedelta(days=10)),
                           (tE, True, AGORA + timedelta(days=5))):
    db.add(TenantCertificate(tenant_id=str(t), cnpj=CNPJ[t], arquivo="x.pfx",
                             senha_encrypted="x", ativo=ativo, validade=validade))
u_adm, u_ope = uuid.uuid4(), uuid.uuid4()
db.add_all([UserProfile(id=u_adm, tenant_id=tA, role="admin", name="Admin"),
            UserProfile(id=u_ope, tenant_id=tA, role="operador", name="Operador")])
db.add(UserTenant(user_id=u_ope, tenant_id=tA))
db.commit()

XML_NFSE = """<?xml version="1.0" encoding="UTF-8"?>
<ConsultarNfseResposta xmlns="http://www.abrasf.org.br/nfse.xsd">
 <ListaNfse><CompNfse><Nfse><InfNfse>
  <Numero>298231</Numero>
  <CodigoVerificacao>ABC123XY</CodigoVerificacao>
  <DataEmissao>2026-09-16T10:22:00</DataEmissao>
  <Competencia>2026-09-01T00:00:00</Competencia>
  <ValoresNfse><BaseCalculo>267.26</BaseCalculo><Aliquota>2.00</Aliquota>
    <ValorIss>5.35</ValorIss><ValorLiquidoNfse>267.26</ValorLiquidoNfse></ValoresNfse>
  <PrestadorServico>
    <IdentificacaoPrestador><Cnpj>03813799000123</Cnpj>
      <InscricaoMunicipal>11994101</InscricaoMunicipal></IdentificacaoPrestador>
    <RazaoSocial>COMLINK S. A.</RazaoSocial>
    <Endereco><Uf>SP</Uf><CodigoMunicipio>3543402</CodigoMunicipio></Endereco>
  </PrestadorServico>
  <DeclaracaoPrestacaoServico><InfDeclaracaoPrestacaoServico>
    <Servico>
      <Valores><ValorServicos>267.26</ValorServicos><ValorPis>1.26</ValorPis>
        <ValorCofins>5.83</ValorCofins><ValorInss>0</ValorInss><ValorIr>0</ValorIr>
        <ValorCsll>0</ValorCsll><OutrasRetencoes>0</OutrasRetencoes>
        <ValorDeducoes>0</ValorDeducoes><DescontoIncondicionado>0</DescontoIncondicionado>
        <DescontoCondicionado>0</DescontoCondicionado></Valores>
      <IssRetido>2</IssRetido>
      <ItemListaServico>0105</ItemListaServico>
      <Discriminacao>SERVICO DE LINK DEDICADO DE INTERNET</Discriminacao>
    </Servico>
    <Tomador>
      <IdentificacaoTomador><CpfCnpj><Cnpj>48135251000100</Cnpj></CpfCnpj></IdentificacaoTomador>
      <RazaoSocial>ALFA SERVICOS</RazaoSocial>
      <Endereco><Uf>ES</Uf><CodigoMunicipio>3203205</CodigoMunicipio></Endereco>
    </Tomador>
  </InfDeclaracaoPrestacaoServico></DeclaracaoPrestacaoServico>
 </InfNfse></Nfse></CompNfse></ListaNfse>
</ConsultarNfseResposta>"""
XML_NFSE2 = XML_NFSE.replace("<Numero>298231</Numero>", "<Numero>298232</Numero>")

CHAMADAS = []


def _fake_baixar(pfx, senha, cnpj, ultimo_nsu=0, max_paginas=40, limite_seg=None,
                 estado=None):
    """ALFA tem 2 notas; BETA esbarra no consumo indevido; o resto, nada."""
    CHAMADAS.append((cnpj, limite_seg))
    if cnpj == CNPJ[tB]:
        raise adn_client.ConsumoIndevido("O ADN adiou a consulta por consumo indevido. Nada se perdeu.")
    if cnpj == CNPJ[tA] and ultimo_nsu < 10:
        yield (10, XML_NFSE)
        yield (11, XML_NFSE2)


adn_client.baixar_lote = _fake_baixar

_carregar_real = certificados.carregar_material


def _fake_carregar(db_, tenant_id, cnpj):
    if tenant_id == str(tC):
        raise ValueError("Certificado marcado como inativo.")
    return b"pfx-falso", "senha"


certificados.carregar_material = _fake_carregar
check("(massa montada)", True)


# ============================================================================
print("\n=== 4. A tela de antes tambem ganha o 'adiada' ===")

r = rota_nfse._buscar_uma_empresa(db, str(tB), CNPJ[tB], b"x", "s", False)
check("consumo indevido vira tipo_erro 'adiada'", r["tipo_erro"] == "adiada", r)
check("com a mensagem do ADN, nao '0 notas' calado", r["erro"] and "consumo indevido" in r["erro"])
cur = db.query(NfseCursor).filter_by(tenant_id=str(tB), cnpj=CNPJ[tB]).first()
# Primeira consulta da empresa: o rollback do erro desfaz também a criação do
# cursor. O que importa é que o NSU não andou.
check("o cursor nao andou", cur is None or cur.nsu == 0, cur.nsu if cur else None)
check("a tela continua com o teto de 60s por empresa",
      CHAMADAS[-1][1] == adn_client.LIMITE_EMPRESA_SEG, CHAMADAS[-1])


# ============================================================================
print("\n=== 5. Escopo e certificados ===")

a_todas = NfseAutomacao(id=str(uuid.uuid4()), nome="Todas", frequencia="diaria", hora="07:00",
                        escopo="todas", ativo=True)
nomes = [t.name for t in svc.empresas_do_escopo(db, a_todas)]
check("'todas' = ativas com certificado (inativo entra, para aparecer como falha)",
      nomes == ["ALFA SERVICOS", "BETA LTDA", "DELTA INDUSTRIA", "GAMA COMERCIO"], nomes)
check("empresa inativa fica fora", "EPSILON INATIVA" not in nomes)
check("empresa sem certificado fica fora do 'todas'", "ZETA SEM CERTIFICADO" not in nomes)
a_sel = NfseAutomacao(id=str(uuid.uuid4()), nome="Sel", frequencia="diaria", hora="07:00",
                      escopo="selecionadas", empresas=f"{tF},{tE},lixo", ativo=True)
nomes_sel = [t.name for t in svc.empresas_do_escopo(db, a_sel)]
check("'selecionadas' respeita a lista, sem inativa e sem id quebrado",
      nomes_sel == ["ZETA SEM CERTIFICADO"], nomes_sel)
venc = svc.certificados_vencendo(db, svc.empresas_do_escopo(db, a_todas), AGORA)
check("certificado vencendo em 10 dias entra no aviso",
      [v["empresa"] for v in venc] == ["DELTA INDUSTRIA"], venc)
check("com os dias contados", venc and venc[0]["dias"] in (9, 10) and not venc[0]["vencido"], venc)


# ============================================================================
print("\n=== 6. Uma rodada: cada empresa com o resultado dela ===")

db.add(a_todas)
db.commit()
CHAMADAS.clear()
eid = svc.executar(SessionLocal, a_todas.id, "manual")
db.expire_all()
e = db.get(NfseAutomacaoExecucao, eid)
check("a execucao foi registrada", e is not None)
check("status NAO e 'ok' com empresa adiada e com falha", e.status == "com_falhas", e.status)
check("contagens: 4 empresas, 2 ok, 1 falha, 1 adiada",
      (e.empresas_total, e.empresas_ok, e.empresas_falha, e.empresas_adiadas) == (4, 2, 1, 1),
      (e.empresas_total, e.empresas_ok, e.empresas_falha, e.empresas_adiadas))
check("2 notas novas (as da ALFA)", e.notas_novas == 2, e.notas_novas)
check("notas gravadas de verdade",
      db.query(NfseDocument).filter_by(tenant_id=str(tA), origem="adn").count() == 2)
check("o teto por empresa da coleta e o da automacao, nao os 60s da tela",
      all(lim == svc.TETO_EMPRESA_SEG for _c, lim in CHAMADAS), CHAMADAS)
check("certificado vencendo contado", e.certificados_vencendo == 1, e.certificados_vencendo)
detalhe = json.loads(e.detalhe)
por_nome = {x["empresa"]: x for x in detalhe["empresas"]}
check("a falha leva o motivo", por_nome["GAMA COMERCIO"]["erro"] == "Certificado marcado como inativo."
      and por_nome["GAMA COMERCIO"]["tipo_erro"] == "certificado", por_nome["GAMA COMERCIO"])
check("a adiada leva o tipo", por_nome["BETA LTDA"]["tipo_erro"] == "adiada")
check("a mensagem resume sem esconder",
      "1 com falha" in e.mensagem and "1 adiada" in e.mensagem, e.mensagem)
check("fim registrado", e.fim is not None and e.fim >= e.inicio)

# Tudo falhando: 'falhou'. Nada no escopo: 'vazia' — nunca 'ok'.
a_so_gama = NfseAutomacao(id=str(uuid.uuid4()), nome="So Gama", frequencia="diaria", hora="07:00",
                          escopo="selecionadas", empresas=str(tC), ativo=True)
a_vazia = NfseAutomacao(id=str(uuid.uuid4()), nome="Vazia", frequencia="diaria", hora="07:00",
                        escopo="selecionadas", empresas=str(tE), ativo=True)
db.add_all([a_so_gama, a_vazia])
db.commit()
e2 = db.get(NfseAutomacaoExecucao, svc.executar(SessionLocal, a_so_gama.id, "manual"))
check("todas falharam: status 'falhou'", e2.status == "falhou", e2.status)
e3 = db.get(NfseAutomacaoExecucao, svc.executar(SessionLocal, a_vazia.id, "manual"))
check("escopo vazio: status 'vazia', com o motivo", e3.status == "vazia"
      and "Nenhuma empresa" in e3.mensagem, (e3.status, e3.mensagem))

# Outra coleta andando: o botão recusa em vez de empilhar chamadas ao ADN.
svc._executando.acquire()
try:
    check("com coleta em andamento, executar sem esperar devolve None",
          svc.executar(SessionLocal, a_todas.id, "manual", esperar=False) is None)
    ok_, texto_ = svc.rodar_agora(SessionLocal, a_todas.id)
    check("e o 'rodar agora' diz que ja ha coleta", not ok_ and "andamento" in texto_, texto_)
finally:
    svc._executando.release()


# ============================================================================
print("\n=== 7. Nao roda em dobro; fora do ar, roda UMA vez ao voltar ===")

a_ag = NfseAutomacao(id=str(uuid.uuid4()), nome="Agenda", frequencia="diaria", hora="07:00",
                     escopo="selecionadas", empresas=str(tD), ativo=True,
                     proxima_execucao=datetime.utcnow() - timedelta(days=3))
db.add(a_ag)
db.commit()

s1, s2 = SessionLocal(), SessionLocal()
x1, x2 = s1.get(NfseAutomacao, a_ag.id), s2.get(NfseAutomacao, a_ag.id)
g1 = svc.reivindicar(s1, x1, datetime.utcnow())
g2 = svc.reivindicar(s2, x2, datetime.utcnow())
check("dois processos reivindicam a mesma vez: so um consegue", (g1, g2) == (True, False), (g1, g2))
s1.close()
s2.close()
db.expire_all()
check("e a proxima foi para o futuro", db.get(NfseAutomacao, a_ag.id).proxima_execucao > datetime.utcnow())

# Volta a proxima para 3 dias atras: o agendador roda UMA vez, nao tres.
db.get(NfseAutomacao, a_ag.id).proxima_execucao = datetime.utcnow() - timedelta(days=3)
db.commit()
antes = db.query(NfseAutomacaoExecucao).filter_by(automacao_id=a_ag.id).count()
n1 = svc.rodar_devidas(SessionLocal, datetime.utcnow())
n2 = svc.rodar_devidas(SessionLocal, datetime.utcnow())
depois = db.query(NfseAutomacaoExecucao).filter_by(automacao_id=a_ag.id).count()
check("tres dias perdidos viram UMA coleta", (n1, n2, depois - antes) == (1, 0, 1), (n1, n2, depois - antes))
ex_ag = (db.query(NfseAutomacaoExecucao).filter_by(automacao_id=a_ag.id)
         .order_by(NfseAutomacaoExecucao.inicio.desc()).first())
check("e ela sai marcada como 'agenda'", ex_ag.gatilho == "agenda", ex_ag.gatilho)

pausada = NfseAutomacao(id=str(uuid.uuid4()), nome="Pausada", frequencia="diaria", hora="07:00",
                        escopo="todas", ativo=False,
                        proxima_execucao=datetime.utcnow() - timedelta(hours=1))
db.add(pausada)
db.commit()
check("automacao pausada nao roda", svc.rodar_devidas(SessionLocal, datetime.utcnow()) == 0)


# ============================================================================
print("\n=== 8. Servidor que cai no meio; o agendador liga so no servidor ===")

db.add(NfseAutomacaoExecucao(id=str(uuid.uuid4()), automacao_id=a_todas.id, gatilho="agenda",
                             status="rodando", inicio=datetime.utcnow()))
db.commit()
check("coleta 'rodando' de antes do reinicio vira 'interrompida'",
      svc.marcar_interrompidas(db) == 1
      and db.query(NfseAutomacaoExecucao).filter_by(status="rodando").count() == 0)

check("importar o app NAO liga o agendador", svc._thread is None)
fonte_main = io.open(os.path.join("backend", "app", "main.py"), encoding="utf-8").read()
check("ele liga no lifespan do servidor", "lifespan=_ciclo_de_vida" in fonte_main
      and "iniciar_agendador" in fonte_main)
os.environ["SYNAX_AGENDADOR"] = "0"
check("SYNAX_AGENDADOR=0 desliga", svc.iniciar_agendador(SessionLocal) is False and svc._thread is None)
os.environ["SYNAX_AGENDADOR"] = "1"
svc.INTERVALO_SEG = 3600      # liga, mas nao chega a rodar um ciclo no teste
check("ligado, sobe UMA thread", svc.iniciar_agendador(SessionLocal) is True
      and svc._thread.is_alive())
check("chamar de novo nao sobe outra", svc.iniciar_agendador(SessionLocal) is False)
svc.parar_agendador()
svc._thread.join(timeout=5)
check("e para quando o servidor desce", not svc._thread.is_alive())
svc._thread = None
os.environ.pop("SYNAX_AGENDADOR", None)


# ============================================================================
print("\n=== 9. A tela: so admin, validacao visivel, historico ===")

TOKENS["adm"] = types.SimpleNamespace(id=str(u_adm), email="admin@synax.com")
TOKENS["ope"] = types.SimpleNamespace(id=str(u_ope), email="oper@synax.com")

from backend.app.api.dependencies.db import get_rls_db as _real_rls  # noqa: E402


def _fake_rls():
    ss = SessionLocal()
    try:
        yield ss
    finally:
        ss.close()


m.app.dependency_overrides[_real_rls] = _fake_rls
c = TestClient(m.app, raise_server_exceptions=False)


def as_(token):
    c.cookies.clear()
    if token:
        c.cookies.set("access_token", f"Bearer {token}")
    return c


PAGINA = io.open(os.path.join("frontend", "templates", "pages", "admin.html"),
                 encoding="utf-8").read()
check("a aba existe na Administracao", 'hx-get="/admin/automacoes"' in PAGINA)

r = as_("adm").get("/admin/automacoes")
check("admin abre a aba", r.status_code == 200 and "Automações" in r.text, r.status_code)
check("o historico aparece", "Histórico de coletas" in r.text and "Concluída com avisos" in r.text)
check("sem o agendador ligado, a tela avisa", "agendador não está ligado" in r.text)
check("operador nao entra", as_("ope").get("/admin/automacoes").status_code == 403)
check("nem grava", as_("ope").post("/admin/automacoes/salvar",
                                   data={"nome": "x", "hora": "07:00"}).status_code == 403)
check("anonimo barrado", as_(None).get("/admin/automacoes").status_code in (401, 403))

r = as_("adm").post("/admin/automacoes/salvar",
                    data={"nome": "Semanal", "frequencia": "semanal", "hora": "07:00", "escopo": "todas"})
check("semanal sem dia: 200 com o motivo no modal (nao 422)",
      r.status_code == 200 and "marque ao menos um dia" in r.text
      and r.headers.get("HX-Retarget") == "#modal-container", (r.status_code, r.headers.get("HX-Retarget")))
r = as_("adm").post("/admin/automacoes/salvar",
                    data={"nome": "Hora ruim", "frequencia": "diaria", "hora": "25:00", "escopo": "todas"})
check("hora invalida volta com o motivo", "HH:MM" in r.text)
r = as_("adm").post("/admin/automacoes/salvar",
                    data={"nome": "Forjada", "frequencia": "diaria", "hora": "07:00",
                          "escopo": "selecionadas", "empresas": [str(uuid.uuid4())]})
check("empresa forjada no formulario nao conta como escolhida", "Marque ao menos uma empresa" in r.text)
r = as_("adm").post("/admin/automacoes/salvar", data={"nome": "", "hora": "07:00"})
check("sem nome: pede o nome", "Dê um nome" in r.text)

r = as_("adm").post("/admin/automacoes/salvar",
                    data={"nome": "<script>alert(1)</script> Diaria", "frequencia": "diaria",
                          "hora": "06:30", "escopo": "selecionadas",
                          "empresas": [str(tA), str(tD)], "ativo": "1"})
nova = db.query(NfseAutomacao).filter(NfseAutomacao.nome.like("%Diaria")).first()
check("valida: grava", r.status_code == 200 and nova is not None, r.status_code)
check("com a proxima execucao calculada", nova and nova.proxima_execucao is not None
      and nova.proxima_execucao > datetime.utcnow())
check("e as empresas escolhidas", nova and set(svc.ids_selecionados(nova)) == {str(tA), str(tD)})
check("o nome e escapado na tela (XSS)",
      "&lt;script&gt;" in r.text and "<script>alert(1)" not in r.text)
check("o modal fecha por OOB e o toast confirma",
      'id="modal-container" hx-swap-oob' in r.text and "Automação criada" in r.text)

r = as_("adm").get(f"/admin/automacoes/{nova.id}/editar")
check("editar abre o modal preenchido", r.status_code == 200 and "Editar automação" in r.text
      and 'value="06:30"' in r.text)

r = as_("adm").post(f"/admin/automacoes/{nova.id}/ativo")
db.expire_all()
check("pausar zera a proxima e avisa",
      not db.get(NfseAutomacao, nova.id).ativo
      and db.get(NfseAutomacao, nova.id).proxima_execucao is None and "pausada" in r.text)
as_("adm").post(f"/admin/automacoes/{nova.id}/ativo")
db.expire_all()
check("reativar recalcula", db.get(NfseAutomacao, nova.id).proxima_execucao is not None)

CHAMADAS.clear()
r = as_("adm").post(f"/admin/automacoes/{nova.id}/rodar")
check("rodar agora responde na hora", r.status_code == 200 and "Coleta iniciada" in r.text)
check("e a tela passa a se atualizar sozinha", 'hx-trigger="every 4s"' in r.text)
svc._ultima_manual.join(timeout=30)
ex = (db.query(NfseAutomacaoExecucao).filter_by(automacao_id=nova.id, gatilho="manual")
      .order_by(NfseAutomacaoExecucao.inicio.desc()).first())
check("a coleta manual rodou e ficou no historico", ex is not None and ex.status != "rodando",
      ex.status if ex else None)

r = as_("adm").get(f"/admin/automacoes/execucao/{eid}")
check("o detalhe traz as tres secoes do informativo",
      all(x in r.text for x in ("Notas que chegaram", "O que não foi consultado",
                                "Certificados vencendo em até 30 dias")))
check("com a adiada nomeada como adiada", "Adiada pelo ADN" in r.text and "BETA LTDA" in r.text)
check("com o valor das notas que entraram (2 x 267,26)", "534,52" in r.text, "valor ausente")
check("e o certificado com o prazo", "DELTA INDUSTRIA" in r.text and "vence em" in r.text)

id_todas = a_todas.id
r = as_("adm").post(f"/admin/automacoes/{id_todas}/excluir")
db.expire_all()
check("excluir remove a automacao", db.get(NfseAutomacao, id_todas) is None)
check("mas mantem o historico", db.query(NfseAutomacaoExecucao).filter_by(automacao_id=id_todas).count() >= 1)
check("que passa a dizer 'automacao excluida'", "(automação excluída)" in r.text)


# ============================================================================
print("\n=== 10. Fase 2: a planilha da rodada, salva na pasta das NFS-e ===")

import openpyxl  # noqa: E402
from backend.app.services.nfse import informativo  # noqa: E402

db.expire_all()
e = db.get(NfseAutomacaoExecucao, eid)
check("a rodada guardou o caminho da planilha", bool(e.relatorio), e.relatorio)
caminho = os.path.join(_TMP_STORAGE, e.relatorio or "")
check("o arquivo esta na pasta das NFS-e, em _relatorios/AAAA-MM",
      os.path.exists(caminho) and e.relatorio.startswith("_relatorios"), e.relatorio)
wb = openpyxl.load_workbook(caminho)
check("as cinco abas", wb.sheetnames == ["Visão Geral", "Notas novas", "Por empresa",
                                         "Não consultadas", "Certificados"], wb.sheetnames)
ws = wb["Notas novas"]
check("as 2 notas novas, com a empresa na frente", ws.max_row == 3 and ws["A2"].value == "ALFA SERVICOS",
      (ws.max_row, ws["A2"].value))
titulos = [c.value for c in ws[1]]
col_toma = titulos.index("CNPJ Tomador") + 1
check("CNPJ sai como TEXTO (sem perder zero nem virar notacao cientifica)",
      ws.cell(row=2, column=col_toma).value == "48135251000100"
      and ws.cell(row=2, column=col_toma).number_format == "@",
      (ws.cell(row=2, column=col_toma).value, ws.cell(row=2, column=col_toma).number_format))
nao = [r[0] for r in wb["Não consultadas"].iter_rows(min_row=2, values_only=True)]
check("'Nao consultadas' traz a falha e a adiada", sorted(nao) == ["BETA LTDA", "GAMA COMERCIO"], nao)
check("'Certificados' traz a DELTA",
      [r[0] for r in wb["Certificados"].iter_rows(min_row=2, values_only=True)] == ["DELTA INDUSTRIA"])
check("sem destinatario, o e-mail fica registrado como nao enviado de proposito",
      e.email_status == "sem_destinatario", e.email_status)

r = as_("adm").get(f"/admin/automacoes/execucao/{eid}/planilha")
check("download: 200, xlsx de verdade", r.status_code == 200 and r.content[:2] == b"PK", r.status_code)
check("baixado, nunca renderizado (attachment + octet-stream + nosniff)",
      "attachment" in r.headers.get("content-disposition", "")
      and r.headers.get("content-type", "").startswith("application/octet-stream")
      and r.headers.get("x-content-type-options") == "nosniff", dict(r.headers))
check("operador nao baixa", as_("ope").get(f"/admin/automacoes/execucao/{eid}/planilha").status_code == 403)
DET = io.open(os.path.join("frontend", "templates", "partials", "admin", "modals",
                           "automacao_execucao.html"), encoding="utf-8").read()
check("o botao de baixar fica FORA do boost (senao o xlsx cai na tela)",
      'hx-boost="false"' in DET and "/planilha" in DET)

# Pasta da rede fora: a coleta segue, e a mensagem diz que a planilha nao foi salva.
_raiz_ok = os.environ["NFSE_STORAGE_DIR"]
_arquivo_no_lugar = os.path.join(_TMP_STORAGE, "nao-e-pasta.txt")
io.open(_arquivo_no_lugar, "w").write("x")
os.environ["NFSE_STORAGE_DIR"] = _arquivo_no_lugar
e_sem_rede = db.get(NfseAutomacaoExecucao, svc.executar(SessionLocal, a_so_gama.id, "manual"))
os.environ["NFSE_STORAGE_DIR"] = _raiz_ok
check("pasta da rede fora: relatorio vazio e a mensagem avisa",
      e_sem_rede.relatorio is None and "pasta da rede indisponível" in (e_sem_rede.mensagem or ""),
      (e_sem_rede.relatorio, e_sem_rede.mensagem))


# ============================================================================
print("\n=== 11. Fase 3: o informativo por e-mail ===")

import smtplib  # noqa: E402
from backend.app.services import email_envio  # noqa: E402
from backend.app.services import configuracoes as cfg  # noqa: E402
from backend.app.domain.models.configuracao import AppSetting  # noqa: E402

ENVIADOS, CHAMADAS_SMTP = [], []


class FakeSMTP:
    falhar_login = False
    recusar_conexao = False

    def __init__(self, host, porta, timeout=None, **kw):
        if FakeSMTP.recusar_conexao:
            raise ConnectionRefusedError("recusado")
        CHAMADAS_SMTP.append(("conectar", host, porta))

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def ehlo(self):
        pass

    def starttls(self, context=None):
        CHAMADAS_SMTP.append(("starttls",))

    def login(self, usuario, senha):
        CHAMADAS_SMTP.append(("login", usuario, senha))
        if FakeSMTP.falhar_login:
            raise smtplib.SMTPAuthenticationError(535, b"5.7.3 Authentication unsuccessful")

    def send_message(self, msg):
        ENVIADOS.append(msg)
        return {}


_smtp_real = email_envio.smtplib.SMTP
email_envio.smtplib.SMTP = FakeSMTP

# --- a configuracao, pela tela
r = as_("adm").get("/admin/configuracoes")
check("a aba Configuracoes tem o cartao do servidor de e-mail",
      "Servidor de e-mail (MCO2)" in r.text and "Não configurado" in r.text)
r = as_("adm").post("/admin/configuracoes/smtp", data={
    "smtp_host": "MCO2", "smtp_porta": "abc", "smtp_seguranca": "nenhuma"})
check("porta invalida volta com o motivo", "porta tem de ser um número" in r.text)
r = as_("adm").post("/admin/configuracoes/smtp", data={
    "smtp_host": "MCO2", "smtp_porta": "587", "smtp_seguranca": "starttls",
    "smtp_verificar_tls": "1", "smtp_usuario": "synax", "smtp_senha": "segredo-do-mco2",
    "smtp_remetente": "synax@vmcontadores.com.br", "smtp_remetente_nome": "Synax"})
check("salva e diz que salvou", r.status_code == 200 and "Servidor de e-mail salvo" in r.text)
check("a senha NUNCA volta para a tela", "segredo-do-mco2" not in r.text)
check("so a mascara", cfg.mascarar("segredo-do-mco2") in r.text)
bruta = db.query(AppSetting).filter(AppSetting.chave == cfg.SMTP_SENHA).first()
check("a senha fica CIFRADA no banco", bruta and bruta.segredo and bruta.valor != "segredo-do-mco2")
as_("adm").post("/admin/configuracoes/smtp", data={
    "smtp_host": "MCO2", "smtp_porta": "587", "smtp_seguranca": "starttls",
    "smtp_verificar_tls": "1", "smtp_usuario": "synax", "smtp_senha": "",
    "smtp_remetente": "synax@vmcontadores.com.br", "smtp_remetente_nome": "Synax"})
check("senha em branco MANTEM a gravada", cfg.config_smtp(db)["senha"] == "segredo-do-mco2")
check("operador nao mexe no servidor de e-mail",
      as_("ope").post("/admin/configuracoes/smtp", data={"smtp_host": "x"}).status_code == 403)

# --- o teste de envio
r = as_("adm").post("/admin/configuracoes/smtp/teste", data={"destinatario_teste": "sandro@vmcontadores.com.br"})
check("o teste sai pela configuracao gravada", "Teste enviado" in r.text and len(ENVIADOS) == 1
      and ("starttls",) in CHAMADAS_SMTP and ("login", "synax", "segredo-do-mco2") in CHAMADAS_SMTP,
      CHAMADAS_SMTP)
FakeSMTP.falhar_login = True
r = as_("adm").post("/admin/configuracoes/smtp/teste", data={"destinatario_teste": "sandro@vmcontadores.com.br"})
check("usuario/senha recusados: a tela diz em portugues", "recusou o usuário ou a senha" in r.text)
FakeSMTP.falhar_login = False
r = as_("adm").post("/admin/configuracoes/smtp/teste", data={"destinatario_teste": "nao-e-email"})
check("destinatario invalido no teste", "UM e-mail válido" in r.text)

# --- a automacao com destinatarios
r = as_("adm").post("/admin/automacoes/salvar", data={
    "nome": "Com e-mail", "frequencia": "diaria", "hora": "07:00", "escopo": "todas",
    "ativo": "1", "destinatarios": "fiscal@vm.com.br, isso-nao-e-email"})
check("destinatario invalido nao some em silencio: volta com o motivo",
      "E-mail inválido: isso-nao-e-email" in r.text)
r = as_("adm").post("/admin/automacoes/salvar", data={
    "nome": "Com e-mail", "frequencia": "diaria", "hora": "07:00", "escopo": "todas",
    "ativo": "1", "destinatarios": "fiscal@vm.com.br; sandro@vm.com.br\nfiscal@vm.com.br"})
com_email = db.query(NfseAutomacao).filter_by(nome="Com e-mail").first()
check("destinatarios gravados sem repeticao",
      com_email and com_email.destinatarios == "fiscal@vm.com.br,sandro@vm.com.br",
      com_email.destinatarios if com_email else None)
check("o cartao diz para quantas pessoas vai", "Informativo para 2 pessoas" in r.text)

ENVIADOS.clear()
e_mail = db.get(NfseAutomacaoExecucao, svc.executar(SessionLocal, com_email.id, "agenda"))
check("o informativo saiu e ficou registrado",
      e_mail.email_status == "enviado" and "fiscal@vm.com.br" in e_mail.email_detalhe
      and len(ENVIADOS) == 1, (e_mail.email_status, e_mail.email_detalhe))
msg = ENVIADOS[-1]
check("para os dois destinatarios", msg["To"] == "fiscal@vm.com.br, sandro@vm.com.br", msg["To"])
check("o assunto diz o essencial (notas e problemas)",
      "[Synax] NFS-e" in msg["Subject"] and "falha" in msg["Subject"]
      and "adiada" in msg["Subject"] and "certificado" in msg["Subject"], msg["Subject"])
html_mail = msg.get_body(preferencelist=("html",)).get_content()
texto_mail = msg.get_body(preferencelist=("plain",)).get_content()
check("o HTML traz as TRES secoes",
      all(s in html_mail for s in ("1. Notas que chegaram", "2. O que não foi consultado",
                                   "3. Certificados vencendo em até 30 dias")))
check("e o texto puro tambem", all(s in texto_mail for s in ("1. NOTAS QUE CHEGARAM",
                                                            "2. O QUE NÃO FOI CONSULTADO",
                                                            "3. CERTIFICADOS VENCENDO")))
check("com a adiada e o certificado nomeados",
      "Adiada pelo ADN" in html_mail and "DELTA INDUSTRIA" in html_mail)
anexos = list(msg.iter_attachments())
check("com a planilha anexada", len(anexos) == 1 and anexos[0].get_filename().endswith(".xlsx")
      and anexos[0].get_content()[:2] == b"PK", [a.get_filename() for a in anexos])

# Falha de envio: a COLETA continua com o status dela; o e-mail registra o motivo.
FakeSMTP.recusar_conexao = True
e_falha = db.get(NfseAutomacaoExecucao, svc.executar(SessionLocal, com_email.id, "agenda"))
FakeSMTP.recusar_conexao = False
check("servidor fora: o e-mail sai 'falhou' com o motivo",
      e_falha.email_status == "falhou" and "Não foi possível conectar a MCO2:587" in e_falha.email_detalhe,
      (e_falha.email_status, e_falha.email_detalhe))
check("e a coleta nao vira 'falhou' por causa do e-mail", e_falha.status == "com_falhas", e_falha.status)
r = as_("adm").get(f"/admin/automacoes/execucao/{e_falha.id}")
check("o detalhe mostra que o informativo nao saiu", "Informativo não enviado" in r.text)
r = as_("adm").get("/admin/automacoes")
check("e o historico marca o e-mail que falhou", "O informativo não saiu" in r.text)

# Rodada so com falha: o assunto nao pode parecer tudo certo.
check("rodada que falhou inteira tem assunto de alerta",
      informativo.assunto(e2).startswith("⚠"), informativo.assunto(e2))

# Sem servidor configurado: a coleta roda, o e-mail diz por que nao saiu.
cfg.gravar(db, cfg.SMTP_HOST, "", "teste")
e_sem_smtp = db.get(NfseAutomacaoExecucao, svc.executar(SessionLocal, com_email.id, "agenda"))
check("sem servidor de e-mail: 'falhou' dizendo onde configurar",
      e_sem_smtp.email_status == "falhou" and "Configurações" in e_sem_smtp.email_detalhe,
      e_sem_smtp.email_detalhe)
r = as_("adm").get("/admin/automacoes")
check("e a aba avisa que ha destinatario sem servidor", "servidor de e-mail não está configurado" in r.text)
email_envio.smtplib.SMTP = _smtp_real

# O HTML do e-mail escapa o que vem de cadastro (nome de empresa).
from types import SimpleNamespace  # noqa: E402
_falsa = SimpleNamespace(id="x" * 36, inicio=datetime.utcnow(), fim=datetime.utcnow(), gatilho="agenda",
                         status="ok", mensagem="ok", notas_novas=1, notas_atualizadas=0,
                         empresas_ok=1, empresas_total=1, empresas_falha=0, empresas_adiadas=0,
                         certificados_vencendo=0, relatorio=None,
                         detalhe=json.dumps({"empresas": [{"tenant_id": "t", "empresa": "<script>x</script>",
                                                           "cnpj": "1", "novos": 1, "atualizados": 0}]}))
_a, _h, _t = informativo.montar(db, _falsa, "Teste", "Concluída")
check("nome de empresa com HTML sai escapado no e-mail", "<script>x" not in _h and "&lt;script&gt;" in _h)


# ============================================================================
print("\n=== 12. Colunas novas chegam a banco que ja tinha as tabelas ===")

from sqlalchemy import create_engine, inspect as _inspect, text  # noqa: E402
from backend.app.infrastructure.database import schema_upgrades  # noqa: E402

CAMINHO_VELHO = "./_automacao_schema_velho.db"
if os.path.exists(CAMINHO_VELHO):
    os.remove(CAMINHO_VELHO)
eng_velho = create_engine("sqlite:///" + CAMINHO_VELHO)
for modelo, tabela in ((NfseAutomacao, "nfse_automacoes"),
                       (NfseAutomacaoExecucao, "nfse_automacao_execucoes")):
    modelo.__table__.create(bind=eng_velho)
    registradas = [n for n, _ddl in schema_upgrades.COLUNAS_NOVAS.get(tabela, [])]
    with eng_velho.begin() as _conn:
        for _nome in registradas:
            _conn.execute(text(f"ALTER TABLE {tabela} DROP COLUMN {_nome}"))
schema_upgrades.aplicar(eng_velho)
for modelo, tabela in ((NfseAutomacao, "nfse_automacoes"),
                       (NfseAutomacaoExecucao, "nfse_automacao_execucoes")):
    faltando = ({c.name for c in modelo.__table__.columns}
                - {c["name"] for c in _inspect(eng_velho).get_columns(tabela)})
    check(f"{tabela}: nenhuma coluna do model falta depois do startup", not faltando, faltando)
eng_velho.dispose()
try:
    os.remove(CAMINHO_VELHO)
except OSError:
    pass


# ============================================================================
adn_client.baixar_lote = rota_nfse.adn_client.baixar_lote  # (mesmo modulo)
certificados.carregar_material = _carregar_real
db.close()
_eng.dispose()
shutil.rmtree(_TMP_STORAGE, ignore_errors=True)
if os.path.exists("_automacao_test.db"):
    try:
        os.remove("_automacao_test.db")
    except OSError:
        pass

print("\n" + "=" * 62)
print(f"RESULTADO: {OK}/{OK + FALHAS} passaram")
for f in _falhou:
    print("  -", f)
sys.exit(1 if FALHAS else 0)
