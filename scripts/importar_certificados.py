"""Importa a planilha de senhas de certificados para o banco, CIFRADAS.

Uso (a partir da raiz do projeto):
    ./venv/Scripts/python.exe scripts/importar_certificados.py --simular
    ./venv/Scripts/python.exe scripts/importar_certificados.py --aplicar

O que faz:
  1. lê a planilha (CNPJ, apelido, razão, senha);
  2. casa cada CNPJ com uma empresa (tenants.cnpj) já cadastrada no sistema;
  3. procura o arquivo .pfx correspondente na pasta de certificados;
  4. grava em tenant_certificates com a senha cifrada (Fernet).

NENHUMA senha é exibida na saída. O modo padrão é apenas simulação.
"""
import argparse
import os
import re
import sys

sys.path.insert(0, os.path.abspath("."))

from dotenv import load_dotenv
load_dotenv()

import openpyxl
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from backend.app.core.security import encrypt_data
from backend.app.domain.models.auth_models import Tenant
from backend.app.domain.models.certificates import TenantCertificate
from backend.app.services.nfse import certificados

PLANILHA_PADRAO = (r"\\srv-dc02\documentos\Contratos e Alterações"
                   r"\__Certificados\_Certificados Senha.xlsx")


def ler_planilha(caminho: str) -> list[dict]:
    ws = openpyxl.load_workbook(caminho, data_only=True).active
    linhas, vistos = [], set()
    # Linha 1 é vazia; o cabeçalho está na 2 (CNPJ | Apelido | RAZÃO | SENHA).
    for r in range(3, ws.max_row + 1):
        cnpj = re.sub(r"\D", "", str(ws.cell(r, 2).value or ""))
        senha = str(ws.cell(r, 5).value or "").strip()
        if len(cnpj) != 14 or not senha or cnpj in vistos:
            continue
        vistos.add(cnpj)
        linhas.append({
            "cnpj": cnpj,
            "apelido": str(ws.cell(r, 3).value or "").strip(),
            "razao": str(ws.cell(r, 4).value or "").strip(),
            "senha": senha,
        })
    return linhas


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--planilha", default=os.environ.get("CERT_PLANILHA", PLANILHA_PADRAO))
    ap.add_argument("--aplicar", action="store_true", help="grava no banco")
    ap.add_argument("--simular", action="store_true", help="apenas relata (padrão)")
    args = ap.parse_args()
    aplicar = args.aplicar and not args.simular

    if not os.path.exists(args.planilha):
        print(f"ERRO: planilha não encontrada em {args.planilha}")
        return 2

    registros = ler_planilha(args.planilha)
    print(f"planilha            : {len(registros)} CNPJ(s) únicos com senha")

    arquivos = certificados.listar_arquivos()
    print(f"pasta de certificados: {len(arquivos)} arquivo(s) em {certificados.pasta()}")
    if not arquivos:
        print("AVISO: nenhum arquivo lido — o vínculo de arquivo ficará vazio.")

    url = os.environ["DATABASE_URL"].replace("postgres://", "postgresql://", 1)
    Sessao = sessionmaker(bind=create_engine(url))
    db = Sessao()

    empresas = {re.sub(r"\D", "", t.cnpj or ""): t for t in db.query(Tenant).all()}
    print(f"empresas no sistema  : {len(empresas)}")
    print()

    novos = atualizados = sem_empresa = sem_arquivo = 0
    for reg in registros:
        empresa = empresas.get(reg["cnpj"])
        if not empresa:
            sem_empresa += 1
            continue

        nome_arquivo = certificados.localizar_por_cnpj(reg["cnpj"], arquivos)
        if not nome_arquivo:
            sem_arquivo += 1

        existente = db.query(TenantCertificate).filter_by(
            tenant_id=str(empresa.id), cnpj=reg["cnpj"]).first()

        if existente:
            atualizados += 1
            if aplicar:
                existente.senha_encrypted = encrypt_data(reg["senha"])
                existente.apelido = reg["apelido"] or existente.apelido
                existente.razao_social = reg["razao"] or existente.razao_social
                if nome_arquivo:
                    existente.arquivo = nome_arquivo
        else:
            novos += 1
            if aplicar:
                db.add(TenantCertificate(
                    tenant_id=str(empresa.id), cnpj=reg["cnpj"],
                    arquivo=nome_arquivo,
                    senha_encrypted=encrypt_data(reg["senha"]),
                    apelido=reg["apelido"], razao_social=reg["razao"],
                    ativo=True,
                ))

    if aplicar:
        db.commit()

    print(f"a criar/criados     : {novos}")
    print(f"a atualizar/atualiz.: {atualizados}")
    print(f"sem empresa no Synax: {sem_empresa}  (CNPJ da planilha não cadastrado)")
    print(f"sem arquivo .pfx    : {sem_arquivo}  (vincular manualmente depois)")
    print()
    print("MODO: " + ("APLICADO no banco" if aplicar else "SIMULAÇÃO — nada gravado"))
    if not aplicar:
        print("Para gravar de fato, rode novamente com --aplicar")

    db.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
